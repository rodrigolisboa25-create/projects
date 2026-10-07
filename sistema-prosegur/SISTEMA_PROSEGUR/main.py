import json
import os
import re
import shutil
import sqlite3
import time
import traceback
import uuid
from collections import defaultdict
from copy import copy
from dataclasses import asdict
from datetime import datetime, timedelta
from typing import Any, Dict, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook

from processador_pdf import DocumentPayload, analisar_pdf_extrair_dados
from agente_auditoria import executar_auditoria

SISTEMA_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(SISTEMA_DIR)
INPUT_DIR = os.path.join(ROOT_DIR, 'Idocs_Prosegur')
OUTPUT_DIR = os.path.join(ROOT_DIR, 'PROSEGUR_PROCESS')
DB_PATH = os.path.join(SISTEMA_DIR, 'documentos.db')
LOG_FILE = os.path.join(SISTEMA_DIR, 'processamento.log')
METRICS_FILE = os.path.join(SISTEMA_DIR, 'vektor_ar_metricas.json')
BASE_FORN_PATH = os.path.join(ROOT_DIR, 'Bases de informações.xlsx')
BASE_FORN_FALLBACK_PATH = os.path.join(SISTEMA_DIR, 'Base_Forn.xlsx')
TEMPLATE_DIR = os.path.join(SISTEMA_DIR, 'templates')
TEMPLATE_SAP = {
    'SBF': os.path.join(TEMPLATE_DIR, 'Planilha_SAP_SBF.xlsx'),
    'FISIA': os.path.join(TEMPLATE_DIR, 'Planilha_SAP_FISIA.xlsx'),
}


def logar(msg: str) -> None:
    ts = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    line = f'[{ts}] {msg}'
    with open(LOG_FILE, 'a', encoding='utf-8') as file:
        file.write(line + '\n')
    print(line)


def safe_float(value: Any) -> float:
    try:
        if value in (None, ''):
            return 0.0
        return float(value)
    except Exception:
        return 0.0


def periodo_atual() -> str:
    return datetime.now().strftime('%m-%Y')


def get_periodo(payload_dict: Dict[str, Any]) -> str:
    for key in ('periodo', 'PERIODO'):
        value = str(payload_dict.get(key) or '').strip()
        if re.fullmatch(r'\d{2}-\d{4}', value):
            return value
    return periodo_atual()


def normalizar_texto_chave(value: Any) -> str:
    text = str(value or '').strip().upper()
    text = re.sub(r'\.0$', '', text)
    return re.sub(r'\W+', '', text, flags=re.UNICODE)


def valor_por_coluna(row: Dict[str, Any], header: Any) -> Any:
    if header in row:
        return row.get(header)

    header_norm = normalizar_texto_chave(header)
    for key, value in row.items():
        if normalizar_texto_chave(key) == header_norm:
            return value

    return ''


def valor_por_alias(row: Dict[str, Any], aliases: list[str]) -> Any:
    for alias in aliases:
        value = valor_por_coluna(row, alias)
        if value not in (None, ''):
            return value
    return ''


def numero_nota_para_check(row: Dict[str, Any]) -> Any:
    return valor_por_alias(
        row,
        ['Nº Nota Fiscal Eletrônica', 'N° Nota Fiscal Eletrônica', 'Nº DOCUMENTO']
    )


def numero_nota_normalizado(value: Any) -> str:
    return re.sub(r'\D', '', str(value or '').strip())


def formatar_cnpj(value: Any) -> str:
    digits = re.sub(r'\D', '', str(value or '').strip())
    if not digits:
        return ''
    digits = digits.zfill(14)[-14:]
    return f'{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}'


def normalizar_empresa(value: Any) -> str:
    text = str(value or '').strip().upper()
    if 'FISIA' in text or 'FÍSIA' in text:
        return 'FISIA'
    if 'SBF' in text or 'CENTAURO' in text:
        return 'SBF'
    return text


def periodo_para_nome_arquivo(periodo: str) -> str:
    match = re.fullmatch(r'(\d{2})-(\d{4})', str(periodo or '').strip())
    if not match:
        return str(periodo or '').strip()
    return f'{match.group(1)}.{match.group(2)}'


def chave_linha_master(sheet_name: str, row: Dict[str, Any]) -> tuple:
    periodo = normalizar_texto_chave(valor_por_alias(row, ['PERIODO']))
    empresa = normalizar_texto_chave(valor_por_alias(row, ['EMPRESA']))

    if sheet_name == 'Notas_e_CTEs':
        tipo_doc = normalizar_texto_chave(valor_por_alias(row, ['TIPO_DOC']))
        numero = normalizar_texto_chave(valor_por_alias(row, ['Nº Nota Fiscal Eletrônica', 'N° Nota Fiscal Eletrônica', 'Nº DOCUMENTO']))
        lote = normalizar_texto_chave(valor_por_alias(row, ['Lote Fatura']))
        return (periodo, empresa, tipo_doc, numero, lote)

    if sheet_name == 'Boletos':
        numero = normalizar_texto_chave(valor_por_alias(row, ['Nº DOCUMENTO', 'N° DOCUMENTO']))
        lote = normalizar_texto_chave(valor_por_alias(row, ['Lote Fatura']))
        return (periodo, empresa, numero, lote)

    if sheet_name == 'Demonstrativos':
        nota = normalizar_texto_chave(valor_por_alias(row, ['Nota Fiscal']))
        loja = normalizar_texto_chave(valor_por_alias(row, ['LOJA']))
        pagina = normalizar_texto_chave(valor_por_alias(row, ['Página Item', 'Pagina Item']))
        guia = normalizar_texto_chave(valor_por_alias(row, ['Guia']))
        data = normalizar_texto_chave(valor_por_alias(row, ['Data']))
        valor = f"{safe_float(valor_por_alias(row, ['Valor a faturar', 'Valor da Nota', 'Valor'])):.2f}"
        return (periodo, empresa, nota, loja, pagina, guia, data, valor)

    if sheet_name == 'Pendencias_Auditoria':
        arquivo = normalizar_texto_chave(valor_por_alias(row, ['Arquivo']))
        return (periodo, empresa, arquivo)

    return tuple(normalizar_texto_chave(value) for value in row.values())


def append_rows_to_master(master_path: str, sheets_rows: Dict[str, list[Dict[str, Any]]]) -> Dict[str, Dict[str, int]]:
    if os.path.exists(master_path):
        wb = load_workbook(master_path)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)

    stats: Dict[str, Dict[str, int]] = {}

    for sheet_name, rows in sheets_rows.items():
        if not rows:
            continue

        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            garantir_colunas_calculadas(ws)
            headers = [cell.value for cell in ws[1]]
        else:
            ws = wb.create_sheet(sheet_name)
            headers = list(rows[0].keys())
            ws.append(headers)
            garantir_colunas_calculadas(ws)
            headers = [cell.value for cell in ws[1]]

        existing_keys = set()
        for values in ws.iter_rows(min_row=2, values_only=True):
            existing_row = {headers[idx]: values[idx] if idx < len(values) else None for idx in range(len(headers))}
            existing_keys.add(chave_linha_master(sheet_name, existing_row))

        inserted = 0
        skipped = 0
        for row in rows:
            row_key = chave_linha_master(sheet_name, row)
            if row_key in existing_keys:
                skipped += 1
                continue

            ws.append([valor_por_coluna(row, header) for header in headers])
            existing_keys.add(row_key)
            inserted += 1

        stats[sheet_name] = {'inserted': inserted, 'skipped_existing': skipped}

    os.makedirs(os.path.dirname(master_path), exist_ok=True)
    wb.save(master_path)
    wb.close()
    return stats


def periodo_da_linha_db(row: pd.Series, doc: Dict[str, Any]) -> str:
    periodo = str(row.get('periodo') or doc.get('periodo') or '').strip()
    if re.fullmatch(r'\d{2}-\d{4}', periodo):
        return periodo
    return periodo_atual()


def montar_mapa_loja_por_nota(docs: list[Dict[str, Any]]) -> Dict[str, Any]:
    mapa: Dict[str, Any] = {}
    for doc in docs:
        if doc.get('document_type') not in ('NFSE', 'CTE_OS'):
            continue

        numero = numero_nota_normalizado(doc.get('numero_documento'))
        loja = doc.get('numero_loja')
        if numero and loja not in (None, ''):
            mapa[numero] = loja

    return mapa


def montar_mapa_cnpj_prestador_por_nota(notas_ctes: list[Dict[str, Any]]) -> Dict[str, str]:
    mapa: Dict[str, str] = {}
    for row in notas_ctes:
        numero = numero_nota_normalizado(valor_por_alias(row, ['CHECK LOJA']))
        cnpj = formatar_cnpj(valor_por_alias(row, ['CNPJ PRESTADOR']))
        if numero and cnpj:
            mapa[numero] = cnpj
    return mapa


def aplicar_cnpj_prestador_demonstrativos(
    demonstrativos: list[Dict[str, Any]],
    cnpj_por_nota: Dict[str, str],
) -> None:
    for row in demonstrativos:
        numero = numero_nota_normalizado(valor_por_alias(row, ['Nota Fiscal']))
        row['CNPJ PRESTADOR'] = cnpj_por_nota.get(numero, '')


def corrigir_formula_lofa_fat(ws) -> None:
    headers = [cell.value for cell in ws[1]]
    if ws.title != 'Demonstrativos' or 'LOFA FAT' not in headers or 'Nota Fiscal' not in headers:
        return

    lofa_col = headers.index('LOFA FAT') + 1
    nota_col_letter = ws.cell(row=1, column=headers.index('Nota Fiscal') + 1).column_letter
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=lofa_col)
        if isinstance(cell.value, str) and cell.value.startswith('=VLOOKUP('):
            cell.value = f'=VLOOKUP({nota_col_letter}{row_idx},Notas_e_CTEs!$C:$D,2,0)'


def garantir_colunas_calculadas(ws) -> None:
    required = {
        'Notas_e_CTEs': [('CHECK LOJA', 3)],
        'Demonstrativos': [('LOFA FAT', 3), ('CNPJ PRESTADOR', 4)],
    }

    for header, position in required.get(ws.title, []):
        headers = [cell.value for cell in ws[1]]
        if header in headers:
            continue

        ws.insert_cols(position)
        ws.cell(row=1, column=position, value=header)

    corrigir_formula_lofa_fat(ws)


def copiar_estilo_celula(origem, destino) -> None:
    if origem.has_style:
        destino._style = copy(origem._style)
    if origem.number_format:
        destino.number_format = origem.number_format
    if origem.alignment:
        destino.alignment = copy(origem.alignment)
    if origem.protection:
        destino.protection = copy(origem.protection)


def copiar_modelo_linha(ws, linha_modelo: int, linha_destino: int, max_col: int) -> None:
    ws.row_dimensions[linha_destino].height = ws.row_dimensions[linha_modelo].height
    for col_idx in range(1, max_col + 1):
        if ws.cell(linha_destino, col_idx).__class__.__name__ == 'MergedCell':
            continue
        copiar_estilo_celula(ws.cell(linha_modelo, col_idx), ws.cell(linha_destino, col_idx))


def limpar_valores(ws, min_row: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, min_col=1, max_col=max_col):
        for cell in row:
            if cell.__class__.__name__ == 'MergedCell':
                continue
            cell.value = None


def carregar_bases_sap() -> tuple[Dict[str, Any], Dict[str, Any], list[list[Any]], list[list[Any]]]:
    fornecedor_por_cnpj: Dict[str, Any] = {}
    cc_fisia_por_loja: Dict[str, Any] = {}
    fornecedores_base: list[list[Any]] = [['CNPJ FORN', 'CODIGO FORN']]
    cc_base: list[list[Any]] = [['LOJA FISIA', 'CC FISIA']]

    base_path = BASE_FORN_PATH if os.path.exists(BASE_FORN_PATH) else BASE_FORN_FALLBACK_PATH
    if not os.path.exists(base_path):
        return fornecedor_por_cnpj, cc_fisia_por_loja, fornecedores_base, cc_base

    wb = load_workbook(base_path, read_only=True, data_only=True)
    try:
        ws_fornecedores = wb['forns prosegur']
        for row in ws_fornecedores.iter_rows(min_row=2, values_only=True):
            cnpj = formatar_cnpj(row[1] or row[0])
            codigo = row[3] if len(row) > 3 else None
            if cnpj and codigo not in (None, ''):
                fornecedor_por_cnpj.setdefault(cnpj, codigo)
                fornecedores_base.append([cnpj, codigo])

        ws_cc = wb[wb.sheetnames[2]]
        for loja, cc, *_ in ws_cc.iter_rows(min_row=2, values_only=True):
            loja_norm = numero_nota_normalizado(loja)
            if loja_norm and cc not in (None, ''):
                cc_fisia_por_loja.setdefault(loja_norm, cc)
                cc_base.append([loja, cc])
    finally:
        wb.close()

    return fornecedor_por_cnpj, cc_fisia_por_loja, fornecedores_base, cc_base


def criar_aba_bases_informacoes(wb, fornecedores_base: list[list[Any]], cc_base: list[list[Any]]) -> None:
    sheet_name = 'Bases de informações'
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(sheet_name)

    for row_idx, row in enumerate(fornecedores_base, start=1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for row_idx, row in enumerate(cc_base, start=1):
        for col_idx, value in enumerate(row, start=4):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col, width in {'A': 20, 'B': 14, 'D': 12, 'E': 12}.items():
        ws.column_dimensions[col].width = width


def remover_planilha_auxiliar_template(wb) -> None:
    if 'Planilha2' in wb.sheetnames:
        wb.remove(wb['Planilha2'])


def referencia_base_informacoes(sheet_name: str, range_ref: str) -> str:
    base_path = BASE_FORN_PATH if os.path.exists(BASE_FORN_PATH) else BASE_FORN_FALLBACK_PATH
    folder = os.path.dirname(base_path)
    filename = os.path.basename(base_path)
    return f"'{folder}\\[{filename}]{sheet_name}'!{range_ref}"


def classificar_tipo_servico_sap(tipo_servico: Any) -> str:
    texto = str(tipo_servico or '').upper()
    if 'MANUSEIO' in texto and 'TESOURARIA' in texto:
        return 'CONTAGEM DE VALORES'
    return 'TRANSPORTE DE VALORES'


def dados_operacao_por_tipo(tipo_serv: str) -> tuple[int, str, str, str]:
    if tipo_serv == 'CONTAGEM DE VALORES':
        return 3002991, 'CONTAGEM E PROCESSAMENTO DE VALORES', 'PROCESSAMENTO DE VALORES', 'RR 0135/12'
    return 3002990, 'SERV DE TRANSPORTE DE VALORES', 'TRANSPORTE DE VALORES', 'RR 1175/11'


def calcular_cc_sbf(loja_rateio: Any) -> str:
    loja = numero_nota_normalizado(loja_rateio)
    if not loja:
        return ''
    return f'7010{int(loja):04d}01'


def normalizar_loja_rateio_sap(empresa: str, loja_rateio: Any) -> Any:
    loja = numero_nota_normalizado(loja_rateio)
    if empresa == 'FISIA' and loja == '5092':
        return 2070
    if loja:
        return int(loja)
    return loja_rateio


def valor_lofa_fat(row: Dict[str, Any], loja_por_nota: Dict[str, Any]) -> Any:
    valor = valor_por_coluna(row, 'LOFA FAT')
    if isinstance(valor, str) and (valor.startswith('=') or valor.strip().upper() in ('#N/D', '#N/A')):
        nota = numero_nota_normalizado(valor_por_coluna(row, 'Nota Fiscal'))
        return loja_por_nota.get(nota, '')
    return valor


def montar_linhas_sap(
    master_path: str,
    periodo: str,
    empresa: str,
    fornecedor_por_cnpj: Dict[str, Any],
    cc_fisia_por_loja: Dict[str, Any],
) -> list[Dict[str, Any]]:
    wb = load_workbook(master_path, read_only=True, data_only=True)
    try:
        ws_notas = wb['Notas_e_CTEs']
        headers_notas = [cell.value for cell in ws_notas[1]]
        rows_notas = []
        for values in ws_notas.iter_rows(min_row=2, values_only=True):
            rows_notas.append({headers_notas[idx]: values[idx] if idx < len(values) else None for idx in range(len(headers_notas))})

        loja_por_nota = {}
        for row in rows_notas:
            nota = numero_nota_normalizado(valor_por_coluna(row, 'CHECK LOJA'))
            loja = valor_por_coluna(row, 'LOJA')
            if nota and loja not in (None, ''):
                loja_por_nota[nota] = loja

        ws_dem = wb['Demonstrativos']
        headers_dem = [cell.value for cell in ws_dem[1]]
        linhas: list[Dict[str, Any]] = []
        empresa_norm = normalizar_empresa(empresa)

        for values in ws_dem.iter_rows(min_row=2, values_only=True):
            row = {headers_dem[idx]: values[idx] if idx < len(values) else None for idx in range(len(headers_dem))}
            if str(valor_por_coluna(row, 'PERIODO') or '').strip() != periodo:
                continue
            if normalizar_empresa(valor_por_coluna(row, 'EMPRESA')) != empresa_norm:
                continue

            cnpj = formatar_cnpj(valor_por_coluna(row, 'CNPJ PRESTADOR'))
            loja_fat = valor_lofa_fat(row, loja_por_nota)
            loja_rateio = normalizar_loja_rateio_sap(empresa_norm, valor_por_coluna(row, 'LOJA'))
            nota = valor_por_coluna(row, 'Nota Fiscal')
            valor = safe_float(valor_por_coluna(row, 'Total a Faturar'))
            tipo_serv = classificar_tipo_servico_sap(valor_por_coluna(row, 'Tipo de Serviço'))
            cod_op, servico, resumao, rr = dados_operacao_por_tipo(tipo_serv)
            codigo_fornecedor = fornecedor_por_cnpj.get(cnpj, '')
            if empresa_norm == 'SBF':
                cc = calcular_cc_sbf(loja_rateio)
                empresa_sap = 'SBF 7010'
            else:
                cc = cc_fisia_por_loja.get(numero_nota_normalizado(loja_rateio), '')
                empresa_sap = 'FISIA 7170'

            linhas.append({
                'EMPRESA': empresa_sap,
                'CNPJ FORN': cnpj,
                'CÓDIGO FORN': codigo_fornecedor,
                'LOJA FAT': loja_fat,
                'LOJA RATEIO': loja_rateio,
                'TIPO SERV': tipo_serv,
                'NF': nota,
                'TIPO NF': 'NF',
                'FATURADO': None,
                'VALOR': valor,
                'CC': cc,
                'FORN': codigo_fornecedor,
                'NÚMERO DA NFE': nota,
                'LOJA': loja_fat,
                'LOJA RATEIO2': loja_rateio,
                'CC2': cc,
                'COD.OP': cod_op,
                'SERVICO': servico,
                'RESUMÃO': resumao,
                'RR': rr,
            })
    finally:
        wb.close()

    return linhas


def escrever_linhas_aba_sap(ws, linhas: list[Dict[str, Any]], empresa: str) -> None:
    headers = [ws.cell(1, col_idx).value for col_idx in range(1, 21)]
    limpar_valores(ws, 2, 20)
    ref_fornecedor = referencia_base_informacoes('forns prosegur', '$B:$D')
    ref_cc_fisia = referencia_base_informacoes('cc físia', '$A:$B')
    for row_offset, row in enumerate(linhas, start=2):
        if row_offset > ws.max_row:
            ws.append([None] * 20)
        copiar_modelo_linha(ws, 2, row_offset, 20)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_offset, column=col_idx)
            if header == 'CÓDIGO FORN':
                cell.value = f'=VLOOKUP(B{row_offset},{ref_fornecedor},3,0)'
            elif header == 'CC' and col_idx == 11 and empresa == 'SBF':
                cell.value = f'=CONCATENATE(7010,TEXT(E{row_offset},"0000"),"01")'
            elif header == 'CC' and col_idx == 11:
                cell.value = f'=VLOOKUP(E{row_offset},{ref_cc_fisia},2,0)'
            elif col_idx == 12:
                cell.value = f'=C{row_offset}'
            elif col_idx == 13:
                cell.value = f'=G{row_offset}'
            elif col_idx == 14:
                cell.value = f'=D{row_offset}'
            elif col_idx == 15:
                cell.value = f'=E{row_offset}'
            elif col_idx == 16:
                cell.value = f'=K{row_offset}'
            else:
                cell.value = row.get(header)

            if col_idx in (4, 5, 7, 13, 14, 15):
                cell.number_format = 'General'


def consolidar_linhas_dinamica(linhas: list[Dict[str, Any]]) -> list[list[Any]]:
    totais: Dict[tuple, float] = defaultdict(float)
    ordem = ['FORN', 'NÚMERO DA NFE', 'LOJA', 'LOJA RATEIO2', 'CC2', 'COD.OP', 'SERVICO', 'RESUMÃO', 'RR']
    for row in linhas:
        key = tuple(row.get(col) for col in ordem)
        totais[key] += safe_float(row.get('VALOR'))

    resultado = []
    for key in sorted(totais.keys(), key=lambda item: tuple(str(part) for part in item)):
        total = round(totais[key], 2)
        if abs(total) < 0.000001:
            continue
        resultado.append(list(key) + [total])
    return resultado


def total_master_demonstrativos(master_path: str, periodo: str, empresa: str) -> tuple[int, float]:
    wb = load_workbook(master_path, read_only=True, data_only=True)
    try:
        ws = wb['Demonstrativos']
        headers = [cell.value for cell in ws[1]]
        total = 0.0
        count = 0
        empresa_norm = normalizar_empresa(empresa)
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = {headers[idx]: values[idx] if idx < len(values) else None for idx in range(len(headers))}
            if str(valor_por_coluna(row, 'PERIODO') or '').strip() != periodo:
                continue
            if normalizar_empresa(valor_por_coluna(row, 'EMPRESA')) != empresa_norm:
                continue
            total += safe_float(valor_por_coluna(row, 'Total a Faturar'))
            count += 1
        return count, round(total, 2)
    finally:
        wb.close()


def validar_total_relatorio_sap(master_path: str, periodo: str, empresa: str, linhas: list[Dict[str, Any]]) -> None:
    master_count, master_total = total_master_demonstrativos(master_path, periodo, empresa)
    sap_count = len(linhas)
    sap_total = round(sum(safe_float(row.get('VALOR')) for row in linhas), 2)
    if master_count != sap_count or abs(master_total - sap_total) > 0.01:
        raise ValueError(
            f'Total SAP divergente do MASTER para {empresa} {periodo}: '
            f'MASTER linhas={master_count} total={master_total:.2f}; '
            f'SAP linhas={sap_count} total={sap_total:.2f}'
        )


def set_formula(ws, row_idx: int, col_letter: str, formula: str) -> None:
    ws[f'{col_letter}{row_idx}'] = formula.format(row=row_idx)


def preencher_parametros_din(ws, row_idx: int, empresa: str, data_procedimento: datetime) -> None:
    data_vencimento = datetime.combine(data_procedimento.date() + timedelta(days=60), datetime.min.time())
    data_emissao = data_procedimento.strftime('%d.%m.%Y')

    parametros_por_empresa = {
        'SBF': {
            'W': 'K',
            'AO': 'OC01',
        },
        'FISIA': {
            'W': 'H',
            'AO': 'FI03',
        },
    }
    parametros = parametros_por_empresa.get(empresa, parametros_por_empresa['SBF'])

    # Bloco lateral da Din. parametrizado, sem copiar a linha inteira do template.
    valores = {
        'L': 'PMP 60DD',
        'O': data_vencimento,
        'R': None,
        'S': None,
        'T': None,
        'W': parametros['W'],
        'X': 'D',
        'Y': None,
        'AA': 1,
        'AB': 'SRV',
        'AC': None,
        'AD': 408,
        'AE': None,
        'AF': None,
        'AG': None,
        'AH': 'D',
        'AI': data_emissao,
        'AK': 500001501,
        'AM': 87942,
        'AO': parametros['AO'],
        'AQ': None,
        'AS': 'WS4047915449',
        'AT': None,
        'AW': None,
        'AZ': 1,
        'BA': 'SRV',
        'BC': 'BRL',
    }
    for col_letter, value in valores.items():
        ws[f'{col_letter}{row_idx}'] = value

    ws[f'O{row_idx}'].number_format = 'd-mmm'

    formulas = {
        'M': '=B{row}',
        'N': '=I{row}',
        'P': '=C{row}',
        'Q': '=A{row}',
        'U': '=IFERROR(IF(AND(Q{row}=Q{prev},M{row}=M{prev}),U{prev},IF(Q{row}=Q{prev},U{prev}+10,10)),"10")',
        'V': '=IF(U{row}=U{prev},"",U{row})',
        'Z': '=H{row}',
        'AJ': '=C{row}',
        'AL': '=AJ{row}',
        'AN': '=A{row}',
        'AP': '=AN{row}',
        'AR': '=AI{row}',
        'AU': '=AP{row}',
        'AV': '=B{row}',
        'AX': '=F{row}',
        'AY': '=G{row}',
        'BB': '=J{row}',
        'BD': '=E{row}',
    }
    for col_letter, formula in formulas.items():
        ws[f'{col_letter}{row_idx}'] = formula.format(row=row_idx, prev=max(row_idx - 1, 2))


def escrever_aba_dinamica(ws, linhas_dinamica: list[list[Any]], empresa: str) -> None:
    max_col = max(ws.max_column, 56)
    data_procedimento = datetime.now()
    limpar_valores(ws, 3, max_col)

    for row_offset, row in enumerate(linhas_dinamica, start=3):
        if row_offset > ws.max_row:
            ws.append([None] * max_col)
        copiar_modelo_linha(ws, 3, row_offset, max_col)

        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_offset, column=col_idx, value=value)

        preencher_parametros_din(ws, row_offset, empresa, data_procedimento)

    last_row = max(3, len(linhas_dinamica) + 2)
    ws.auto_filter.ref = f'L2:AD{last_row}'


def atualizar_pivots_template(ws_din, sheet_name: str, qtd_linhas: int) -> None:
    source_ref = f'A1:T{max(qtd_linhas + 1, 2)}'
    for pivot in getattr(ws_din, '_pivots', []):
        cache = getattr(pivot, 'cache', None)
        cache_source = getattr(cache, 'cacheSource', None)
        worksheet_source = getattr(cache_source, 'worksheetSource', None)
        if worksheet_source is not None:
            worksheet_source.ref = source_ref
            worksheet_source.sheet = sheet_name
        if cache is not None:
            cache.recordCount = qtd_linhas
            cache.refreshOnLoad = True


def gerar_relatorio_sap_empresa(master_path: str, periodo: str, empresa: str) -> str | None:
    empresa_norm = normalizar_empresa(empresa)
    template_path = TEMPLATE_SAP.get(empresa_norm)
    if not template_path or not os.path.exists(template_path):
        logar(f'[AVISO] Template SAP nao encontrado para {empresa_norm}: {template_path}')
        return None

    fornecedor_por_cnpj, cc_fisia_por_loja, _, _ = carregar_bases_sap()
    linhas = montar_linhas_sap(master_path, periodo, empresa_norm, fornecedor_por_cnpj, cc_fisia_por_loja)
    if not linhas:
        return None
    validar_total_relatorio_sap(master_path, periodo, empresa_norm, linhas)

    wb = load_workbook(template_path)
    sheet_name = empresa_norm
    try:
        if sheet_name not in wb.sheetnames:
            logar(f'[AVISO] Template SAP sem aba {sheet_name}: {template_path}')
            return None

        remover_planilha_auxiliar_template(wb)
        ws = wb[sheet_name]
        escrever_linhas_aba_sap(ws, linhas, empresa_norm)

        linhas_dinamica = consolidar_linhas_dinamica(linhas)
        ws_din = wb['Din.']
        escrever_aba_dinamica(ws_din, linhas_dinamica, empresa_norm)
        atualizar_pivots_template(ws_din, sheet_name, len(linhas))

        nome_periodo = periodo_para_nome_arquivo(periodo)
        output_dir = os.path.join(OUTPUT_DIR, empresa_norm, periodo)
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, f'Planilha {nome_periodo} {empresa_norm}.xlsx')
        wb.save(output_path)
        return output_path
    finally:
        wb.close()


def gerar_relatorios_sap_prosegur(master_path: str, periodo: str) -> list[str]:
    output_paths = []
    for empresa in ('SBF', 'FISIA'):
        path = gerar_relatorio_sap_empresa(master_path, periodo, empresa)
        if path:
            output_paths.append(path)
            logar(f'[SAP][{empresa}] Relatorio gerado em: {path}')
    return output_paths


def normalizar_tipo_servico_nome_arquivo(tipo_servico: Any) -> str:
    texto = str(tipo_servico or '').strip().upper()
    texto = re.sub(r'\s+', ' ', texto)

    if texto == 'SERVICOS DE MANUSEIO, ARRUMACAO E CONTAGEM DE VALORES EM TESOURARIA':
        return 'CONTAGEM DE VALORES'

    if texto == 'SERVICOS MUNICIPAIS DE COLETA/ENTREGA VALORES A EMPRESAS':
        return 'TRANSPORTE DE VALORES'

    # fallback defensivo
    if 'MANUSEIO' in texto or 'CONTAGEM' in texto:
        return 'CONTAGEM DE VALORES'

    if 'COLETA/ENTREGA' in texto or 'TRANSPORTE' in texto:
        return 'TRANSPORTE DE VALORES'

    return 'SERVICO'


def sanitize_filename(value: Any) -> str:
    text = str(value or 'DESC').strip().upper()
    text = re.sub(r'[^A-Z0-9._-]+', '_', text)
    text = re.sub(r'_+', '_', text).strip('_')
    return text or 'DESC'


def inicializar_db() -> None:
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    cur.execute(
        '''
        CREATE TABLE IF NOT EXISTS documentos (
            source_filename TEXT PRIMARY KEY,
            raw_payload TEXT NOT NULL,
            document_type TEXT,
            periodo TEXT,
            empresa TEXT,
            numero_documento TEXT,
            updated_at TEXT NOT NULL
        )
        '''
    )

    # Garante a nova coluna execution_id em bases já existentes
    cols = [row[1] for row in cur.execute("PRAGMA table_info(documentos)").fetchall()]
    if 'execution_id' not in cols:
        cur.execute("ALTER TABLE documentos ADD COLUMN execution_id TEXT")

    conn.commit()
    conn.close()


def salvar_no_db(payload: DocumentPayload, execution_id: str) -> None:
    inicializar_db()

    payload_dict = asdict(payload)
    payload_dict['execution_id'] = execution_id
    periodo = get_periodo(payload_dict)
    payload_dict['periodo'] = periodo

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            INSERT OR REPLACE INTO documentos (
                source_filename, raw_payload, document_type, periodo, empresa, numero_documento, updated_at, execution_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ''',
            (
                payload.source_filename,
                json.dumps(payload_dict, ensure_ascii=False, default=str),
                payload.document_type,
                periodo,
                payload.empresa_nome,
                payload.numero_documento,
                datetime.now().isoformat(sep=' '),
                execution_id,
            ),
        )
        conn.commit()
    finally:
        conn.close()


def taxes_to_map(doc: Dict[str, Any]) -> Dict[str, float]:
    result: Dict[str, float] = {}
    for tax in doc.get('taxes') or []:
        tax_type = str(tax.get('tax_type') or '').strip()
        if tax_type:
            result[tax_type] = safe_float(tax.get('amount'))
    return result


def map_notas_ctes(doc: Dict[str, Any]) -> Dict[str, Any]:
    taxes = taxes_to_map(doc)
    is_nfse = doc.get('document_type') == 'NFSE'
    prestador = doc.get('prestador') or {}
    tomador = doc.get('tomador') or {}

    row = {
        'PERIODO': get_periodo(doc),
        'EMPRESA': doc.get('empresa_nome'),
        'CHECK LOJA': '',
        'LOJA': doc.get('numero_loja'),
        'TIPO_DOC': doc.get('document_type'),
        'Nº Nota Fiscal Eletrônica': doc.get('numero_documento') if is_nfse else '',
        'Nº DOCUMENTO': '' if is_nfse else doc.get('numero_documento'),
        'DATA EMISSÃO': doc.get('data_emissao'),
        'COMPETÊNCIA': doc.get('competencia'),
        'VENCIMENTO': doc.get('vencimento'),
        'CNPJ PRESTADOR': prestador.get('tax_id'),
        'CNPJ/CPF TOMADOR': tomador.get('tax_id'),
        'CÓDIGO FISCAL DE OPERAÇÕES E PRESTAÇÕES - NATUREZA DA OPERAÇÃO': doc.get('natureza_operacao'),
        'DESCRIÇÃO DO SERVIÇO PRESTADO': doc.get('descricao_servico'),
        'VALOR TOTAL DA PRESTAÇÃO DO SERVIÇO': safe_float(doc.get('valor_total_documento')) if not is_nfse else 0.0,
        'VALOR DA NOTA (R$)': safe_float(doc.get('valor_total_documento')) if is_nfse else 0.0,
        'VALOR LÍQUIDO (R$)': safe_float(doc.get('valor_liquido')) if is_nfse else safe_float(doc.get('valor_liquido')),
        'VALOR A RECEBER': safe_float(doc.get('valor_a_receber') or doc.get('valor_liquido')),
        'VALOR DO ICMS': safe_float(doc.get('valor_icms')),
        'PIS': taxes.get('PIS', 0.0),
        'COFINS': taxes.get('COFINS', 0.0),
        'IR': taxes.get('IR', 0.0),
        'CSLL': taxes.get('CSLL', 0.0),
        'INSS': taxes.get('INSS', 0.0),
        'ISS': taxes.get('ISS', 0.0),
        'Lote Fatura': doc.get('lote_fatura'),
    }
    row['CHECK LOJA'] = numero_nota_para_check(row)
    return row


def map_boletos(doc: Dict[str, Any]) -> Dict[str, Any]:
    taxes = taxes_to_map(doc)
    return {
        'PERIODO': get_periodo(doc),
        'EMPRESA': doc.get('empresa_nome'),
        'LOJA': doc.get('numero_loja'),
        'TIPO_DOC': doc.get('document_type'),
        'Nº DOCUMENTO': doc.get('numero_documento'),
        'Vencimento': doc.get('vencimento'),
        'Agência/Código Beneficiário': doc.get('agencia_beneficiario'),
        'Nosso Número': doc.get('nosso_numero'),
        'Valor do Documento': safe_float(doc.get('valor_total_documento')),
        'VALOR FATURA': safe_float(doc.get('valor_fatura')),
        'IRRF': taxes.get('IRRF', 0.0),
        'ISSQN': taxes.get('ISSQN', 0.0),
        'INSS': taxes.get('INSS', 0.0),
        'PIS/COF/CSLL': taxes.get('PIS/COF/CSLL', 0.0),
        'Lote Fatura': doc.get('lote_fatura'),
        'Valor Líquido': safe_float(doc.get('valor_liquido')),
        'Código de barras': doc.get('codigo_barras'),
        'CNPJ PRESTADOR': (doc.get('prestador') or {}).get('tax_id'),
        'CNPJ TOMADOR': (doc.get('tomador') or {}).get('tax_id'),
    }


def map_demonstrativos(
    doc: Dict[str, Any],
    item: Optional[Dict[str, Any]] = None,
    exibir_total_pagina: bool = False,
    exibir_total_documento: bool = False,
    loja_fat: Any = None,
) -> Dict[str, Any]:
    tomador = doc.get('tomador') or {}
    row = {
        'PERIODO': get_periodo(doc),
        'EMPRESA': doc.get('empresa_nome'),
        'LOFA FAT': loja_fat or doc.get('numero_loja'),
        'CNPJ PRESTADOR': '',
        'LOJA': doc.get('numero_loja'),
        'Cod Único cliente': doc.get('numero_loja'),
        'Nota Fiscal': doc.get('numero_documento'),
        'Total a Faturar': 0.0,
        'Valor da Nota': safe_float(doc.get('valor_fatura')) if exibir_total_documento else 0.0,
        'Tipo de Serviço': doc.get('tipo_servico'),
        'Modalidade': doc.get('modalidade'),
        'Cliente Faturar': doc.get('cliente_faturar'),
        'Cliente Origem': doc.get('cliente_origem'),
        'Endereço': tomador.get('address'),
        'Página Item': None,
        'Data': None,
        'Guia': None,
        'Coleta / Entrega': None,
        'Chegada / Saída': None,
        'Qtd Milheiro': None,
        'Valor': 0.0,
        'Repasse': 0.0,
        'Montante Total': 0.0,
        'Adv.': 0.0,
        'Tempo Exc.': 0.0,
        'Qtd Malote': None,
        'Malote': 0.0,
        'Cust.': 0.0,
        'Qtd Emb': None,
        'Valor Extr.': 0.0,
        'Valor a faturar': 0.0,
    }

    if item:
        row.update({
            'LOJA': item.get('numero_loja') or doc.get('numero_loja'),
            'Cod Único cliente': item.get('numero_loja') or doc.get('numero_loja'),
            'Tipo de Serviço': item.get('tipo_servico') or doc.get('tipo_servico'),
            'Modalidade': item.get('modalidade') or doc.get('modalidade'),
            'Cliente Faturar': item.get('cliente_faturar') or doc.get('cliente_faturar'),
            'Cliente Origem': item.get('cliente_origem') or doc.get('cliente_origem'),
            'Página Item': item.get('page_number'),
            'Total a Faturar': safe_float(item.get('total_a_faturar_pagina')) if exibir_total_pagina else 0.0,
            'Valor da Nota': safe_float(item.get('valor_nota_documento')) if exibir_total_documento else 0.0,
            'Data': item.get('data'),
            'Guia': item.get('guia'),
            'Coleta / Entrega': item.get('coleta_entrega'),
            'Chegada / Saída': item.get('chegada_saida'),
            'Qtd Milheiro': item.get('qtd_milheiro'),
            'Valor': safe_float(item.get('valor')),
            'Repasse': safe_float(item.get('repasse')),
            'Montante Total': safe_float(item.get('montante_total')),
            'Adv.': safe_float(item.get('advalorem')),
            'Tempo Exc.': safe_float(item.get('tempo_excedente')),
            'Qtd Malote': item.get('malote_qtd'),
            'Malote': safe_float(item.get('malote_valor')),
            'Cust.': safe_float(item.get('custodia')),
            'Qtd Emb': item.get('qtd_emb'),
            'Valor Extr.': safe_float(item.get('valor_extra')),
            'Valor a faturar': safe_float(item.get('valor_a_faturar')),
        })
    return row


def gerar_excels(execution_id: str) -> tuple[str | None, int]:
    periodo_processamento = periodo_atual()
    logar('Consolidando relatório MASTER...')
    # Garante a tabela antes de tentar ler (Prevenção para Cloud Drives)
    inicializar_db()

    # --- CHAMADA DO AGENTE DE AUDITORIA ---
    erros = executar_auditoria(execution_id)
    if erros > 0:
        logar(f"[CONCLUÍDO] Processamento finalizado com {erros} divergências encontradas.")
    else:
        logar("[CONCLUÍDO] Processamento finalizado com sucesso. Tudo conferido!")
    # --------------------------------------

    conn = sqlite3.connect(DB_PATH)
    try:
        df_db = pd.read_sql_query(
            'SELECT source_filename, raw_payload, periodo FROM documentos ORDER BY source_filename',
            conn
        )
    except Exception:
        logar('Erro ao ler a tabela documentos. Certifique-se de que o processamento gerou dados.')
        conn.close()
        return None, erros
    conn.close()

    if df_db.empty:
        logar('Banco vazio. Nada para consolidar.')
        return None, erros

    notas_ctes = []
    boletos = []
    demonstrativos = []
    pendencias = []  # Aba separada para UNKNOWN
    docs_periodo = []

    for _, row in df_db.iterrows():
        doc = json.loads(row['raw_payload'])
        doc['source_filename'] = doc.get('source_filename') or row.get('source_filename')
        doc['periodo'] = periodo_da_linha_db(row, doc)
        if get_periodo(doc) != periodo_processamento:
            continue

        docs_periodo.append(doc)

    loja_por_nota = montar_mapa_loja_por_nota(docs_periodo)

    for doc in docs_periodo:
        doc_type = doc.get('document_type')
        if doc_type in ('NFSE', 'CTE_OS'):
            notas_ctes.append(map_notas_ctes(doc))
        elif doc_type == 'UNKNOWN':
            # UNKNOWN vai para aba própria, não contamina as Notas
            pendencias.append({
                'PERIODO': get_periodo(doc),
                'Arquivo': doc.get('source_filename'),
                'Tipo Detectado': doc_type,
                'Empresa': doc.get('empresa_nome'),
                'Lote': doc.get('lote_fatura'),
                'Observação': 'Classificação não identificada - revisão manual necessária',
            })
        elif doc_type == 'BOL':
            boletos.append(map_boletos(doc))
        elif doc_type == 'DEM':
            loja_fat = loja_por_nota.get(numero_nota_normalizado(doc.get('numero_documento')), doc.get('numero_loja'))
            items = doc.get('dem_items') or []
            if items:
                paginas_ja_escritas = set()
                primeiro_item_documento = True

                for item in items:
                    chave_pagina = (doc.get('source_filename'), item.get('page_number'), item.get('numero_loja'))

                    exibir_total_pagina = chave_pagina not in paginas_ja_escritas
                    exibir_total_documento = primeiro_item_documento

                    demonstrativos.append(
                        map_demonstrativos(
                            doc,
                            item,
                            exibir_total_pagina=exibir_total_pagina,
                            exibir_total_documento=exibir_total_documento,
                            loja_fat=loja_fat,
                        )
                    )

                    paginas_ja_escritas.add(chave_pagina)
                    primeiro_item_documento = False
            else:
                demonstrativos.append(
                    map_demonstrativos(
                        doc,
                        exibir_total_pagina=True,
                        exibir_total_documento=True,
                        loja_fat=loja_fat,
                    )
                )

    aplicar_cnpj_prestador_demonstrativos(
        demonstrativos,
        montar_mapa_cnpj_prestador_por_nota(notas_ctes),
    )

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    master_path = os.path.join(OUTPUT_DIR, 'MASTER_Relatorio_Prosegur.xlsx')
    sheets_rows = {
        'Notas_e_CTEs': notas_ctes,
        'Boletos': boletos,
        'Demonstrativos': demonstrativos,
        'Pendencias_Auditoria': pendencias,
    }

    if not any(sheets_rows.values()):
        logar(f'Nenhum documento do periodo {periodo_processamento} para incluir no MASTER.')
        return master_path if os.path.exists(master_path) else None, erros

    stats = append_rows_to_master(master_path, sheets_rows)
    for sheet_name, sheet_stats in stats.items():
        logar(
            f"[MASTER][{sheet_name}] periodo={periodo_processamento} | "
            f"novas_linhas={sheet_stats['inserted']} | "
            f"ja_existiam={sheet_stats['skipped_existing']}"
        )

    if pendencias:
        logar(f'[AVISO] {len(pendencias)} documentos com classificacao UNKNOWN movidos para aba Pendencias_Auditoria.')

    try:
        gerar_relatorios_sap_prosegur(master_path, periodo_processamento)
    except Exception:
        logar('[AVISO] Falha ao gerar relatorios SAP Prosegur. O MASTER foi preservado.')
        logar(traceback.format_exc())

    logar(f'[OK] Relatório gerado em: {master_path}')
    return master_path, erros

def extrair_valor_analitico(doc: Dict[str, Any]) -> float:
    doc_type = str(doc.get('document_type') or '').strip().upper()

    # Mesmo criterio da aba Notas_e_CTEs do MASTER:
    # CTE_OS entra em "VALOR TOTAL DA PRESTACAO DO SERVICO"
    # NFSE entra em "VALOR DA NOTA (R$)".
    if doc_type in ('NFSE', 'CTE_OS'):
        return safe_float(doc.get('valor_total_documento'))

    return 0.0


def gerar_metricas_vektor(execution_id: str) -> str | None:
    logar('Gerando JSON analítico para o Vektor...')

    inicializar_db()

    conn = sqlite3.connect(DB_PATH)
    try:
        df_db = pd.read_sql_query(
            '''
            SELECT
                source_filename,
                raw_payload,
                document_type,
                periodo,
                empresa,
                numero_documento,
                updated_at,
                execution_id
            FROM documentos
            ORDER BY updated_at
            ''',
            conn
        )
    finally:
        conn.close()

    if df_db.empty:
        logar('Banco vazio. JSON analítico não gerado.')
        return None

    rows = []
    for _, row in df_db.iterrows():
        try:
            doc = json.loads(row['raw_payload'])
        except Exception:
            continue

        periodo = str(row.get('periodo') or doc.get('periodo') or '')
        empresa = str(row.get('empresa') or doc.get('empresa_nome') or 'DESCONHECIDA').strip().upper()
        document_type = str(row.get('document_type') or doc.get('document_type') or 'UNKNOWN').strip().upper()
        execution = str(row.get('execution_id') or '')
        valor = extrair_valor_analitico(doc)

        rows.append({
            'source_filename': str(row.get('source_filename') or ''),
            'periodo': periodo,
            'empresa': empresa,
            'document_type': document_type,
            'numero_documento': str(row.get('numero_documento') or ''),
            'updated_at': str(row.get('updated_at') or ''),
            'execution_id': execution,
            'valor': valor
        })

    def agregar(data_rows):
        by_periodo_empresa = {}
        by_tipo = {}
        cards = {
            'total_docs': 0,
            'total_valor': 0.0,
            'empresas': {},
            'periodos': {}
        }

        for item in data_rows:
            periodo = item['periodo'] or 'SEM_PERIODO'
            empresa = item['empresa'] or 'DESCONHECIDA'
            tipo = item['document_type'] or 'UNKNOWN'
            valor = safe_float(item['valor'])

            cards['total_docs'] += 1
            cards['total_valor'] += valor

            if empresa not in cards['empresas']:
                cards['empresas'][empresa] = {'docs': 0, 'valor': 0.0}
            cards['empresas'][empresa]['docs'] += 1
            cards['empresas'][empresa]['valor'] += valor

            if periodo not in cards['periodos']:
                cards['periodos'][periodo] = {'docs': 0, 'valor': 0.0}
            cards['periodos'][periodo]['docs'] += 1
            cards['periodos'][periodo]['valor'] += valor

            k = f'{periodo}__{empresa}'
            if k not in by_periodo_empresa:
                by_periodo_empresa[k] = {
                    'periodo': periodo,
                    'empresa': empresa,
                    'docs': 0,
                    'valor': 0.0
                }
            by_periodo_empresa[k]['docs'] += 1
            by_periodo_empresa[k]['valor'] += valor

            k2 = f'{periodo}__{empresa}__{tipo}'
            if k2 not in by_tipo:
                by_tipo[k2] = {
                    'periodo': periodo,
                    'empresa': empresa,
                    'document_type': tipo,
                    'docs': 0,
                    'valor': 0.0
                }
            by_tipo[k2]['docs'] += 1
            by_tipo[k2]['valor'] += valor

        return {
            'cards': cards,
            'by_periodo_empresa': list(by_periodo_empresa.values()),
            'by_tipo': list(by_tipo.values())
        }

    current_period = periodo_atual()
    rows_latest = [r for r in rows if r['execution_id'] == execution_id]
    rows_current_period = [r for r in rows if r['periodo'] == current_period]
    resumo_latest = agregar(rows_latest)
    resumo_current_period = agregar(rows_current_period)
    resumo_full = agregar(rows)

    payload = {
        'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'execution_id': execution_id,
        'current_period': current_period,
        'source': {
            'db_path': DB_PATH,
            'log_file': LOG_FILE
        },
        'latest_execution': resumo_latest,
        'current_period_base': resumo_current_period,
        'full_base': resumo_full
    }

    with open(METRICS_FILE, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    logar(f'[OK] JSON analítico gerado em: {METRICS_FILE}')
    return METRICS_FILE


def copiar_e_renomear_arquivo(origem: str, payload_dict: Dict[str, Any]) -> None:
    periodo = get_periodo(payload_dict)
    empresa = sanitize_filename(payload_dict.get('empresa_nome'))
    if payload_dict.get('document_type') == 'BOL':
        subpasta = 'Boletos'
    elif payload_dict.get('document_type') == 'DEM':
        subpasta = 'Demonstrativos'
    else:
        subpasta = 'Notas Fiscais'

    dest_dir = os.path.join(OUTPUT_DIR, empresa, periodo, subpasta)
    os.makedirs(dest_dir, exist_ok=True)

    tipo_servico = sanitize_filename(
        normalizar_tipo_servico_nome_arquivo(payload_dict.get('tipo_servico'))
    )
    loja = sanitize_filename(payload_dict.get('numero_loja') or 'DESC')
    fornecedor = sanitize_filename(payload_dict.get('codigo_forn') or 'DESC')
    numero_doc = sanitize_filename(payload_dict.get('numero_documento') or 'DESC')
    doc_type = payload_dict.get('document_type')

    if doc_type == 'BOL':
        target_name = f'{fornecedor}_{numero_doc}.PDF'
    elif doc_type in ('NFSE', 'CTE_OS'):
        target_name = f'{tipo_servico}_{loja}_{fornecedor}_{numero_doc}.PDF'
    else:
        target_name = f'{tipo_servico}_{loja}_{fornecedor}_{numero_doc}.PDF'

    destino = os.path.join(dest_dir, target_name)
    shutil.copy2(origem, destino)
    os.remove(origem)


def executar_pipeline() -> None:
    start_time = time.time()
    logar('=== ROBO PROSEGUR ===')
    inicializar_db()

    execution_id = datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:8]
    logar(f'[EXECUTION_ID] {execution_id}')

    if not os.path.exists(INPUT_DIR):
        os.makedirs(INPUT_DIR, exist_ok=True)
        logar(f'Pasta de entrada criada em {INPUT_DIR}.')
        return

    arquivos = sorted([file for file in os.listdir(INPUT_DIR) if file.lower().endswith('.pdf')])
    if not arquivos:
        logar('Nenhum PDF encontrado na pasta de entrada.')
        return

    grupos = {}
    for file_name in arquivos:
        parts = file_name.split('_')
        key = f'{parts[0]}_{parts[3]}' if len(parts) >= 4 else f'{parts[0]}_DESC'
        grupos.setdefault(key, []).append(file_name)

    novos = 0
    pulados = 0
    falhas_parse = 0

    # Pré-filtra: separa arquivos novos dos já processados
    arquivos_novos = []
    for file_name in arquivos:
        conn = sqlite3.connect(DB_PATH)
        exists = conn.execute(
            'SELECT 1 FROM documentos WHERE source_filename = ?',
            (file_name,)
        ).fetchone()
        conn.close()
        if exists:
            pulados += 1
        else:
            arquivos_novos.append(file_name)

    total_novos = len(arquivos_novos)

    if pulados > 0:
        logar(f'{pulados} arquivo(s) já processados anteriormente serão ignorados.')
    if total_novos == 0:
        logar('Nenhum arquivo novo para processar nesta leva.')

    for group_key, batch in grupos.items():
        logar(f'Sub-lote: {group_key}')
        batch_payloads = []

        for file_name in batch:
            file_path = os.path.join(INPUT_DIR, file_name)

            # Trava de Duplicidade
            conn = sqlite3.connect(DB_PATH)
            exists = conn.execute(
                'SELECT 1 FROM documentos WHERE source_filename = ?',
                (file_name,)
            ).fetchone()
            conn.close()

            if exists:
                logar(f'[PULADO] Já processado: {file_name}')
                os.remove(file_path)
                continue

            novos += 1
            logar(f'[{novos}/{total_novos}] Processando novo arquivo: {file_name}')
            try:
                payload = analisar_pdf_extrair_dados(file_path, file_name)
                batch_payloads.append(payload)
            except Exception as exc:
                falhas_parse += 1
                logar(f'[ERRO] Falha ao processar {file_name}: {exc}')


        reference = next((item for item in batch_payloads if item.numero_loja and item.numero_loja != 'MULTI'), None)

        for payload in batch_payloads:
            if reference:
                if not payload.numero_loja:
                    payload.numero_loja = reference.numero_loja
                if payload.empresa_nome == 'Desconhecida':
                    payload.empresa_nome = reference.empresa_nome
                if not payload.codigo_forn:
                    payload.codigo_forn = reference.codigo_forn
                if not payload.competencia:
                    payload.competencia = reference.competencia

            try:
                salvar_no_db(payload, execution_id)
                copiar_e_renomear_arquivo(os.path.join(INPUT_DIR, payload.source_filename), asdict(payload))
                logar(f'[OK] Arquivo finalizado: {payload.source_filename}')
            except Exception as exc:
                logar(f'[ERRO] Falha ao salvar/copiar {payload.source_filename}: {exc}')

    master_path, erros_auditoria = gerar_excels(execution_id)
    metrics_path = gerar_metricas_vektor(execution_id)
    elapsed = int(time.time() - start_time)
    logar(
        f'Resumo da execução | novos={novos} | '
        f'já processados/pulados={pulados} | '
        f'falhas_parse={falhas_parse}'
    )
    if metrics_path:
        logar(f'Arquivo analítico Vektor: {metrics_path}')
    logar(f'Finalizado em {elapsed}s.')



if __name__ == '__main__':
    executar_pipeline()
