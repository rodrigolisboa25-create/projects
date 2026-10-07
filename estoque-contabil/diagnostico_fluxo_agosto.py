from pathlib import Path

from openpyxl import load_workbook

from ops_contabil.db import connect
from ops_contabil.ingestion.schema_resolver import load_contracts, resolve_headers
from ops_contabil.settings import load_settings


def main() -> None:
    settings = load_settings()
    with connect(settings.path("database")) as connection:
        queries = {
            "inv lengths": """
                select length(material), length(style_color), count(*)
                  from inventory_rows where period='2026-08'
                 group by 1,2 order by 3 desc limit 20
            """,
            "mb lengths": """
                select length(material), length(plant), count(*)
                  from mb59_rows where period='2026-08'
                 group by 1,2 order by 3 desc limit 20
            """,
            "matched examples": """
                select i.plant,i.material,i.style_color,m.material
                  from inventory_rows i join mb59_rows m
                    on m.period=i.period and m.center_material_key=i.center_material_key
                 where i.period='2026-08' limit 10
            """,
            "unmatched examples": """
                select i.plant,i.material,i.style_color
                  from inventory_rows i left join mb59_rows m
                    on m.period=i.period and m.center_material_key=i.center_material_key
                 where i.period='2026-08' and m.center_material_key is null limit 20
            """,
            "mb examples": """
                select plant,material,center_material_key,posting_date
                  from mb59_rows where period='2026-08' limit 20
            """,
            "inv styles": """
                select count(distinct material),count(distinct style_color),
                       count(distinct substr(material,1,10)),min(material),max(material)
                  from inventory_rows where period='2026-08'
            """,
        }
        for name, query in queries.items():
            print(f"\n{name}")
            print(connection.execute(query).fetchall())

    files = [
        ("MB59", Path(r"C:\Users\75649\AppData\Local\OpsContabil\landing\mb59\2026-08\MB59_2026-08.xlsx"), "mb59_base"),
        ("ZMM", Path(r"C:\Users\75649\AppData\Local\OpsContabil\landing\zmm119\2026-08\ZMM119_2026-08.xlsx"), "zmm119_base_orig"),
        ("AB", Path(r"C:\Users\75649\AppData\Local\OpsContabil\landing\all_brazil\2026-08\ALL_BRAZIL_31_08.xlsx"), "all_brazil_data"),
    ]
    contracts = load_contracts(settings.root / "config" / "schemas.yaml")
    for name, path, contract_name in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            print(f"\n{name}: {path} {workbook.sheetnames}")
            for worksheet in workbook.worksheets:
                preview = [
                    list(row)
                    for row in worksheet.iter_rows(
                        min_row=1,
                        max_row=contracts[contract_name].header_search_rows,
                        values_only=True,
                    )
                ]
                try:
                    resolution = resolve_headers(preview, contracts[contract_name])
                except Exception as exc:
                    print(worksheet.title, "not resolved", exc)
                    continue
                print(
                    worksheet.title,
                    "header",
                    resolution.header_row,
                    "found",
                    resolution.positions,
                    "missing",
                    resolution.optional_missing,
                )
                row = next(
                    worksheet.iter_rows(
                        min_row=resolution.header_row + 1,
                        max_row=resolution.header_row + 1,
                        values_only=True,
                    )
                )
                print("sample mapped", {field: row[pos] for field, pos in resolution.positions.items()})
        finally:
            workbook.close()


if __name__ == "__main__":
    main()
