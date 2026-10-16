"""PASSO A PASSO pela Base GR do mês: SKU, entrada mais recente sem intercompany, Data de lançamento e coleção do pedido
(Base GR ou export da SQVI DOC_COMPRA). Competências antigas continuam na regra por CE & MAT."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from ops_contabil import mb59
from ops_contabil import closing_agenda as agenda
from ops_contabil.db import connect
from ops_contabil.gr_collection import collection_status, import_collection, orders_for_sqvi, read_collection_file
from ops_contabil.mb59 import enrich_inventory_pass_step, ensure_mb59_schema, import_mb59_xlsx
from test_optimus_import import ADMIN, env  # noqa: F401 - fixture reaproveitada

HEADER = ["Material", "Texto breve material", "Centro", "Depósito", "Tipo de movimento", "Doc.material", "Data de lançamento",
          "Qtd.  UM registro", "Nome do usuário", "Cód.débito/crédito", "Data de entrada", "Referência", "Pedido", "Item",
          "Data do documento", "Fornecedor", "Montante em MI", "Quantidade", "SKU", "PO", "Season", "Year"]
# material, centro, doc material, data, pedido, fornecedor, season, year
ENTRIES = [
    ("AAA111-001-M", "1081", "5000000001", date(2026, 9, 5), "4800000001", "8000000815", "FA", 2026),
    ("AAA111-001-L", "1082", "5000000002", date(2026, 9, 20), "4800000002", "8000000781", "HO", 2026),
    ("BBB222-002-M", "1081", "5000000003", date(2026, 9, 10), "4700000003", "152254", "SU", 2026),
    ("BBB222-002-M", "1083", "5000000004", date(2026, 9, 25), "43000004", "7008002", None, None),  # intercompany
    ("CCC333-003-L", "1083", "5000000005", date(2026, 9, 28), "42000005", "7008004", None, None),  # intercompany
]


@pytest.fixture(autouse=True)
def no_all_brazil(monkeypatch):
    monkeypatch.setattr(mb59, "pass_step_lookup_from_all_brazil", lambda *args, **kwargs: {"lookup": {}})


def base_gr(path: Path, season_formula: str | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "BASE"
    ws.append([None] * len(HEADER))  # linha 1 (subtotal no arquivo real)
    ws.append(HEADER)
    for i, (material, plant, doc, posted, order, vendor, season, year) in enumerate(ENTRIES, start=3):
        row = [material, "PRODUTO", plant, "1000", "101", doc, posted, 1, "FF_VENDAS05", "S", posted, "REF", order, "10",
               posted, vendor, 100.0, 1, material[:10], order[:2], season, year]
        if season_formula:
            row[20] = season_formula.format(i=i)
        ws.append(row)
    pe = wb.create_sheet("PE")
    pe.append(["Material", "SKU", "Data de lançamento", "Season", "Year"])
    wb.save(path)
    return path


def sqvi_export(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Doc.compra", "Empr", "C", "Tipo", "Criado em", "Fornecedor", "Ano coleção", "Fim coleção",
               "Inicio da coleção", "Código coleção"])
    ws.append(["4800000001", "7170", "F", "YIMP", "01.06.2026", "8000000815", "2026", "19.12.2026", "20.06.2025", "FA"])
    ws.append(["4800000002", "7170", "F", "YIMP", "01.06.2026", "8000000781", "2026", "19.12.2026", "20.06.2025", "HO"])
    ws.append(["4700000003", "7170", "F", "YL4L", "01.06.2026", "152254", "2026", "19.12.2026", "20.06.2025", "SU"])
    ws.append(["4799999999", "9999", "F", "YL4L", "01.06.2026", "152254", "2025", "19.12.2026", "20.06.2025", "SP"])  # outra empresa
    wb.save(path)
    return path


def seed_position(database) -> None:
    with connect(database) as connection:
        ensure_mb59_schema(connection)
        for period in ("2026-08", "2026-09"):
            for row, (plant, material) in enumerate(
                [("1081", "AAA111-001-S"), ("1081", "BBB222-002-M"), ("1083", "CCC333-003-L"), ("1081", "ZZZ999-009-M")], start=1
            ):
                connection.execute(
                    """insert into inventory_rows(period,source_row,company,plant,material,center_material_key,
                           unrestricted_quantity,fiscal_total_amount) values (?,?,'7170',?,?,?,1,10)""",
                    [period, row, plant, material, plant + material],
                )
        connection.execute("""update inventory_rows set season='SP', year=2025, product_offer_end_date='2025-03-01'
                              where period='2026-08' and material in ('CCC333-003-L','AAA111-001-S')""")


def position(database) -> dict[str, tuple]:
    with connect(database) as connection:
        rows = connection.execute("""select i.material, i.product_offer_end_date, i.season, i.year, s.season
                                     from inventory_rows i join pass_step_sources s on s.period=i.period and s.source_row=i.source_row
                                     where i.period='2026-09'""").fetchall()
    return {r[0]: r[1:] for r in rows}


def test_sku_rule_uses_latest_entry_without_intercompany_and_the_order_collection(env, tmp_path) -> None:  # noqa: F811
    settings, database, *_ = env
    seed_position(database)
    file = base_gr(tmp_path / "5-MB59_BASE GR_ENTRADAS SETEMBRO - SEASON.xlsx")

    loaded = import_mb59_xlsx(settings, "2026-09", file)
    assert loaded["sku_entries"] == 2 and loaded["orders"] == 5 and loaded["intercompany_orders"] == 2
    assert orders_for_sqvi(settings, "2026-09") == ["4700000003", "4800000001", "4800000002"]

    # Sem a coleção: a data já vem da Base GR (primeira entrada do Material); Season/Year seguem para o mês anterior.
    result = enrich_inventory_pass_step(settings, "2026-09", date(2026, 9, 30))
    assert result["pass_step_rule"] == "sku" and result["gr_collection_imported"] is False
    assert result["offer_season_rule"] == "material"
    rows = position(database)
    assert rows["BBB222-002-M"][0] == date(2026, 9, 10)  # primeira entrada do Material no arquivo (mesmo Material, qualquer centro)
    # O tamanho S não teve entrada no mês: nada vem de outro tamanho do mesmo modelo (PROCV pelo Material).
    assert rows["AAA111-001-S"][:3] == (date(2025, 3, 1), "SP", 2025)

    summary = import_collection(settings, "2026-09", file, imported_by=ADMIN)
    assert summary["kind"] == "base_gr" and summary["orders"] == 3 and summary["intercompany_lines"] == 2
    assert summary["coverage"]["pct"] == 100.0
    enrich_inventory_pass_step(settings, "2026-09", date(2026, 9, 30))
    rows = position(database)
    # Tudo pelo Material (como no fechamento): o tamanho S não teve entrada no mês, então fica com o mês anterior.
    assert rows["AAA111-001-S"] == (date(2025, 3, 1), "SP", 2025, "Mês anterior (2026-08)")
    # A entrada intercompany mais nova (25/09) é ignorada: vale a de 10/09 do pedido 4700000003. Esse pedido é
    # nacional (47…), então o Year não vem dele (só pedidos de importação) e, sem mês anterior nem All Brazil, fica vazio.
    assert rows["BBB222-002-M"] == (date(2026, 9, 10), "SU", None, "MB59")
    # Só entrada intercompany no mês: vai para o mês anterior (por CE & MAT).
    assert rows["CCC333-003-L"][1:3] == ("SP", 2025) and rows["CCC333-003-L"][3].startswith("Mês anterior")
    assert rows["ZZZ999-009-M"][3] == "Não localizado"
    item = next(i for i in agenda.checklist(settings, "2026-09") if i["key"] == "gr_collection")
    assert item["done"] is True and "3 pedidos" in item["detail"]


def test_sqvi_export_is_accepted_and_other_company_ignored(env, tmp_path) -> None:  # noqa: F811
    settings, database, *_ = env
    seed_position(database)
    import_mb59_xlsx(settings, "2026-09", base_gr(tmp_path / "base.xlsx"))
    parsed = read_collection_file(sqvi_export(tmp_path / "Suporte SQVI.xlsx"), "2026-09")
    assert parsed["kind"] == "sqvi" and parsed["orders"] == 3 and parsed["other_company_lines"] == 1
    import_collection(settings, "2026-09", tmp_path / "Suporte SQVI.xlsx")
    status = collection_status(settings, "2026-09")
    assert status["imported"] and status["source_kind"] == "sqvi" and status["coverage"]["with_collection"] == 3


def test_base_gr_with_season_from_sku_is_refused(env, tmp_path) -> None:  # noqa: F811
    settings, *_ = env
    file = base_gr(tmp_path / "agosto.xlsx", season_formula="=VLOOKUP(S{i},PE!B:E,3,0)")
    with pytest.raises(ValueError, match="buscada pelo SKU na aba PE"):
        read_collection_file(file, "2026-09")
    other_month = base_gr(tmp_path / "setembro.xlsx")
    with pytest.raises(ValueError, match="nenhum da competência 2026-10"):
        read_collection_file(other_month, "2026-10")


def test_routes_status_orders_and_upload_detection(env, tmp_path) -> None:  # noqa: F811
    settings, database, _downloads, client_for, *_ = env
    seed_position(database)
    import_mb59_xlsx(settings, "2026-09", base_gr(tmp_path / "base.xlsx"))
    client, csrf = client_for(ADMIN)
    status = client.get("/api/sap/gr-collection/status", params={"period": "2026-09"}).json()
    assert status["imported"] is False and status["mb59_ready"] is True and status["coverage"]["mb59_orders"] == 3
    orders = client.get("/api/sap/gr-collection/orders", params={"period": "2026-09"}).json()
    assert orders["count"] == 3 and orders["text"].split("\r\n") == ["4700000003", "4800000001", "4800000002"]
    # Upload: arquivo que não é Base GR nem SQVI é recusado na hora, sem criar execução.
    wb = Workbook()
    wb.active.append(["Qualquer", "Coisa"])
    wb.save(tmp_path / "outro.xlsx")
    response = client.post("/api/sap/upload", params={"transaction": "BASEGR", "period": "2026-09", "file_name": "outro.xlsx"},
                           headers={"X-CSRF-Token": csrf, "Content-Type": "application/octet-stream"},
                           content=(tmp_path / "outro.xlsx").read_bytes())
    assert response.status_code == 422 and "nem o export da SQVI" in response.json()["detail"]


def test_month_status_feeds_the_three_steps_of_the_sap_page(env, tmp_path) -> None:  # noqa: F811
    settings, database, _downloads, client_for, *_ = env
    seed_position(database)
    file = base_gr(tmp_path / "base.xlsx")
    import_mb59_xlsx(settings, "2026-09", file)
    client, _csrf = client_for(ADMIN)
    status = client.get("/api/sap/month-status", params={"period": "2026-09"}).json()
    assert status["mb59"] and status["gr"]["imported"] is False
    import_collection(settings, "2026-09", file)
    status = client.get("/api/sap/month-status", params={"period": "2026-09"}).json()
    assert status["gr"]["imported"] is True and status["gr"]["pct"] == 100
    assert client.get("/api/sap/month-status", params={"period": "09-2026"}).status_code == 422
    # A página mostra as três etapas, um quadro por vez, e mantém Execuções atuais.
    html = (Path(__file__).resolve().parents[1] / "src" / "ops_contabil" / "web" / "templates" / "dashboard.html").read_text(encoding="utf-8")
    assert [f'data-sap-tab="{k}"' in html for k in ("z", "m", "gr")] == [True] * 3
    assert all(f'data-sap-panel="{k}"' in html for k in ("z", "m", "gr")) and 'class="card jobs"' in html


def test_governance_control_shows_the_collection_coverage(env, tmp_path) -> None:  # noqa: F811
    from ops_contabil.health_center import governance_status

    settings, database, *_ = env
    seed_position(database)
    file = base_gr(tmp_path / "base.xlsx")
    import_mb59_xlsx(settings, "2026-09", file)

    def control():
        return next(c for c in governance_status(settings, "2026-09")["controls"]
                    if c["title"] == "Season e Year dos pedidos (Base GR / SQVI)")

    assert control()["badge"] == "Não carregada" and "3 pedidos" in control()["detail"]
    import_collection(settings, "2026-09", file)
    assert control()["level"] == "healthy" and control()["badge"] == "100,0% dos pedidos"
    assert "3 de 3 pedidos da MB59 com coleção" in control()["detail"]


def test_old_month_mb59_is_prepared_when_the_collection_arrives(env, tmp_path) -> None:  # noqa: F811
    settings, database, *_ = env
    seed_position(database)
    file = base_gr(tmp_path / "julho_mb59.xlsx")
    import_mb59_xlsx(settings, "2026-09", file)
    with connect(database) as connection:  # simula um mês carregado antes da regra por SKU
        connection.execute("delete from mb59_sku_entries where period='2026-09'")
        connection.execute("delete from mb59_orders where period='2026-09'")
    # O botão "Copiar pedidos da MB59" já funciona no mês antigo (lê a MB59 salva, sem gravar nada).
    assert orders_for_sqvi(settings, "2026-09") == ["4700000003", "4800000001", "4800000002"]
    summary = import_collection(settings, "2026-09", file)
    assert summary["mb59_prepared"] is True and summary["coverage"]["with_collection"] == 3
    enrich_inventory_pass_step(settings, "2026-09", date(2026, 9, 30))
    assert position(database)["AAA111-001-S"] == (date(2025, 3, 1), "SP", 2025, "Mês anterior (2026-08)")  # tudo pelo Material


def test_old_periods_keep_the_ce_mat_rule(env) -> None:  # noqa: F811
    settings, database, *_ = env
    seed_position(database)
    with connect(database) as connection:  # MB59 carregada antes desta regra: só mb59_rows, sem entradas por SKU
        connection.execute("""insert into mb59_rows(period,center_material_key,material,plant,product_offer_end_date,season,year,
                              season_year,posting_date,source_row) values ('2026-09','1081BBB222-002-M','BBB222-002-M','1081',
                              '2026-12-31','HO',2027,'HO2027','2026-09-10',1)""")
    result = enrich_inventory_pass_step(settings, "2026-09", date(2026, 9, 30))
    assert result["pass_step_rule"] == "ce_mat"
    assert position(database)["BBB222-002-M"] == (date(2026, 12, 31), "HO", 2027, "MB59")
