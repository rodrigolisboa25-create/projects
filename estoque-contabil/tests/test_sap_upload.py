"""Carga manual (upload) das planilhas ZMM119 e MB59 pela página Extrações SAP."""

from __future__ import annotations

import shutil
import time
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import ops_contabil.web.app as web_app
from ops_contabil.auth import SESSION_COOKIE, create_session_token
from ops_contabil.db import connect
from ops_contabil.ingestion.schema_resolver import load_contracts
from ops_contabil.inventory import ZMM119_POSITIONAL_CONTRACT, ensure_inventory_schema
from ops_contabil.settings import Settings

PROJECT_ROOT = Path(__file__).resolve().parents[1]
ADMIN = "upload.admin@gruposbf.com.br"
VIEWER = "upload.viewer@gruposbf.com.br"


@pytest.fixture()
def env(tmp_path: Path, monkeypatch):
    (tmp_path / "config").mkdir()
    shutil.copy(PROJECT_ROOT / "config" / "schemas.yaml", tmp_path / "config" / "schemas.yaml")
    database = tmp_path / "processed" / "ops.duckdb"
    settings = Settings(
        root=tmp_path,
        raw={
            "project": {
                "database": str(database),
                "processed": str(database.parent),
                "landing": str(tmp_path / "landing"),
                "output": str(tmp_path / "output"),
            },
            "sources": {"all_brazil_materials": {}},
            "runtime": {"sap_timeout_seconds": 300},
        },
    )
    with connect(database) as connection:
        ensure_inventory_schema(connection)
        for email, role, pages in ((ADMIN, "admin", "[]"), (VIEWER, "user", '["sap"]')):
            connection.execute(
                "insert into allowed_users(email,is_active,role,allowed_pages,identity_subject) values (?,true,?,?,?)",
                [email, role, pages, f"windows:{email}"],
            )
    # Enriquecimentos dependem do All Brazil no Drive; aqui o foco é a carga.
    enrich_calls: list[tuple[str, str]] = []
    monkeypatch.setattr(
        web_app, "enrich_inventory_from_all_brazil",
        lambda _s, period, _d: enrich_calls.append(("all_brazil", period)) or {"coverage_pct": 99.5},
    )
    monkeypatch.setattr(
        web_app, "enrich_inventory_pass_step",
        lambda _s, period, _d: enrich_calls.append(("pass_step", period)) or {"status": "enriched", "message": "PASSO A PASSO ok."},
    )
    monkeypatch.setattr(web_app, "apply_mapping_rules", lambda _s, period, _d: enrich_calls.append(("mapping", period)))
    web_app.SAP_JOBS.clear()
    app = web_app.create_app(settings)

    def client_for(email: str) -> tuple[TestClient, str]:
        client = TestClient(app)
        client.cookies.set(SESSION_COOKIE, create_session_token(settings, email, f"windows:{email}"))
        return client, client.get("/api/auth/session").json()["csrf_token"]

    return settings, database, client_for, enrich_calls


def zmm119_workbook(path: Path, rows: list[tuple[str, str, str, float]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([label for _, label in ZMM119_POSITIONAL_CONTRACT])
    for company, plant, material, quantity in rows:
        ws.append([company, plant, material, "TENIS TESTE", "PAR", "64041100", quantity, 0, 10.0, 12.0,
                   quantity * 12.0, 0, 0, "N", quantity * 10.0, quantity * 12.0, 11.0, None][: len(ZMM119_POSITIONAL_CONTRACT)])
    wb.save(path)
    return path


def mb59_workbook(path: Path, settings: Settings, posting: date, rows: int = 3) -> Path:
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["mb59_base"]
    header = [column.aliases[0] if column.aliases else column.canonical for column in contract.columns]
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for index in range(rows):
        record = []
        for column in contract.columns:
            if column.canonical == "material":
                record.append(f"MAT{index:09d}")
            elif column.canonical == "plant":
                record.append("1081")
            elif column.data_type == "date":
                record.append(posting)
            elif column.data_type == "decimal":
                record.append(1.0)
            else:
                record.append(f"{column.canonical}-{index}")
        ws.append(record)
    wb.save(path)
    return path


def upload(client: TestClient, csrf: str, transaction: str, period: str, path: Path, name: str | None = None, **extra):
    params = {"transaction": transaction, "period": period, "file_name": name or path.name, **extra}
    return client.post(
        "/api/sap/upload",
        params=params,
        content=path.read_bytes(),
        headers={"X-CSRF-Token": csrf, "Content-Type": "application/octet-stream"},
    )


def wait_job(client: TestClient, job_id: str, timeout: float = 60) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/sap/jobs/{job_id}").json()
        if job["status"] in {"COMPLETED", "FAILED", "CANCELLED", "ACTION_REQUIRED"}:
            return job
        time.sleep(0.1)
    raise AssertionError(f"Job {job_id} não terminou: {job}")


def inventory_count(database: Path, period: str) -> int:
    with connect(database) as connection:
        return connection.execute("select count(*) from inventory_rows where period=?", [period]).fetchone()[0]


def test_zmm119_upload_imports_through_the_same_pipeline(env, tmp_path: Path) -> None:
    settings, database, client_for, enrich_calls = env
    client, csrf = client_for(ADMIN)
    source = zmm119_workbook(tmp_path / "export_sap.xlsx", [
        ("7170", "1081", "MAT000000001", 10), ("7170", "1080", "MAT000000002", 5), ("8000", "1081", "MAT000000003", 7),
    ])

    response = upload(client, csrf, "ZMM119", "2026-09", source, as_of="2026-09-30")
    assert response.status_code == 200, response.text
    assert response.json()["source"] == "upload"
    job = wait_job(client, response.json()["job_id"])

    assert job["status"] == "COMPLETED", job
    assert "ZMM119 carregada por upload (2 linhas)" in job["message"]
    assert inventory_count(database, "2026-09") == 2  # somente empresa 7170
    target = settings.path("landing") / "zmm119" / "2026-09" / "ZMM119_2026-09.xlsx"
    assert target.is_file()
    assert ("all_brazil", "2026-09") in enrich_calls and ("mapping", "2026-09") in enrich_calls
    assert not list((settings.path("landing") / "uploads").iterdir())


def test_mb59_upload_is_indexed(env, tmp_path: Path) -> None:
    settings, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    source = mb59_workbook(tmp_path / "mb59.xlsx", settings, date(2026, 9, 15), rows=4)

    job = wait_job(client, upload(client, csrf, "MB59", "2026-09", source).json()["job_id"])

    assert job["status"] == "COMPLETED", job
    assert "MB59 carregada por upload e indexada (4 chaves CE & MAT)" in job["message"]
    with connect(database) as connection:
        assert connection.execute("select count(*) from mb59_rows where period='2026-09'").fetchone()[0] == 4


def test_wrong_file_is_rejected_without_touching_existing_data(env, tmp_path: Path) -> None:
    settings, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    good = zmm119_workbook(tmp_path / "good.xlsx", [("7170", "1081", "MAT000000001", 10)])
    assert wait_job(client, upload(client, csrf, "ZMM119", "2026-09", good).json()["job_id"])["status"] == "COMPLETED"
    target = settings.path("landing") / "zmm119" / "2026-09" / "ZMM119_2026-09.xlsx"
    original_bytes = target.read_bytes()

    # MB59 enviada no lugar da ZMM119: recusada já no envio, sem criar execução.
    wrong = mb59_workbook(tmp_path / "mb59.xlsx", settings, date(2026, 9, 15))
    response = upload(client, csrf, "ZMM119", "2026-09", wrong)

    assert response.status_code == 422
    assert "Esta planilha é da MB59, não da ZMM119" in response.json()["detail"]
    assert inventory_count(database, "2026-09") == 1
    assert target.read_bytes() == original_bytes
    assert all(job.get("file_name") != "mb59.xlsx" for job in client.get("/api/sap/jobs").json())


def test_zmm119_sent_as_mb59_is_rejected_on_upload(env, tmp_path: Path) -> None:
    settings, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    zmm = zmm119_workbook(tmp_path / "ZMM119_setembro.xlsx", [("7170", "1081", "MAT000000001", 10)])

    response = upload(client, csrf, "MB59", "2026-09", zmm)

    assert response.status_code == 422
    assert "Esta planilha é da ZMM119, não da MB59" in response.json()["detail"]
    uploads = settings.path("landing") / "uploads"
    assert not uploads.exists() or not list(uploads.iterdir())
    assert not (settings.path("landing") / "mb59").exists()


@pytest.mark.parametrize("transaction", ["ZMM119", "MB59"])
def test_unrelated_spreadsheet_is_rejected_for_both_transactions(env, tmp_path: Path, transaction: str) -> None:
    _, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    wb = Workbook()
    ws = wb.active
    ws.append(["Fornecedor", "CNPJ", "Valor", "Vencimento"])
    ws.append(["ACME", "00.000.000/0001-00", 1500.0, date(2026, 9, 30)])
    other = tmp_path / "contas_a_pagar.xlsx"
    wb.save(other)

    response = upload(client, csrf, transaction, "2026-09", other)

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "não é uma exportação da" in detail and "ZMM119" in detail and "MB59" in detail
    assert inventory_count(database, "2026-09") == 0


def test_mb59_from_another_month_is_rejected(env, tmp_path: Path) -> None:
    settings, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    august = mb59_workbook(tmp_path / "mb59_agosto.xlsx", settings, date(2026, 8, 20))

    job = wait_job(client, upload(client, csrf, "MB59", "2026-09", august).json()["job_id"])

    assert job["status"] == "FAILED"
    assert "2026-08" in job["message"] and "2026-09" in job["message"]
    with connect(database) as connection:
        # Recusada antes de qualquer gravação: a tabela da MB59 nem chegou a ser criada.
        exists = connection.execute(
            "select count(*) from information_schema.tables where table_name='mb59_rows'"
        ).fetchone()[0]
        assert not exists or connection.execute("select count(*) from mb59_rows").fetchone()[0] == 0


def test_zmm119_without_company_7170_is_rejected(env, tmp_path: Path) -> None:
    _, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    other = zmm119_workbook(tmp_path / "outra_empresa.xlsx", [("8000", "1081", "MAT000000001", 10)])

    job = wait_job(client, upload(client, csrf, "ZMM119", "2026-09", other).json()["job_id"])

    assert job["status"] == "FAILED"
    assert "7170" in job["message"]
    assert inventory_count(database, "2026-09") == 0


def test_invalid_uploads_are_refused_immediately(env, tmp_path: Path) -> None:
    settings, _, client_for, _ = env
    client, csrf = client_for(ADMIN)
    not_xlsx = tmp_path / "planilha.xlsx"
    not_xlsx.write_text("isto não é um xlsx", encoding="utf-8")
    empty = tmp_path / "vazio.xlsx"
    empty.write_bytes(b"")
    good = zmm119_workbook(tmp_path / "good.xlsx", [("7170", "1081", "MAT000000001", 10)])

    assert upload(client, csrf, "ZMM119", "2026-09", not_xlsx).status_code == 422
    assert upload(client, csrf, "ZMM119", "2026-09", empty).status_code == 422
    assert upload(client, csrf, "ZMM119", "2026-09", good, name="export.xls").status_code == 422
    assert upload(client, csrf, "ME2N", "2026-09", good).status_code == 422
    assert upload(client, csrf, "ZMM119", "2026-13x", good).status_code == 422
    uploads = settings.path("landing") / "uploads"
    assert not uploads.exists() or not list(uploads.iterdir())


def test_upload_is_admin_only_and_requires_csrf(env, tmp_path: Path) -> None:
    _, database, client_for, _ = env
    good = zmm119_workbook(tmp_path / "good.xlsx", [("7170", "1081", "MAT000000001", 10)])
    viewer, viewer_csrf = client_for(VIEWER)
    assert upload(viewer, viewer_csrf, "ZMM119", "2026-09", good).status_code == 403

    admin, _ = client_for(ADMIN)
    assert upload(admin, "token-invalido", "ZMM119", "2026-09", good).status_code == 403
    assert inventory_count(database, "2026-09") == 0


def test_failed_import_restores_previous_landing_file(env, tmp_path: Path, monkeypatch) -> None:
    settings, database, client_for, _ = env
    client, csrf = client_for(ADMIN)
    good = zmm119_workbook(tmp_path / "good.xlsx", [("7170", "1081", "MAT000000001", 10)])
    assert wait_job(client, upload(client, csrf, "ZMM119", "2026-09", good).json()["job_id"])["status"] == "COMPLETED"
    target = settings.path("landing") / "zmm119" / "2026-09" / "ZMM119_2026-09.xlsx"
    original_bytes = target.read_bytes()

    def broken_import(*_args, **_kwargs):
        raise ValueError("Falha simulada durante a gravação.")

    monkeypatch.setattr(web_app, "import_zmm119_xlsx", broken_import)
    newer = zmm119_workbook(tmp_path / "newer.xlsx", [("7170", "1081", "MAT000000009", 99)])
    job = wait_job(client, upload(client, csrf, "ZMM119", "2026-09", newer).json()["job_id"])

    assert job["status"] == "FAILED"
    assert "Falha simulada" in job["message"]
    assert target.read_bytes() == original_bytes
    assert inventory_count(database, "2026-09") == 1
