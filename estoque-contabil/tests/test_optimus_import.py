"""Optimus importando planilhas ZMM119/MB59 já baixadas do SAP (contingência), sempre com confirmação."""

from __future__ import annotations

import json
import os
import re
import shutil
import time
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import ops_contabil.sap_upload as sap_upload
import ops_contabil.web.app as web_app
from ops_contabil.auth import SESSION_COOKIE, create_session_token
from ops_contabil.db import connect
from ops_contabil.ingestion.schema_resolver import load_contracts
from ops_contabil.inventory import ZMM119_POSITIONAL_CONTRACT, ensure_inventory_schema
from ops_contabil.settings import Settings

PROJECT_ROOT = Path(__file__).resolve().parents[1]
ADMIN = "optimus.import.admin@gruposbf.com.br"
USER = "optimus.import.user@gruposbf.com.br"


def zmm119_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append([label for _, label in ZMM119_POSITIONAL_CONTRACT])
    for plant, material, quantity in (("1081", "MAT000000001", 10), ("1080", "MAT000000002", 5)):
        ws.append(["7170", plant, material, "TENIS TESTE", "PAR", "64041100", quantity, 0, 10.0, 12.0,
                   quantity * 12.0, 0, 0, "N", quantity * 10.0, quantity * 12.0, 11.0, None][: len(ZMM119_POSITIONAL_CONTRACT)])
    wb.save(path)
    return path


@pytest.fixture()
def env(tmp_path: Path, monkeypatch):
    (tmp_path / "config").mkdir()
    shutil.copy(PROJECT_ROOT / "config" / "schemas.yaml", tmp_path / "config" / "schemas.yaml")
    database = tmp_path / "processed" / "ops.duckdb"
    settings = Settings(
        root=tmp_path,
        raw={
            "project": {"database": str(database), "processed": str(database.parent),
                        "landing": str(tmp_path / "landing"), "output": str(tmp_path / "output")},
            "sources": {"all_brazil_materials": {}},
            "runtime": {"sap_timeout_seconds": 300},
        },
    )
    with connect(database) as connection:
        ensure_inventory_schema(connection)
        for email, role, pages in ((ADMIN, "admin", "[]"), (USER, "user", '["overview","agent"]')):
            connection.execute(
                "insert into allowed_users(email,is_active,role,allowed_pages,identity_subject) values (?,true,?,?,?)",
                [email, role, pages, f"windows:{email}"],
            )
    downloads = tmp_path / "Downloads"
    downloads.mkdir()
    # A busca fica restrita a uma pasta de teste (nunca vasculha o computador real).
    monkeypatch.setattr(sap_upload, "export_search_folders", lambda: [downloads])
    monkeypatch.setattr(web_app, "export_search_folders", lambda: [downloads])
    monkeypatch.setattr(web_app, "enrich_inventory_from_all_brazil", lambda _s, _p, _d: {"coverage_pct": 100})
    monkeypatch.setattr(web_app, "enrich_inventory_pass_step", lambda _s, _p, _d: {"status": "enriched", "message": "ok"})
    monkeypatch.setattr(web_app, "apply_mapping_rules", lambda *args, **kwargs: None)
    web_app.SAP_JOBS.clear()
    web_app.SAP_CONFIRMATIONS.clear()

    sent: list[dict] = []
    replies: list = []

    class FakeClient:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def post(self, url, json=None, headers=None):
            sent.append(json)
            reply = replies.pop(0) if replies else "ok"
            text = reply(json) if callable(reply) else reply

            class Response:
                status_code = 200

                @staticmethod
                def json():
                    return {"output": text}

            return Response()

    monkeypatch.setenv("N8N_CHAT_URL", "http://n8n.invalid/webhook")
    monkeypatch.setattr(web_app.httpx, "AsyncClient", FakeClient)
    app = web_app.create_app(settings)

    def client_for(email: str) -> tuple[TestClient, str]:
        client = TestClient(app)
        client.cookies.set(SESSION_COOKIE, create_session_token(settings, email, f"windows:{email}"))
        return client, client.get("/api/auth/session").json()["csrf_token"]

    return settings, database, downloads, client_for, sent, replies


def action(payload: dict) -> str:
    return f"[[OPS_ACTION]]\n{json.dumps(payload)}\n[[/OPS_ACTION]]"


def tool_result(message: dict) -> dict:
    return json.loads(re.search(r"\[RESULTADO_FERRAMENTA_LOCAL\]\n(.*)\n\[/RESULTADO_FERRAMENTA_LOCAL\]", message["chatInput"], re.S).group(1))


def ask(client: TestClient, csrf: str, message: str) -> dict:
    response = client.post(
        "/api/optimus/chat",
        json={"message": message, "session_id": "sessao-importacao", "period": "2026-09"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200, response.text
    return response.json()


def wait_job(client: TestClient, job_id: str, timeout: float = 60) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/sap/jobs/{job_id}").json()
        if job["status"] in {"COMPLETED", "FAILED", "CANCELLED", "ACTION_REQUIRED"}:
            return job
        time.sleep(0.1)
    raise AssertionError(f"Job {job_id} não terminou")


def mb59_workbook(path: Path, settings: Settings) -> Path:
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["mb59_base"]
    wb = Workbook()
    ws = wb.active
    ws.append([column.aliases[0] if column.aliases else column.canonical for column in contract.columns])
    ws.append([
        "MAT000000001" if column.canonical == "material" else "1081" if column.canonical == "plant"
        else date(2026, 9, 15) if column.data_type == "date" else 1.0 if column.data_type == "decimal"
        else f"{column.canonical}-1"
        for column in contract.columns
    ])
    wb.save(path)
    return path


def test_candidates_recognize_sap_exports_and_ignore_other_files(env) -> None:
    settings, _db, downloads, *_ = env
    zmm119_workbook(downloads / "export_zmm119.xlsx")
    mb59 = mb59_workbook(downloads / "export_mb59.xlsx", settings)
    # A leitura rápida (só cabeçalho) reconhece igual à leitura completa da planilha.
    assert sap_upload.quick_identify_sap_export(settings, mb59) == sap_upload.identify_sap_export(settings, mb59) == "MB59"
    assert sap_upload.quick_identify_sap_export(settings, downloads / "export_zmm119.xlsx") == "ZMM119"
    other = Workbook()
    other.active.append(["qualquer", "coisa"])
    other.save(downloads / "orcamento.xlsx")
    (downloads / "~$aberto.xlsx").write_bytes(b"lock")
    old = zmm119_workbook(downloads / "antiga.xlsx")
    stamp = time.time() - 60 * 86400
    os.utime(old, (stamp, stamp))

    found = {item["file_name"]: item["transaction"] for item in sap_upload.find_sap_export_candidates(settings)}
    assert found == {"export_zmm119.xlsx": "ZMM119", "export_mb59.xlsx": "MB59", "orcamento.xlsx": None}


def test_optimus_lists_prepares_and_imports_only_after_confirmation(env) -> None:
    settings, database, downloads, client_for, sent, replies = env
    zmm119_workbook(downloads / "export_zmm119.xlsx")
    client, csrf = client_for(ADMIN)

    replies[:] = [action({"type": "find_sap_exports"}), "Encontrei a export_zmm119.xlsx."]
    ask(client, csrf, "importe a planilha da ZMM119 que baixei")
    listing = tool_result(sent[-1])
    assert listing["ok"] and [item["file_name"] for item in listing["files"]] == ["export_zmm119.xlsx"]
    assert listing["files"][0]["transaction"] == "ZMM119"

    replies[:] = [action({"type": "import_sap_file", "file_name": "export_zmm119.xlsx", "period": "2026-09"}), "Confirma?"]
    ask(client, csrf, "pode importar para setembro")
    prepared = tool_result(sent[-1])
    assert prepared["status"] == "CONFIRMATION_REQUIRED"
    assert "export_zmm119.xlsx" in prepared["summary"] and "ZMM119 da competência 2026-09" in prepared["summary"]
    with connect(database) as connection:
        assert connection.execute("select count(*) from inventory_rows").fetchone()[0] == 0  # nada importado ainda

    confirmation_id = prepared["confirmation_id"]
    replies[:] = [action({"type": "confirm_sap", "confirmation_id": confirmation_id}), "Importação iniciada."]
    ask(client, csrf, "sim, confirma")
    started = tool_result(sent[-1])
    assert started["ok"] and started["status"] == "STARTED"
    job = wait_job(client, started["result"]["job_id"])
    assert job["status"] == "COMPLETED", job
    assert job["source"] == "upload"
    with connect(database) as connection:
        assert connection.execute("select count(*) from inventory_rows where period='2026-09'").fetchone()[0] == 2
    assert (downloads / "export_zmm119.xlsx").is_file()  # o arquivo original do usuário não é movido nem apagado


def test_file_changed_after_summary_is_not_imported(env) -> None:
    settings, database, downloads, client_for, sent, replies = env
    source = zmm119_workbook(downloads / "export_zmm119.xlsx")
    client, csrf = client_for(ADMIN)
    replies[:] = [action({"type": "import_sap_file", "file_name": "export_zmm119.xlsx", "period": "2026-09"}), "Confirma?"]
    ask(client, csrf, "importe a export_zmm119.xlsx")
    confirmation_id = tool_result(sent[-1])["confirmation_id"]

    zmm119_workbook(source)  # regravada depois do resumo
    stamp = time.time() + 5
    os.utime(source, (stamp, stamp))
    replies[:] = [action({"type": "confirm_sap", "confirmation_id": confirmation_id}), "ok"]
    ask(client, csrf, "sim")
    assert tool_result(sent[-1])["status"] == "FILE_CHANGED"
    with connect(database) as connection:
        assert connection.execute("select count(*) from inventory_rows").fetchone()[0] == 0


def test_failed_extraction_makes_optimus_offer_the_downloaded_file(env, monkeypatch) -> None:
    settings, database, downloads, client_for, sent, replies = env

    class FailingRobot:
        def __init__(self, *args, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def run_zmm119(self, **kwargs):
            # O SAP chegou a salvar a exportação, mas o robô falhou depois.
            zmm119_workbook(downloads / "EXPORT_20261005_101500.XLSX")
            raise web_app.SapGuiError("O SAP não confirmou a gravação no destino oficial.")

    monkeypatch.setattr(web_app, "SapGuiRobot", FailingRobot)
    web_app.OPTIMUS_OFFERS.clear()
    client, csrf = client_for(ADMIN)
    run = client.post("/api/sap/run", json={"transaction": "ZMM119", "period": "2026-09", "as_of": "2026-09-30"},
                      headers={"X-CSRF-Token": csrf})
    assert run.status_code == 200, run.text
    assert wait_job(client, run.json()["job_id"])["status"] == "FAILED"

    # O chat busca os eventos proativos: o AGENTE escreve a mensagem a partir dos fatos (sem texto pronto).
    def agent_writes(message: dict) -> str:
        event = json.loads(re.search(r"\[EVENTO_PROATIVO\]\n(.*)\n\[/EVENTO_PROATIVO\]", message["chatInput"], re.S).group(1))
        assert event["type"] == "sap_extraction_failed" and event["transaction"] == "ZMM119"
        assert event["saved_export"]["file_name"] == "EXPORT_20261005_101500.XLSX"
        assert "[PERGUNTA_USUARIO]" not in message["chatInput"]
        return "Texto escrito pelo agente: a extração falhou, mas posso importar a planilha salva. Confirma?"

    replies[:] = [agent_writes]
    deadline = time.monotonic() + 30
    offer: list = []
    while not offer and time.monotonic() < deadline:  # a oferta sai logo depois do aviso de falha
        offer = client.get("/api/optimus/proactive", params={"session_id": "sessao-importacao"}).json()["messages"]
        time.sleep(0.2)
    assert len(offer) == 1 and offer[0]["text"].startswith("Texto escrito pelo agente")
    assert offer[0]["notify"] is False  # a notificação da falha já foi enviada pelo servidor
    assert client.get("/api/optimus/proactive", params={"session_id": "sessao-importacao"}).json()["messages"] == []

    def confirm_from_context(message: dict) -> str:
        context = json.loads(re.search(r"\[CONTEXTO_OPS_CONTABIL\]\n(.*)\n\[/CONTEXTO_OPS_CONTABIL\]", message["chatInput"], re.S).group(1))
        pending = context["pending_sap_confirmation"]
        assert pending["origin"] == "optimus_proactive" and pending["parameters"]["kind"] == "upload"
        return action({"type": "confirm_sap", "confirmation_id": pending["confirmation_id"]})

    replies[:] = [confirm_from_context, "Importação iniciada."]
    ask(client, csrf, "sim")
    started = tool_result(sent[-1])
    assert started["status"] == "STARTED", started
    assert wait_job(client, started["result"]["job_id"])["status"] == "COMPLETED"
    with connect(database) as connection:
        assert connection.execute("select count(*) from inventory_rows where period='2026-09'").fetchone()[0] == 2


def test_common_user_cannot_list_or_import(env) -> None:
    _settings, _db, downloads, client_for, sent, replies = env
    zmm119_workbook(downloads / "export_zmm119.xlsx")
    client, csrf = client_for(USER)
    replies[:] = [action({"type": "find_sap_exports"}), "Sem permissão."]
    ask(client, csrf, "liste as planilhas baixadas")
    result = tool_result(sent[-1])
    assert result["status"] == "FORBIDDEN" and "files" not in result
