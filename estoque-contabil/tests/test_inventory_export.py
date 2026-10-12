"""Exportação da Base de Estoque: XLSX completo, gravado em fluxo e fora da pasta TEMP."""

from __future__ import annotations

import io
import json
from datetime import date

from openpyxl import load_workbook

from ops_contabil.db import connect
from ops_contabil.xlsx_stream import write_xlsx
from test_optimus_import import ADMIN, env  # noqa: F401 - fixture reaproveitada


def test_stream_writer_keeps_types_and_special_text(tmp_path) -> None:
    target = tmp_path / "saida.xlsx"
    rows = [(1, "CE & MAT <teste> \x07ok", 10.5, date(2026, 9, 30), None, 2026.0), (2, "", -3, None, "Ç", 0)]
    assert write_xlsx(target, "Base de Estoque", ["#", "Texto", "Valor", "Data", "Vazio", "Year"], rows) == 2
    sheet = load_workbook(target).active
    assert sheet.title == "Base de Estoque" and sheet.freeze_panes == "A2"
    assert [c.value for c in sheet[1]] == ["#", "Texto", "Valor", "Data", "Vazio", "Year"]
    first = [c.value for c in sheet[2]]
    assert first[0] == 1 and first[1] == "CE & MAT <teste> ok" and first[2] == 10.5 and first[3].date() == date(2026, 9, 30)
    assert first[4] is None and first[5] == 2026
    assert sheet["D2"].number_format == "dd/mm/yyyy"
    assert [c.value for c in sheet[3]][:3] == [2, None, -3] and sheet["E3"].value == "Ç"
    assert not list(tmp_path.glob("*.parcial"))


def test_export_route_returns_every_filtered_row(env) -> None:  # noqa: F811
    settings, database, _downloads, client_for, _sent, _replies = env
    with connect(database) as connection:
        connection.executemany(
            """insert into inventory_rows(period,source_row,material,plant,unrestricted_quantity,fiscal_total_amount,product_offer_end_date)
               values ('2026-09',?,?,?,?,?,?)""",
            [(i, f"MAT-{i}", "2115" if i % 2 else "1081", i, i * 1.5, date(2026, 1, 1)) for i in range(1, 2501)],
        )
    client, _csrf = client_for(ADMIN)
    response = client.get("/api/inventory/export", params={"period": "2026-09", "filters": json.dumps({"plant": "2115"})})
    assert response.status_code == 200, response.text
    assert response.headers["X-Exported-Rows"] == "1250"
    sheet = load_workbook(io.BytesIO(response.content), read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) == 1251 and rows[0][0] == "#" and rows[1][0] == 1
    exports = settings.path("processed").parent / "exports"
    assert exports.is_dir()  # temporário na pasta do sistema (não na TEMP)
