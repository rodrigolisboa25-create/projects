"""Inventaria planilhas Excel sem alterar os arquivos de origem.

Uso:
    python tools/inspect_workbooks.py arquivo1.xlsx arquivo2.xlsb --output mapa.json

Arquivos XLSB são lidos via Microsoft Excel/COM, em modo somente leitura.
Arquivos XLSX são lidos com openpyxl em modo streaming.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def trim_row(values: Iterable[Any]) -> list[Any]:
    row = [serializable(value) for value in values]
    while row and row[-1] is None:
        row.pop()
    return row


def normalize_formula(formula: str) -> str:
    formula = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "<CELL>", formula.upper())
    formula = re.sub(r"\d+(?:[.,]\d+)?", "<N>", formula)
    return formula[:500]


def inspect_xlsx(path: Path, sample_rows: int, sample_cols: int) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    result: dict[str, Any] = {
        "file": str(path.resolve()),
        "format": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "sheet_names": wb.sheetnames,
        "defined_names": sorted(str(name) for name in wb.defined_names),
        "sheets": [],
    }
    try:
        for ws in wb.worksheets:
            formula_count = 0
            formula_patterns: Counter[str] = Counter()
            formula_samples: list[dict[str, Any]] = []
            preview: list[list[Any]] = []
            nonempty_by_row: list[tuple[int, int]] = []

            for row_index, row in enumerate(ws.iter_rows(), start=1):
                nonempty = 0
                if row_index <= sample_rows:
                    preview.append(trim_row(cell.value for cell in row[:sample_cols]))
                for cell in row:
                    value = cell.value
                    if value is not None:
                        nonempty += 1
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        formula = str(value)
                        formula_count += 1
                        formula_patterns[normalize_formula(formula)] += 1
                        if len(formula_samples) < 100:
                            formula_samples.append({"cell": cell.coordinate, "formula": formula})
                if nonempty:
                    nonempty_by_row.append((row_index, nonempty))

            likely_header_rows = [
                {"row": row_number, "nonempty_cells": count}
                for row_number, count in sorted(nonempty_by_row[:50], key=lambda item: item[1], reverse=True)[:5]
            ]
            result["sheets"].append(
                {
                    "name": ws.title,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "formula_count": formula_count,
                    "top_formula_patterns": formula_patterns.most_common(30),
                    "formula_samples": formula_samples,
                    "likely_header_rows": likely_header_rows,
                    "preview": preview,
                }
            )
    finally:
        wb.close()
    return result


def _to_matrix(value: Any) -> list[list[Any]]:
    if value is None:
        return []
    if isinstance(value, tuple):
        if value and isinstance(value[0], tuple):
            return [trim_row(row) for row in value]
        return [trim_row(value)]
    return [[serializable(value)]]


def inspect_with_excel(path: Path, sample_rows: int, sample_cols: int) -> dict[str, Any]:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    excel.AutomationSecurity = 3
    workbook = None
    try:
        workbook = excel.Workbooks.Open(
            str(path.resolve()),
            UpdateLinks=0,
            ReadOnly=True,
            IgnoreReadOnlyRecommended=True,
            AddToMru=False,
            Notify=False,
        )
        result: dict[str, Any] = {
            "file": str(path.resolve()),
            "format": path.suffix.lower(),
            "size_bytes": path.stat().st_size,
            "sheet_names": [sheet.Name for sheet in workbook.Worksheets],
            "defined_names": [],
            "connections": [],
            "external_links": [],
            "sheets": [],
        }
        try:
            result["defined_names"] = [name.Name for name in workbook.Names]
        except Exception as exc:
            result["defined_names_error"] = str(exc)
        try:
            result["connections"] = [connection.Name for connection in workbook.Connections]
        except Exception as exc:
            result["connections_error"] = str(exc)
        try:
            links = workbook.LinkSources(1)
            if links:
                result["external_links"] = list(links) if isinstance(links, tuple) else [str(links)]
        except Exception as exc:
            result["external_links_error"] = str(exc)

        for ws in workbook.Worksheets:
            used = ws.UsedRange
            row_count = int(used.Rows.Count)
            col_count = int(used.Columns.Count)
            first_row = int(used.Row)
            first_col = int(used.Column)
            preview_rows = min(sample_rows, row_count)
            preview_cols = min(sample_cols, col_count)
            preview_range = ws.Range(
                ws.Cells(first_row, first_col),
                ws.Cells(first_row + preview_rows - 1, first_col + preview_cols - 1),
            )
            preview = _to_matrix(preview_range.Value2)

            formula_count = 0
            formula_samples: list[dict[str, Any]] = []
            formula_patterns: Counter[str] = Counter()
            try:
                formula_cells = used.SpecialCells(-4123)
                formula_count = int(formula_cells.CountLarge)
                sample_limit = min(formula_count, 250)
                for index in range(1, sample_limit + 1):
                    cell = formula_cells.Item(index)
                    formula = str(cell.Formula)
                    formula_samples.append({"cell": cell.Address(False, False), "formula": formula})
                    formula_patterns[normalize_formula(formula)] += 1
            except Exception:
                pass

            tables = []
            try:
                tables = [
                    {"name": table.Name, "range": table.Range.Address(False, False)}
                    for table in ws.ListObjects
                ]
            except Exception:
                pass
            query_tables = []
            try:
                query_tables = [table.Name for table in ws.QueryTables]
            except Exception:
                pass

            result["sheets"].append(
                {
                    "name": ws.Name,
                    "visible": int(ws.Visible),
                    "used_range": used.Address(False, False),
                    "first_row": first_row,
                    "first_column": first_col,
                    "max_row": first_row + row_count - 1,
                    "max_column": first_col + col_count - 1,
                    "formula_count": formula_count,
                    "top_formula_patterns": formula_patterns.most_common(30),
                    "formula_samples": formula_samples,
                    "tables": tables,
                    "query_tables": query_tables,
                    "preview": preview,
                }
            )
        return result
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()
        pythoncom.CoUninitialize()


def inspect(path: Path, sample_rows: int, sample_cols: int) -> dict[str, Any]:
    if path.suffix.lower() == ".xlsx":
        return inspect_xlsx(path, sample_rows, sample_cols)
    if path.suffix.lower() in {".xlsb", ".xls", ".xlsm"}:
        return inspect_with_excel(path, sample_rows, sample_cols)
    raise ValueError(f"Formato não suportado: {path.suffix}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+")
    parser.add_argument("--output", required=True)
    parser.add_argument("--sample-rows", type=int, default=30)
    parser.add_argument("--sample-cols", type=int, default=80)
    args = parser.parse_args()

    payload = []
    for file_name in args.files:
        path = Path(file_name)
        try:
            payload.append(inspect(path, args.sample_rows, args.sample_cols))
        except Exception as exc:
            payload.append({"file": str(path.resolve()), "error": f"{type(exc).__name__}: {exc}"})

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
