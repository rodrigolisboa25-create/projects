from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from ops_contabil.db import connect
from ops_contabil.ingestion.schema_resolver import load_contracts, resolve_headers
from ops_contabil.settings import load_settings


def main() -> None:
    settings = load_settings()
    period = "2026-08"
    source = Path(r"C:\Users\75649\AppData\Local\OpsContabil\landing\all_brazil\2026-08\ALL_BRAZIL_31_08.xlsx")
    with connect(settings.path("database")) as connection:
        inventory = connection.execute(
            """select upper(trim(material)), upper(substr(trim(material),1,10))
                 from inventory_rows where period=?""",
            [period],
        ).fetchall()
    full_counts = Counter(row[0] for row in inventory if row[0])
    style_counts = Counter(row[1] for row in inventory if row[1])
    wanted_full = set(full_counts)
    wanted_style = set(style_counts)

    contract = load_contracts(settings.root / "config" / "schemas.yaml")["all_brazil_data"]
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook["Data"]
        preview = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=contract.header_search_rows,
                values_only=True,
            )
        ]
        resolution = resolve_headers(preview, contract)
        positions = {
            field: resolution.positions[field]
            for field in (
                "material_number",
                "product_offer_end_date",
                "collection_code",
                "collection_year",
            )
        }
        found_full: dict[str, tuple[object, object, object]] = {}
        found_style: dict[str, tuple[object, object, object]] = {}
        full_conflicts: set[str] = set()
        style_conflicts: set[str] = set()
        for index, row in enumerate(
            worksheet.iter_rows(min_row=resolution.header_row + 1, values_only=True),
            start=1,
        ):
            raw_key = row[positions["material_number"]]
            if raw_key in (None, ""):
                continue
            key = str(raw_key).strip().upper()
            values = (
                row[positions["product_offer_end_date"]],
                row[positions["collection_code"]],
                row[positions["collection_year"]],
            )
            if key in wanted_full:
                if key in found_full and found_full[key] != values:
                    full_conflicts.add(key)
                else:
                    found_full[key] = values
            if key in wanted_style:
                if key in found_style and found_style[key] != values:
                    style_conflicts.add(key)
                else:
                    found_style[key] = values
            if index % 50000 == 0:
                print(f"{index:,} linhas analisadas", flush=True)
    finally:
        workbook.close()

    exact_rows = sum(full_counts[key] for key in found_full)
    style_rows = sum(style_counts[key] for key in found_style)
    hybrid_rows = sum(
        count
        for key, count in full_counts.items()
        if key in found_full or key[:10] in found_style
    )
    style_fallback_rows = sum(
        count
        for key, count in full_counts.items()
        if key not in found_full and key[:10] in found_style
    )
    print(
        {
            "inventory_rows": len(inventory),
            "inventory_full_keys": len(full_counts),
            "inventory_style_keys": len(style_counts),
            "exact_keys": len(found_full),
            "style_keys": len(found_style),
            "exact_rows": exact_rows,
            "style_rows": style_rows,
            "hybrid_rows": hybrid_rows,
            "hybrid_missing_rows": len(inventory) - hybrid_rows,
            "style_fallback_rows": style_fallback_rows,
            "full_conflicts": len(full_conflicts),
            "style_conflicts": len(style_conflicts),
            "sample_exact": list(found_full.items())[:3],
            "sample_style": list(found_style.items())[:3],
        },
        flush=True,
    )


if __name__ == "__main__":
    main()
