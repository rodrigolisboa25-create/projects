"""Coleção (Season/Year) dos pedidos de compra para o PASSO A PASSO.

No fechamento manual, a Base GR (MB59 do mês) recebe Season e Year por PROCV do pedido de compra no
export da SQVI DOC_COMPRA (cabeçalho do pedido, ME23N): "Código coleção" = Season e "Ano coleção" = Year.
Aqui o usuário envia um destes dois arquivos, do jeito que já produz:

* a própria Base GR montada (aba BASE, com Pedido, Season e Year já calculados), ou
* o export da SQVI DOC_COMPRA (Doc.compra, Código coleção, Ano coleção).

O sistema guarda a coleção por pedido (gr_collection_rows) e o PASSO A PASSO a usa na entrada mais recente
de cada SKU (mb59.enrich_inventory_pass_step). Pedidos intercompany não têm coleção e ficam de fora.
"""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import pyarrow as pa
from openpyxl import load_workbook

from .audit import sha256
from .db import connect
from .mb59 import INTERCOMPANY_VENDORS, _date, _text, ensure_mb59_schema, normalize_vendor

SEASONS = ("FA", "HO", "SP", "SU")
HEADER_SEARCH_ROWS = 15
SQVI_COMPANY = "7170"


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


# Cabeçalhos aceitos (normalizados: minúsculas, sem acento e sem pontuação).
BASE_GR_COLUMNS = {
    "purchase_order": ("pedido",),
    "season": ("season",),
    "year": ("year",),
    "vendor": ("fornecedor",),
    "posting_date": ("data de lancamento",),
    "material": ("material",),
}
SQVI_COLUMNS = {
    "purchase_order": ("doc compra", "n do documento de compras", "documento de compras", "doc compras"),
    "season": ("codigo colecao", "cod colecao"),
    "year": ("ano colecao",),
    "company": ("empr", "empresa"),
}


def _find_header(rows: list[list[Any]], columns: dict[str, tuple[str, ...]], required: tuple[str, ...]):
    for index, row in enumerate(rows):
        names = [_norm(value) for value in row]
        positions = {}
        for field, aliases in columns.items():
            for position, name in enumerate(names):
                if name in aliases:
                    positions[field] = position
                    break
        if all(field in positions for field in required):
            return index + 1, positions
    return None


def _season(value: Any) -> str | None:
    text = _text(value)
    text = text.strip().upper() if text else None
    return text if text in SEASONS else None


def _year(value: Any) -> int | None:
    try:
        number = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None
    return number if 2000 <= number <= 2100 else None


def identify_collection_file(path: Path) -> dict[str, Any] | None:
    """Reconhece a Base GR (aba BASE) ou o export da SQVI DOC_COMPRA pelas colunas. None se não for nenhum dos dois."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = sorted(workbook.worksheets, key=lambda ws: 0 if ws.title.strip().upper() == "BASE" else 1)
        for worksheet in sheets:
            preview = [list(row) for row in worksheet.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True)]
            found = _find_header(preview, BASE_GR_COLUMNS, ("purchase_order", "season", "year", "material"))
            if found:
                return {"kind": "base_gr", "sheet": worksheet.title, "header_row": found[0], "positions": found[1]}
            found = _find_header(preview, SQVI_COLUMNS, ("purchase_order", "season", "year"))
            if found:
                return {"kind": "sqvi", "sheet": worksheet.title, "header_row": found[0], "positions": found[1]}
        return None
    finally:
        workbook.close()


def _season_formula_by_sku(path: Path, sheet: str, header_row: int, season_position: int) -> bool:
    """True se a Season da Base GR foi buscada pelo SKU numa aba da própria planilha (ex.: PE), e não pelo pedido."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook[sheet]
        for row in worksheet.iter_rows(min_row=header_row + 1, max_row=header_row + 200, values_only=True):
            value = row[season_position] if season_position < len(row) else None
            if isinstance(value, str) and value.startswith("="):
                formula = value.upper()
                # PROCV do pedido no export da SQVI aponta para outro arquivo ([n]...); pelo SKU aponta para a aba PE.
                return "PE!" in formula or "'PE'!" in formula
        return False
    finally:
        workbook.close()


def read_collection_file(path: Path, period: str) -> dict[str, Any]:
    """Lê e valida o arquivo sem gravar nada. ValueError com a causa quando não puder ser usado."""
    datetime.strptime(period, "%Y-%m")
    detected = identify_collection_file(path)
    if detected is None:
        raise ValueError(
            "O arquivo não é a Base GR (aba BASE com Pedido, Season e Year) nem o export da SQVI DOC_COMPRA "
            "(Doc.compra, Código coleção e Ano coleção). Nada foi alterado."
        )
    kind, positions = detected["kind"], detected["positions"]
    if kind == "base_gr" and _season_formula_by_sku(path, detected["sheet"], detected["header_row"], positions["season"]):
        raise ValueError(
            "Na aba BASE deste arquivo, a Season foi buscada pelo SKU na aba PE, e não pelo pedido de compra na SQVI "
            "DOC_COMPRA. Assim a coleção pode vir de outro pedido. Envie a Base GR com o PROCV da SQVI pelo pedido "
            "(ou o próprio export da SQVI). Nada foi alterado."
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[detected["sheet"]]

        def cell(row: tuple[Any, ...], field: str) -> Any:
            position = positions.get(field)
            return None if position is None or position >= len(row) else row[position]

        values: dict[str, Counter] = {}
        lines = intercompany_lines = missing_lines = other_company = 0
        missing_orders: set[str] = set()
        months: Counter[str] = Counter()
        for row in worksheet.iter_rows(min_row=detected["header_row"] + 1, values_only=True):
            order = _text(cell(row, "purchase_order"))
            if not order:
                continue
            if kind == "base_gr" and not _text(cell(row, "material")):
                continue
            lines += 1
            if kind == "sqvi" and "company" in positions:
                company = _text(cell(row, "company"))
                if company and company.lstrip("0") != SQVI_COMPANY:
                    other_company += 1
                    continue
            if kind == "base_gr":
                posting = _date(cell(row, "posting_date"))
                if posting:
                    months[posting.strftime("%Y-%m")] += 1
                if normalize_vendor(cell(row, "vendor")) in INTERCOMPANY_VENDORS:
                    intercompany_lines += 1
                    continue
            season, year = _season(cell(row, "season")), _year(cell(row, "year"))
            if season is None or year is None:
                missing_lines += 1
                missing_orders.add(order)
                continue
            values.setdefault(order, Counter())[(season, year)] += 1
    finally:
        workbook.close()
    if not lines:
        raise ValueError("O arquivo não tem nenhuma linha com pedido de compra após o cabeçalho. Nada foi alterado.")
    if kind == "base_gr" and months and not months.get(period):
        found = ", ".join(month for month, _ in months.most_common(3))
        raise ValueError(
            f"A Base GR tem lançamentos de {found} e nenhum da competência {period}. "
            "Confira a competência escolhida ou o arquivo. Nada foi alterado."
        )
    if not values:
        raise ValueError(
            "Nenhum pedido do arquivo tem Season (FA, HO, SP ou SU) e Year (4 dígitos) válidos. Confira se o PROCV "
            "da SQVI foi calculado antes de salvar (sem #N/D). Nada foi alterado."
        )
    conflicts = {order: counts for order, counts in values.items() if len(counts) > 1}
    rows = [
        {"period": period, "purchase_order": order, "season": counts.most_common(1)[0][0][0],
         "year": counts.most_common(1)[0][0][1]}
        for order, counts in values.items()
    ]
    return {
        "kind": kind,
        "sheet": detected["sheet"],
        "rows": rows,
        "lines": lines,
        "orders": len(rows),
        "intercompany_lines": intercompany_lines,
        "missing_lines": missing_lines,
        "missing_orders": sorted(missing_orders - set(values))[:50],
        "missing_orders_count": len(missing_orders - set(values)),
        "other_company_lines": other_company,
        "conflicts": [
            {"purchase_order": order, "values": [f"{s}{y}" for (s, y), _ in counts.most_common()]}
            for order, counts in list(conflicts.items())[:20]
        ],
        "conflicts_count": len(conflicts),
    }


def ensure_month_entries(settings: Any, period: str, report: Callable[[str, str], None] | None = None) -> dict[str, Any] | None:
    """Mês com MB59 carregada antes da regra por SKU: reprocessa a MB59 já salva (sem extrair de novo) para montar as
    entradas por SKU e os pedidos. Nada a fazer se o mês já tem as entradas ou se a MB59 nunca foi carregada."""
    from .mb59 import import_mb59_xlsx

    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        if connection.execute("select count(*) from mb59_sku_entries where period=?", [period]).fetchone()[0]:
            return None
        row = connection.execute("select source_path from mb59_imports where period=?", [period]).fetchone()
    candidates = [Path(row[0])] if row and row[0] else []
    candidates.append(settings.path("landing") / "mb59" / period / f"MB59_{period}.xlsx")
    source = next((path for path in candidates if path.is_file()), None)
    if source is None:
        return None
    if report:
        report("IMPORTING", f"Preparando a MB59 já carregada de {period} para a regra por SKU (sem extrair de novo).")
    return import_mb59_xlsx(settings, period, source)


def import_collection(
    settings: Any,
    period: str,
    path: Path,
    imported_by: str | None = None,
    report: Callable[[str, str], None] | None = None,
) -> dict[str, Any]:
    """Valida e grava a coleção dos pedidos da competência (substitui a carga anterior do mesmo mês)."""
    path = Path(path).resolve()
    if report:
        report("IMPORTING", "Lendo a coleção (Season/Year) de cada pedido de compra.")
    parsed = read_collection_file(path, period)
    summary = {key: value for key, value in parsed.items() if key not in ("rows",)}
    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        connection.execute("begin")
        try:
            connection.execute("delete from gr_collection_rows where period=?", [period])
            connection.register("gr_collection_batch", pa.Table.from_pylist(parsed["rows"]))
            connection.execute(
                """insert into gr_collection_rows(period, purchase_order, season, year)
                   select period, purchase_order, season, year from gr_collection_batch"""
            )
            connection.execute("delete from gr_collection_imports where period=?", [period])
            connection.execute(
                """insert into gr_collection_imports(period, source_kind, file_name, source_path, source_sha256,
                       orders, lines, summary_json, imported_by) values (?,?,?,?,?,?,?,?,?)""",
                [period, parsed["kind"], path.name, str(path), sha256(path), parsed["orders"], parsed["lines"],
                 json.dumps(summary, ensure_ascii=False, default=str), imported_by],
            )
            connection.execute("commit")
        except Exception:
            connection.execute("rollback")
            raise
    # Competências antigas (julho, agosto, setembro…): a MB59 já carregada passa para a regra por SKU.
    prepared = ensure_month_entries(settings, period, report)
    status = collection_status(settings, period)
    return {**summary, "status": "imported", "period": period, "coverage": status["coverage"],
            "mb59_prepared": bool(prepared)}


def collection_status(settings: Any, period: str) -> dict[str, Any]:
    """Situação da coleção dos pedidos na competência: arquivo carregado e cobertura dos pedidos da MB59."""
    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        imported = connection.execute(
            """select source_kind, file_name, orders, lines, summary_json, imported_by, imported_at
                 from gr_collection_imports where period=?""",
            [period],
        ).fetchone()
        mb59_orders, covered = connection.execute(
            """select count(*), count(collection.purchase_order)
                 from mb59_orders orders
                 left join gr_collection_rows collection
                   on collection.period = orders.period and collection.purchase_order = orders.purchase_order
                where orders.period=? and not orders.intercompany""",
            [period],
        ).fetchone()
        missing = [row[0] for row in connection.execute(
            """select orders.purchase_order from mb59_orders orders
                 left join gr_collection_rows collection
                   on collection.period = orders.period and collection.purchase_order = orders.purchase_order
                where orders.period=? and not orders.intercompany and collection.purchase_order is null
                order by orders.lines desc, orders.purchase_order limit 30""",
            [period],
        ).fetchall()]
        sku_entries = int(connection.execute(
            "select count(*) from mb59_sku_entries where period=?", [period]).fetchone()[0])
    coverage = {
        "mb59_orders": int(mb59_orders or 0),
        "with_collection": int(covered or 0),
        "without_collection": int((mb59_orders or 0) - (covered or 0)),
        "pct": round(100 * covered / mb59_orders, 1) if mb59_orders else None,
        "missing_sample": missing,
    }
    result: dict[str, Any] = {
        "period": period,
        "imported": imported is not None,
        "mb59_ready": sku_entries > 0,
        "coverage": coverage,
    }
    if imported:
        result.update(
            source_kind=imported[0],
            source_label="Base GR (aba BASE)" if imported[0] == "base_gr" else "Export da SQVI DOC_COMPRA",
            file_name=imported[1],
            orders=int(imported[2]),
            lines=int(imported[3]),
            summary=json.loads(imported[4]) if imported[4] else {},
            imported_by=imported[5],
            imported_at=imported[6].isoformat() if hasattr(imported[6], "isoformat") else str(imported[6]),
        )
    return result


def orders_for_sqvi(settings: Any, period: str) -> list[str]:
    """Pedidos de compra da MB59 do mês (sem intercompany), para colar na seleção múltipla da SQVI DOC_COMPRA.
    Mês carregado antes da regra por SKU: lê os pedidos da MB59 já salva, sem gravar nada."""
    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        orders = [row[0] for row in connection.execute(
            "select purchase_order from mb59_orders where period=? and not intercompany order by purchase_order",
            [period],
        ).fetchall()]
        row = connection.execute("select source_path from mb59_imports where period=?", [period]).fetchone()
    if orders:
        return orders
    candidates = [Path(row[0])] if row and row[0] else []
    candidates.append(settings.path("landing") / "mb59" / period / f"MB59_{period}.xlsx")
    source = next((path for path in candidates if path.is_file()), None)
    return _orders_from_file(source) if source else []


def _orders_from_file(path: Path) -> list[str]:
    columns = {"purchase_order": ("pedido", "documento de compra", "ebeln"), "vendor": ("fornecedor", "lifnr")}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = sorted(workbook.worksheets, key=lambda ws: 0 if ws.title.strip().upper() == "BASE" else 1)
        for worksheet in sheets:
            preview = [list(r) for r in worksheet.iter_rows(min_row=1, max_row=30, values_only=True)]
            found = _find_header(preview, columns, ("purchase_order",))
            if not found:
                continue
            header_row, positions = found
            orders: set[str] = set()
            for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                order = _text(row[positions["purchase_order"]]) if positions["purchase_order"] < len(row) else None
                vendor = row[positions["vendor"]] if "vendor" in positions and positions["vendor"] < len(row) else None
                if order and normalize_vendor(vendor) not in INTERCOMPANY_VENDORS:
                    orders.add(order)
            return sorted(orders)
        return []
    finally:
        workbook.close()
