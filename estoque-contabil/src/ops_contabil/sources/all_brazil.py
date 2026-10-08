from __future__ import annotations

import hashlib
import os
import re
import shutil
from dataclasses import dataclass
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable

import pyarrow as pa
from openpyxl import load_workbook

from ..db import connect
from ..ingestion.schema_resolver import load_contracts, resolve_headers


MONTH_NAMES = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril", 5: "Maio", 6: "Junho",
    7: "Julho", 8: "Agosto", 9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro",
}
SNAPSHOT_PATTERN = re.compile(r"^ALL[_ ]BRAZIL[_ ](?P<day>\d{2})[_-](?P<month>\d{2})\.(?:xlsx|xlsb)$", re.I)


@dataclass(frozen=True)
class AllBrazilSnapshot:
    path: Path
    snapshot_date: date
    period: str


@lru_cache(maxsize=8)
def _discover_all_brazil_root(configured_root: str) -> Path:
    """Localiza o espelho do Drive sem assumir uma letra de unidade fixa."""
    candidates: list[Path] = []
    explicit = os.getenv("OPS_ALL_BRAZIL_ROOT", "").strip()
    if explicit:
        candidates.append(Path(os.path.expandvars(explicit)))

    expanded = os.path.expandvars(configured_root).strip()
    if expanded and "%" not in expanded:
        candidates.append(Path(expanded))

    if os.name == "nt":
        for letter in "CDEFGHIJKLMNOPQRSTUVWXYZ":
            drive = Path(f"{letter}:\\")
            candidates.extend(
                [
                    drive / "Drives compartilhados" / "@All Brazil Materials",
                    drive / "Shared drives" / "@All Brazil Materials",
                ]
            )

    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate).casefold()
        if key in seen:
            continue
        seen.add(key)
        try:
            if candidate.is_dir():
                return candidate.resolve()
        except OSError:
            continue

    return Path(expanded or explicit or "@All Brazil Materials")


def resolve_all_brazil_root(settings: Any) -> Path:
    config = settings.raw["sources"]["all_brazil_materials"]
    return _discover_all_brazil_root(str(config.get("local_root", "")))


MONTH_FOLDER_PATTERN = re.compile(r"^(?P<month>\d{2})\.")


def discover_all_brazil_months(root: Path) -> dict[int, dict[int, Path]]:
    """Anos e meses que existem de fato no Drive (pastas "AAAA" e "MM. Mês"), sem lista fixa.

    Um mês novo criado no Drive (ex.: 10. Outubro) aparece sozinho. Pastas fora do padrão são ignoradas.
    """
    found: dict[int, dict[int, Path]] = {}
    try:
        years = [item for item in root.iterdir() if item.is_dir() and item.name.isdigit() and len(item.name) == 4]
    except OSError:
        return found
    for year_folder in years:
        try:
            months = [item for item in year_folder.iterdir() if item.is_dir()]
        except OSError:
            continue
        for month_folder in months:
            match = MONTH_FOLDER_PATTERN.match(month_folder.name)
            if match and 1 <= int(match.group("month")) <= 12:
                found.setdefault(int(year_folder.name), {})[int(match.group("month"))] = month_folder
    return found


def resolve_all_brazil_month(settings: Any, as_of: date) -> list[AllBrazilSnapshot]:
    """Retorna todos os snapshots válidos do mês, do mais recente para o mais antigo."""
    root = resolve_all_brazil_root(settings)
    month_folder = root / str(as_of.year) / f"{as_of.month:02d}. {MONTH_NAMES[as_of.month]}"
    if not month_folder.exists():
        raise FileNotFoundError(f"Pasta All Brazil não encontrada para {as_of:%Y-%m}: {month_folder}")

    candidates: list[AllBrazilSnapshot] = []
    for path in month_folder.iterdir():
        if not path.is_file():
            continue
        match = SNAPSHOT_PATTERN.match(path.name)
        if not match:
            continue
        try:
            snapshot_date = date(as_of.year, int(match.group("month")), int(match.group("day")))
        except ValueError:
            continue
        if snapshot_date.month == as_of.month and snapshot_date <= as_of:
            candidates.append(AllBrazilSnapshot(path, snapshot_date, as_of.strftime("%Y-%m")))

    if not candidates:
        raise FileNotFoundError(
            f"Nenhum snapshot All Brazil válido em {month_folder} com data menor ou igual a {as_of:%d/%m/%Y}."
        )

    candidates.sort(
        key=lambda item: (item.snapshot_date, item.path.stat().st_mtime),
        reverse=True,
    )
    return candidates


def resolve_all_brazil(settings: Any, as_of: date) -> AllBrazilSnapshot:
    """Compatibilidade: retorna o snapshot diário mais recente da competência."""
    return resolve_all_brazil_month(settings, as_of)[0]


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stage_snapshot(settings: Any, snapshot: AllBrazilSnapshot) -> Path:
    landing = settings.path("landing") / "all_brazil" / snapshot.period
    landing.mkdir(parents=True, exist_ok=True)
    target = landing / snapshot.path.name
    if not target.exists() or target.stat().st_size != snapshot.path.stat().st_size:
        shutil.copy2(snapshot.path, target)
    return target


def _aging_bucket(days: int | None) -> str | None:
    if days is None:
        return None
    if days < 0:
        return "0) Futures"
    limits = (
        (104, "1) 0-3 meses"), (194, "2) 3-6 meses"), (284, "3) 6-9 meses"),
        (374, "4) 9-12 meses"), (554, "5) 12-18 meses"), (734, "6) 18-24 meses"),
        (1814, "7) 2-5 anos"), (19995, "8) >5 anos"),
    )
    return next((label for limit, label in limits if days <= limit), "9) Fora da faixa")


def resolve_material_fallback(
    material_keys: set[str],
    source_by_material: dict[str, dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    """Resolve Material exato antes do cadastro agregado por Estilo-Cor."""
    resolved: dict[str, dict[str, Any]] = {}
    for material_key in material_keys:
        if material_key in source_by_material:
            resolved[material_key] = source_by_material[material_key]
        elif material_key[:10] in source_by_material:
            resolved[material_key] = source_by_material[material_key[:10]]
    return resolved


def enrich_inventory_from_all_brazil(
    settings: Any,
    period: str,
    as_of: date,
    progress_callback: Callable[[float, str], None] | None = None,
) -> dict[str, Any]:
    """Enriquece a posição usando somente o último All Brazil válido do mês.

    Regra homologada:
    - selecionar um único arquivo: o snapshot de data mais recente da pasta mensal,
      desde que sua data seja menor ou igual à data-base;
    - relacionar somente por Material, nunca por Centro;
    - procurar primeiro o Material completo no Material Nbr;
    - quando o SKU completo não existir, usar o Estilo-Cor de 10 caracteres.
      Essa granularidade existe no cadastro All Brazil e é a regra da planilha
      legada; o fallback só ocorre depois da tentativa exata;
    - não usar arquivos anteriores do mesmo mês como fallback.
    """
    snapshot = resolve_all_brazil(settings, as_of)
    if snapshot.period != period:
        raise ValueError("A competência da posição e a data-base do All Brazil são divergentes.")

    staged = stage_snapshot(settings, snapshot)
    source_hash = _file_sha256(staged)
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["all_brazil_data"]

    with connect(settings.path("database")) as connection:
        keys = {
            str(row[0]).strip()
            for row in connection.execute(
                "select distinct upper(trim(material)) from inventory_rows where period=? and material is not null",
                [period],
            ).fetchall()
            if row[0]
        }
        style_keys = {key[:10] for key in keys}

    wanted_fields = (
        "division_description",
        "material_origin_description",
        "lifecycle",
        "lifecycle_description",
    )
    found: dict[str, dict[str, Any]] = {}
    duplicate_keys: set[str] = set()
    conflict_keys: set[str] = set()
    conflict_fields: set[tuple[str, str]] = set()
    seen_counts: dict[str, int] = {}

    workbook = load_workbook(staged, read_only=True, data_only=True)
    try:
        worksheet = next(
            (workbook[name] for name in contract.sheet_aliases if name in workbook.sheetnames),
            None,
        )
        if worksheet is None:
            raise ValueError(f"Aba Data não encontrada em {staged.name}")

        preview = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=contract.header_search_rows,
                values_only=True,
            )
        ]
        resolution = resolve_headers(preview, contract)
        key_position = resolution.positions["material_number"]
        positions = {field: resolution.positions.get(field) for field in wanted_fields}
        total_source_rows = max(1, int(getattr(worksheet, "max_row", 1) or 1) - resolution.header_row)
        if progress_callback is not None:
            progress_callback(0.0, f"Lendo {snapshot.path.name}: {total_source_rows:,} linhas disponíveis.")

        for processed, row in enumerate(
            worksheet.iter_rows(min_row=resolution.header_row + 1, values_only=True),
            start=1,
        ):
            if progress_callback is not None and (processed % 5000 == 0 or processed >= total_source_rows):
                progress_callback(
                    min(processed / total_source_rows, 0.86),
                    f"All Brazil: {processed:,} de {total_source_rows:,} linhas analisadas.",
                )
            if key_position >= len(row) or row[key_position] in (None, ""):
                continue
            key = str(row[key_position]).strip().upper()
            if key not in keys and key not in style_keys:
                continue

            seen_counts[key] = seen_counts.get(key, 0) + 1
            if seen_counts[key] > 1:
                duplicate_keys.add(key)

            bucket = found.setdefault(key, {})
            for field, position in positions.items():
                if position is None or position >= len(row):
                    continue
                raw = row[position]
                if raw in (None, ""):
                    continue
                value = str(raw).strip()
                if not value:
                    continue
                if field in bucket and bucket[field] != value:
                    conflict_fields.add((key, field))
                    conflict_keys.add(key)
                else:
                    bucket[field] = value
    finally:
        workbook.close()

    for key, field in conflict_fields:
        found.get(key, {}).pop(field, None)

    resolved = resolve_material_fallback(keys, found)
    matched_keys = set(resolved)
    exact_matches = sum(1 for key in matched_keys if key in found)
    style_fallback_matches = len(matched_keys) - exact_matches
    batch_rows = [
        {
            "material_key": key,
            "division_description": values.get("division_description"),
            "material_origin": values.get("material_origin_description"),
            "lifecycle": values.get("lifecycle"),
            "lifecycle_description": values.get("lifecycle_description"),
        }
        for key, values in resolved.items()
    ]

    coverage = round(100 * len(matched_keys) / len(keys), 2) if keys else 0.0
    if progress_callback is not None:
        progress_callback(0.90, f"All Brazil lido. Aplicando {len(matched_keys):,} materiais encontrados à posição.")

    with connect(settings.path("database")) as connection:
        connection.execute(
            """create table if not exists all_brazil_imports (
                period varchar primary key, snapshot_date date, source_path varchar, source_sha256 varchar,
                inventory_keys bigint, matched_keys bigint, coverage_pct double, duplicate_keys bigint,
                conflict_keys bigint, schema_fingerprint varchar, imported_at timestamp default current_timestamp
            )"""
        )

        # Sempre reaplica o enriquecimento. Isso evita cache falso quando a ZMM119
        # recria as linhas da competência com o mesmo conjunto de chaves.
        connection.execute(
            """update inventory_rows set
                 division_description=null,
                 material_origin=null,
                 lifecycle=null,
                 lifecycle_description=null
               where period=?""",
            [period],
        )

        if batch_rows:
            batch = pa.Table.from_pylist(
                batch_rows,
                schema=pa.schema([
                    ("material_key", pa.string()),
                    ("division_description", pa.string()),
                    ("material_origin", pa.string()),
                    ("lifecycle", pa.string()),
                    ("lifecycle_description", pa.string()),
                ]),
            )
            connection.register("ab_batch", batch)
            connection.execute(
                """update inventory_rows as inventory set
                    division_description=ab.division_description,
                    material_origin=ab.material_origin,
                    lifecycle=ab.lifecycle,
                    lifecycle_description=ab.lifecycle_description
                  from ab_batch as ab
                  where inventory.period=?
                    and upper(trim(inventory.material))=ab.material_key""",
                [period],
            )

        connection.execute("delete from all_brazil_imports where period=?", [period])
        connection.execute(
            """insert into all_brazil_imports(period,snapshot_date,source_path,source_sha256,inventory_keys,
               matched_keys,coverage_pct,duplicate_keys,conflict_keys,schema_fingerprint)
               values (?,?,?,?,?,?,?,?,?,?)""",
            [
                period,
                snapshot.snapshot_date,
                str(snapshot.path),
                source_hash,
                len(keys),
                len(matched_keys),
                coverage,
                len(duplicate_keys),
                len(conflict_keys),
                resolution.fingerprint,
            ],
        )

    if progress_callback is not None:
        progress_callback(1.0, f"All Brazil aplicado. Cobertura de {coverage}%.")

    return {
        "status": "enriched",
        "period": period,
        "snapshot": snapshot.path.name,
        "snapshot_date": snapshot.snapshot_date.isoformat(),
        "snapshots_available": 1,
        "snapshots_scanned": 1,
        "source_path": str(snapshot.path),
        "selection_rule": "latest_snapshot_on_or_before_as_of_only",
        "inventory_keys": len(keys),
        "matched_keys": len(matched_keys),
        "exact_material_matches": exact_matches,
        "style_color_fallback_matches": style_fallback_matches,
        "coverage_pct": coverage,
        "duplicate_keys": len(duplicate_keys),
        "conflict_keys": len(conflict_keys),
    }


def pass_step_lookup_from_all_brazil(
    settings: Any,
    as_of: date,
    material_keys: set[str],
    progress_callback: Callable[[float, str], None] | None = None,
) -> dict[str, Any]:
    """Retorna Product Offer End, Season e Year usando só o último All Brazil do mês.

    A chave não usa Centro. Primeiro tenta o Material completo da Base de Estoque
    contra Material Nbr. Se o SKU completo não existir no cadastro, tenta o
    Estilo-Cor de 10 caracteres. Arquivos anteriores do mesmo mês não são usados.
    """
    wanted = {str(value).strip().upper() for value in material_keys if value not in (None, "")}
    wanted_styles = {value[:10] for value in wanted}
    if not wanted:
        return {
            "lookup": {},
            "snapshot": None,
            "snapshot_date": None,
            "snapshots_available": 0,
            "snapshots_scanned": 0,
            "matched_keys": 0,
            "conflict_keys": 0,
        }

    snapshot = resolve_all_brazil(settings, as_of)
    staged = stage_snapshot(settings, snapshot)
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["all_brazil_data"]
    fields = ("product_offer_end_date", "collection_code", "collection_year")
    found: dict[str, dict[str, Any]] = {}
    conflicts: set[str] = set()
    conflict_fields: set[tuple[str, str]] = set()

    workbook = load_workbook(staged, read_only=True, data_only=True)
    try:
        worksheet = next(
            (workbook[name] for name in contract.sheet_aliases if name in workbook.sheetnames),
            None,
        )
        if worksheet is None:
            raise ValueError(f"Aba Data não encontrada em {staged.name}")

        preview = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=contract.header_search_rows,
                values_only=True,
            )
        ]
        resolution = resolve_headers(preview, contract)
        key_position = resolution.positions["material_number"]
        positions = {field: resolution.positions.get(field) for field in fields}
        total_source_rows = max(1, int(getattr(worksheet, "max_row", 1) or 1) - resolution.header_row)
        if progress_callback is not None:
            progress_callback(0.0, f"Consultando PASSO A PASSO no {snapshot.path.name}.")

        for processed, row in enumerate(
            worksheet.iter_rows(min_row=resolution.header_row + 1, values_only=True),
            start=1,
        ):
            if progress_callback is not None and (processed % 5000 == 0 or processed >= total_source_rows):
                progress_callback(
                    min(processed / total_source_rows, 0.96),
                    f"Consultando All Brazil: {processed:,} de {total_source_rows:,} linhas.",
                )
            if key_position >= len(row) or row[key_position] in (None, ""):
                continue
            key = str(row[key_position]).strip().upper()
            if key not in wanted and key not in wanted_styles:
                continue

            bucket = found.setdefault(key, {})
            for field, position in positions.items():
                if position is None or position >= len(row):
                    continue
                raw = row[position]
                if raw in (None, ""):
                    continue

                if field == "product_offer_end_date":
                    value = raw.date() if hasattr(raw, "date") else raw
                    if not isinstance(value, date):
                        continue
                elif field == "collection_year":
                    try:
                        value = int(float(raw))
                    except (TypeError, ValueError):
                        continue
                else:
                    value = str(raw).strip()
                    if not value:
                        continue

                if field in bucket and bucket[field] != value:
                    conflict_fields.add((key, field))
                    conflicts.add(key)
                else:
                    bucket[field] = value
    finally:
        workbook.close()

    for key, field in conflict_fields:
        found.get(key, {}).pop(field, None)

    resolved = resolve_material_fallback(wanted, found)
    lookup = {
        material_key: {
            "product_offer_end_date": values.get("product_offer_end_date"),
            "season": values.get("collection_code"),
            "year": values.get("collection_year"),
        }
        for material_key, values in resolved.items()
    }
    matched_keys = set(lookup)
    exact_matches = sum(1 for key in matched_keys if key in found)
    style_fallback_matches = len(matched_keys) - exact_matches

    if progress_callback is not None:
        progress_callback(1.0, f"Consulta PASSO A PASSO no All Brazil concluída: {len(matched_keys):,} materiais encontrados.")

    return {
        "lookup": lookup,
        "snapshot": snapshot.path.name,
        "snapshot_date": snapshot.snapshot_date.isoformat(),
        "snapshots_available": 1,
        "snapshots_scanned": 1,
        "matched_keys": len(matched_keys),
        "exact_material_matches": exact_matches,
        "style_color_fallback_matches": style_fallback_matches,
        "conflict_keys": len(conflicts),
        "selection_rule": "latest_snapshot_on_or_before_as_of_only",
    }


def list_all_brazil_snapshots(
    settings: Any,
    *,
    year: int | None = None,
    month: int | None = None,
    limit: int = 500,
) -> dict[str, Any]:
    """Lista os snapshots All Brazil disponíveis no espelho local do Drive."""
    root = resolve_all_brazil_root(settings)
    snapshots: list[dict[str, Any]] = []

    if not root.is_dir():
        return {"root": str(root), "count": 0, "snapshots": []}

    for year_folder in root.iterdir():
        if not year_folder.is_dir() or not year_folder.name.isdigit():
            continue
        snapshot_year = int(year_folder.name)
        if year is not None and snapshot_year != year:
            continue

        for month_folder in year_folder.iterdir():
            if not month_folder.is_dir():
                continue
            month_match = re.match(r"^(?P<month>\d{2})\.", month_folder.name)
            if not month_match:
                continue
            snapshot_month = int(month_match.group("month"))
            if month is not None and snapshot_month != month:
                continue

            for path in month_folder.iterdir():
                if not path.is_file():
                    continue
                match = SNAPSHOT_PATTERN.match(path.name)
                if not match:
                    continue
                try:
                    snapshot_date = date(
                        snapshot_year,
                        int(match.group("month")),
                        int(match.group("day")),
                    )
                except ValueError:
                    continue

                snapshots.append(
                    {
                        "snapshot_date": snapshot_date.isoformat(),
                        "period": snapshot_date.strftime("%Y-%m"),
                        "file_name": path.name,
                        "source_path": str(path),
                        "size_bytes": path.stat().st_size,
                    }
                )

    snapshots.sort(
        key=lambda item: (
            item["snapshot_date"],
            item["file_name"].casefold(),
        ),
        reverse=True,
    )
    safe_limit = max(1, min(int(limit), 2000))
    return {
        "root": str(root),
        "count": len(snapshots),
        "returned": min(len(snapshots), safe_limit),
        "snapshots": snapshots[:safe_limit],
    }


def query_all_brazil_snapshot(
    settings: Any,
    *,
    as_of: date,
    material_numbers: list[str] | None = None,
    q: str = "",
    limit: int = 50,
) -> dict[str, Any]:
    """Consulta atributos canônicos em um snapshot histórico All Brazil."""
    snapshot = resolve_all_brazil(settings, as_of)
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["all_brazil_data"]
    wanted_materials = {
        str(value).strip()
        for value in (material_numbers or [])
        if value not in (None, "")
    }
    query = str(q or "").strip().casefold()
    safe_limit = max(1, min(int(limit), 100))

    workbook = load_workbook(snapshot.path, read_only=True, data_only=True)
    try:
        worksheet = next(
            (workbook[name] for name in contract.sheet_aliases if name in workbook.sheetnames),
            None,
        )
        if worksheet is None:
            raise ValueError(f"Aba Data não encontrada em {snapshot.path.name}")

        preview = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=contract.header_search_rows,
                values_only=True,
            )
        ]
        resolution = resolve_headers(preview, contract)
        available_fields = [
            column.canonical
            for column in contract.columns
            if column.canonical in resolution.positions
        ]
        key_position = resolution.positions["material_number"]

        rows: list[dict[str, Any]] = []
        matched_total = 0

        for source_row in worksheet.iter_rows(
            min_row=resolution.header_row + 1,
            values_only=True,
        ):
            if key_position >= len(source_row):
                continue
            key_raw = source_row[key_position]
            if key_raw in (None, ""):
                continue
            material_number = str(key_raw).strip()

            if wanted_materials and material_number not in wanted_materials:
                continue

            record: dict[str, Any] = {}
            for field in available_fields:
                position = resolution.positions[field]
                value = source_row[position] if position < len(source_row) else None
                if hasattr(value, "isoformat"):
                    try:
                        value = value.isoformat()
                    except Exception:
                        value = str(value)
                record[field] = value

            if query:
                searchable = " | ".join(
                    "" if value is None else str(value)
                    for value in record.values()
                ).casefold()
                if query not in searchable:
                    continue

            matched_total += 1
            if len(rows) < safe_limit:
                rows.append(record)

            if wanted_materials and not query:
                found_keys = {
                    str(item.get("material_number", "")).strip()
                    for item in rows
                }
                if wanted_materials.issubset(found_keys):
                    break

        return {
            "snapshot_date": snapshot.snapshot_date.isoformat(),
            "period": snapshot.period,
            "file_name": snapshot.path.name,
            "source_path": str(snapshot.path),
            "selection_rule": "single_historical_snapshot_latest_on_or_before_as_of",
            "join_rule": (
                "Material completo da Base de Estoque = Material Nbr do All Brazil, "
                "sem Centro e sem corte para Estilo-Cor"
            ),
            "requested_as_of": as_of.isoformat(),
            "material_numbers": sorted(wanted_materials),
            "q": q,
            "available_fields": available_fields,
            "matched_total": matched_total,
            "returned": len(rows),
            "truncated": matched_total > len(rows),
            "rows": rows,
        }
    finally:
        workbook.close()
