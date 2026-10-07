"""Carga manual das planilhas ZMM119/MB59 exportadas do SAP.

Contingência ao robô SAP GUI: o usuário exporta a transação por conta própria e
envia o XLSX. O arquivo é validado por inteiro antes de substituir o destino
oficial em landing e antes de qualquer alteração no DuckDB; a importação em si
reutiliza exatamente as mesmas funções da extração automática.
"""

from __future__ import annotations

import os
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .ingestion.schema_resolver import load_contracts, resolve_headers
from .inventory import INVENTORY_COMPANY, _clean_value, _zmm119_positional_resolution
from .mb59 import _date, _text

SAP_UPLOAD_MAX_BYTES = 500 * 1024 * 1024
SAP_UPLOAD_TRANSACTIONS = ("ZMM119", "MB59")


def landing_target(settings: Any, transaction: str, period: str) -> Path:
    """Destino oficial usado também pelo robô e pelo 'Recarregar base'."""
    name = transaction.upper()
    return settings.path("landing") / name.lower() / period / f"{name}_{period}.xlsx"


def is_xlsx_archive(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        return "[Content_Types].xml" in names and "xl/workbook.xml" in names
    except (OSError, zipfile.BadZipFile):
        return False


def identify_sap_export(settings: Any, path: Path) -> str | None:
    """Reconhece pela estrutura de colunas se o XLSX é uma ZMM119, uma MB59 ou nenhuma."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.worksheets:
            try:
                _zmm119_positional_resolution(workbook.worksheets[0])
                return "ZMM119"
            except ValueError:
                pass
        contract = load_contracts(settings.root / "config" / "schemas.yaml")["mb59_base"]
        for worksheet in workbook.worksheets:
            preview = [
                list(row)
                for row in worksheet.iter_rows(min_row=1, max_row=contract.header_search_rows, values_only=True)
            ]
            try:
                resolve_headers(preview, contract)
                return "MB59"
            except Exception:
                continue
        return None
    finally:
        workbook.close()


def wrong_transaction_message(expected: str, detected: str | None) -> str:
    if detected:
        return (
            f"Esta planilha é da {detected}, não da {expected}. Nada foi alterado. "
            f"Para carregá-la, use o quadro de upload da {detected}; para a {expected}, "
            f"selecione o arquivo exportado da {expected}."
        )
    other = "MB59" if expected == "ZMM119" else "ZMM119"
    return (
        f"Esta planilha não é uma exportação da {expected} nem da {other}: as colunas não "
        f"correspondem ao layout das transações do SAP. Nada foi alterado. Confira o arquivo selecionado."
    )


def validate_sap_upload(settings: Any, transaction: str, period: str, path: Path) -> dict[str, Any]:
    """Valida o XLSX enviado sem gravar nada. Lança ValueError com a causa."""
    transaction = transaction.upper()
    if transaction not in SAP_UPLOAD_TRANSACTIONS:
        raise ValueError(f"Transação não suportada para upload: {transaction}.")
    datetime.strptime(period, "%Y-%m")
    if not is_xlsx_archive(path):
        raise ValueError(
            "O arquivo enviado não é uma planilha XLSX válida. No SAP, exporte a grade "
            "no formato XLSX (Excel 2007 ou superior) e envie esse arquivo."
        )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if transaction == "ZMM119":
            return _validate_zmm119(workbook)
        return _validate_mb59(settings, period, workbook)
    finally:
        workbook.close()


def _validate_zmm119(workbook: Any) -> dict[str, Any]:
    if not workbook.worksheets:
        raise ValueError("A planilha enviada não possui nenhuma aba.")
    worksheet = workbook.worksheets[0]
    try:
        _zmm119_positional_resolution(worksheet)
    except ValueError as exc:
        raise ValueError(
            f"A planilha não corresponde ao layout da ZMM119. {exc} "
            "Confira se o arquivo é a exportação da ZMM119 (e não da MB59)."
        ) from exc
    for row in worksheet.iter_rows(min_row=2, max_col=3, values_only=True):
        # Mesma normalização da importação (o Excel pode trazer 7170 como número).
        if len(row) >= 3 and row[2] not in (None, "") and str(_clean_value("company", row[0]) or "").strip() == INVENTORY_COMPANY:
            return {"transaction": "ZMM119", "sheet": str(worksheet.title)}
    raise ValueError(
        f"A planilha ZMM119 não possui nenhum registro da empresa {INVENTORY_COMPANY}. "
        "Nada foi alterado."
    )


def _validate_mb59(settings: Any, period: str, workbook: Any) -> dict[str, Any]:
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["mb59_base"]
    best: tuple[Any, Any] | None = None
    for worksheet in workbook.worksheets:
        preview = [
            list(row)
            for row in worksheet.iter_rows(min_row=1, max_row=contract.header_search_rows, values_only=True)
        ]
        try:
            candidate = resolve_headers(preview, contract)
        except Exception:
            continue
        if best is None or len(candidate.positions) > len(best[1].positions):
            best = (worksheet, candidate)
    if best is None:
        raise ValueError(
            "A planilha não corresponde ao layout da MB59 (Base GR): as colunas obrigatórias "
            "não foram encontradas. Confira se o arquivo é a exportação da MB59 (e não da ZMM119)."
        )
    worksheet, resolution = best

    def value(row: tuple[Any, ...], field: str) -> Any:
        position = resolution.positions.get(field)
        return None if position is None or position >= len(row) else row[position]

    movements = dated = in_period = 0
    other_months: Counter[str] = Counter()
    for row in worksheet.iter_rows(min_row=resolution.header_row + 1, values_only=True):
        if not _text(value(row, "material")) or not _text(value(row, "plant")):
            continue
        movements += 1
        posting_date = _date(value(row, "posting_date"))
        if posting_date is None:
            continue
        dated += 1
        month = posting_date.strftime("%Y-%m")
        if month == period:
            in_period += 1
        else:
            other_months[month] += 1
    if not movements:
        raise ValueError(
            "A planilha MB59 não possui movimentos (Material e Centro) após o cabeçalho. Nada foi alterado."
        )
    if dated and not in_period:
        found = ", ".join(month for month, _ in other_months.most_common(3))
        raise ValueError(
            f"A planilha MB59 contém lançamentos de {found} e nenhum da competência {period}. "
            "Confira a competência selecionada ou o arquivo enviado. Nada foi alterado."
        )
    return {"transaction": "MB59", "sheet": str(worksheet.title), "movements": movements, "in_period": in_period}


def install_upload(staging: Path, target: Path) -> Path | None:
    """Coloca o XLSX validado no destino oficial, preservando o anterior."""
    target.parent.mkdir(parents=True, exist_ok=True)
    backup: Path | None = None
    if target.exists():
        backup = target.with_name(f"{target.stem}.anterior{target.suffix}")
        os.replace(target, backup)
    os.replace(staging, target)
    return backup


def restore_previous(target: Path, backup: Path | None) -> None:
    """Desfaz install_upload quando a importação falha antes de gravar no banco."""
    if backup is not None and backup.exists():
        os.replace(backup, target)
    else:
        target.unlink(missing_ok=True)


# ----------------------------------------------- planilhas já baixadas (Optimus)

_XML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _column_index(reference: str) -> int:
    number = 0
    for char in reference:
        if not char.isalpha():
            break
        number = number * 26 + (ord(char.upper()) - 64)
    return number - 1


def _sheet_paths(archive: zipfile.ZipFile) -> list[str]:
    """Caminhos das abas na ordem do arquivo (a primeira é a que a ZMM119 usa)."""
    try:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {relation.get("Id"): str(relation.get("Target") or "") for relation in relations}
        paths = []
        for sheet in workbook.iter(f"{_XML_NS}sheet"):
            target = targets.get(sheet.get(f"{_REL_NS}id"), "").lstrip("/")
            if target:
                paths.append(target if target.startswith("xl/") else f"xl/{target}")
        if paths:
            return paths
    except (KeyError, ET.ParseError):
        pass
    return sorted(name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))


def _shared_strings(archive: zipfile.ZipFile, last_index: int) -> list[str]:
    """Lê só os primeiros textos compartilhados (os do cabeçalho vêm primeiro no arquivo)."""
    values: list[str] = []
    try:
        handle = archive.open("xl/sharedStrings.xml")
    except KeyError:
        return values
    with handle:
        for _event, element in ET.iterparse(handle, events=("end",)):
            if element.tag == f"{_XML_NS}si":
                values.append("".join(text.text or "" for text in element.iter(f"{_XML_NS}t")))
                element.clear()
                if len(values) > last_index:
                    break
    return values


def _preview_rows(archive: zipfile.ZipFile, sheet_path: str, max_rows: int) -> tuple[list[list[Any]], int]:
    """Primeiras linhas de uma aba, lidas direto do XML sem carregar a planilha inteira."""
    rows: list[dict[int, Any]] = []
    needed: set[int] = set()
    declared_columns = 0
    with archive.open(sheet_path) as handle:
        for event, element in ET.iterparse(handle, events=("start", "end")):
            if event == "start" and element.tag == f"{_XML_NS}dimension":
                declared_columns = _column_index(str(element.get("ref") or "A1").split(":")[-1]) + 1
            elif event == "end" and element.tag == f"{_XML_NS}row":
                cells: dict[int, Any] = {}
                for cell in element.iter(f"{_XML_NS}c"):
                    reference = cell.get("r")
                    index = _column_index(reference) if reference else len(cells)
                    kind = cell.get("t")
                    value: Any = None
                    if kind == "inlineStr":
                        value = "".join(text.text or "" for text in cell.iter(f"{_XML_NS}t"))
                    else:
                        raw = cell.find(f"{_XML_NS}v")
                        if raw is not None and raw.text is not None:
                            if kind == "s":
                                value = ("shared", int(raw.text))
                                needed.add(int(raw.text))
                            elif kind in ("str", "e"):
                                value = raw.text
                            elif kind == "b":
                                value = raw.text == "1"
                            else:
                                try:
                                    number = float(raw.text)
                                    value = int(number) if number.is_integer() else number
                                except ValueError:
                                    value = raw.text
                    cells[index] = value
                rows.append(cells)
                element.clear()
                if len(rows) >= max_rows:
                    break
    shared = _shared_strings(archive, max(needed)) if needed else []
    width = max([declared_columns, *[max(cells) + 1 for cells in rows if cells]])
    preview = []
    for cells in rows:
        row: list[Any] = [None] * width
        for index, value in cells.items():
            if isinstance(value, tuple):
                value = shared[value[1]] if value[1] < len(shared) else None
            row[index] = value
        preview.append(row)
    return preview, width


class _PreviewSheet:
    """O mínimo de uma aba do openpyxl para reaproveitar a validação posicional da ZMM119."""

    def __init__(self, rows: list[list[Any]], max_column: int) -> None:
        self.rows = rows
        self.max_column = max_column

    def iter_rows(self, min_row: int = 1, max_row: int | None = None, min_col: int = 1, max_col: int | None = None,
                  values_only: bool = True):
        for row in self.rows[min_row - 1:max_row]:
            padded = row + [None] * max(0, (max_col or len(row)) - len(row))
            yield tuple(padded[min_col - 1:max_col])


def quick_identify_sap_export(settings: Any, path: Path) -> str | None:
    """Igual a identify_sap_export, mas lendo só o cabeçalho (segundos em vez de minutos em planilhas grandes)."""
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["mb59_base"]
    max_rows = max(1, int(contract.header_search_rows))
    with zipfile.ZipFile(path) as archive:
        for position, sheet_path in enumerate(_sheet_paths(archive)):
            rows, width = _preview_rows(archive, sheet_path, max_rows)
            if position == 0:
                try:
                    _zmm119_positional_resolution(_PreviewSheet(rows, width))
                    return "ZMM119"
                except ValueError:
                    pass
            try:
                resolve_headers(rows, contract)
                return "MB59"
            except Exception:  # noqa: BLE001 - mesma regra de identify_sap_export
                continue
    return None


SAP_EXPORT_SEARCH_DAYS = 45
SAP_EXPORT_SEARCH_LIMIT = 15


def export_search_folders() -> list[Path]:
    """Pastas onde o SAP costuma gravar as exportações (as mesmas que o robô vasculha)."""
    home = Path.home()
    folders = [home / "Downloads", home / "Desktop", home / "Documents"]
    for onedrive in sorted(home.glob("OneDrive*")):
        folders.extend(onedrive / name for name in ("Desktop", "Área de Trabalho", "Documents", "Documentos"))
    for env_name in ("TEMP", "TMP"):
        value = os.getenv(env_name, "").strip()
        if value:
            folders.append(Path(value))
    for letter in "CDEFGH":
        folders.extend((Path(f"{letter}:\\TEMP"), Path(f"{letter}:\\TMP")))
    unique: list[Path] = []
    seen: set[str] = set()
    for folder in folders:
        try:
            if not folder.is_dir():
                continue
            key = str(folder.resolve()).casefold()
        except OSError:
            continue
        if key not in seen:
            seen.add(key)
            unique.append(folder)
    return unique


def find_sap_export_candidates(
    settings: Any,
    *,
    days: int = SAP_EXPORT_SEARCH_DAYS,
    limit: int = SAP_EXPORT_SEARCH_LIMIT,
    folders: list[Path] | None = None,
) -> list[dict[str, Any]]:
    """Planilhas .xlsx recentes dessas pastas, com a transação reconhecida pelo layout (ou None)."""
    cutoff = datetime.now().timestamp() - days * 86400
    found: list[tuple[float, int, Path]] = []
    for folder in folders if folders is not None else export_search_folders():
        try:
            entries = list(folder.iterdir())
        except OSError:
            continue
        for path in entries:
            if not path.name.casefold().endswith(".xlsx") or path.name.startswith("~$"):
                continue
            try:
                stat = path.stat()
            except OSError:
                continue
            if not path.is_file() or stat.st_mtime < cutoff or not 0 < stat.st_size <= SAP_UPLOAD_MAX_BYTES:
                continue
            found.append((stat.st_mtime, stat.st_size, path))
    found.sort(key=lambda item: item[0], reverse=True)
    result: list[dict[str, Any]] = []
    for modified, size, path in found[:limit]:
        transaction = None
        if is_xlsx_archive(path):
            try:
                transaction = quick_identify_sap_export(settings, path)
            except Exception:  # noqa: BLE001 - XML fora do comum: usa a leitura completa
                try:
                    transaction = identify_sap_export(settings, path)
                except Exception:  # noqa: BLE001 - planilha corrompida ou aberta no Excel
                    transaction = None
        result.append({
            "path": str(path),
            "file_name": path.name,
            "folder": str(path.parent),
            "size_bytes": size,
            "modified_at": datetime.fromtimestamp(modified).isoformat(timespec="seconds"),
            "transaction": transaction,
        })
    return result
