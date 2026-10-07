from __future__ import annotations

import re
import unicodedata
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalize_column(value: Any, position: int) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()
    return text or f"column_{position}"


def iter_xlsx(path: Path, sheet_name: str, header_row: int = 1) -> Iterator[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        for _ in range(header_row - 1):
            next(rows, None)
        raw_header = next(rows)
        header = [normalize_column(value, index) for index, value in enumerate(raw_header, start=1)]
        for source_row, values in enumerate(rows, start=header_row + 1):
            if not any(value is not None for value in values):
                continue
            record = dict(zip(header, values, strict=False))
            record["_source_row"] = source_row
            yield record
    finally:
        workbook.close()
