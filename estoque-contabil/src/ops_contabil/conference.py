"""Conferência com o fechamento manual (agente de conferência, Fases 1 e 2).

Fase 1 — Conferência: a usuária envia a Posição de Estoque fechada à mão (.xlsb ou .xlsx, aba Base Orig.)
e o sistema compara, linha a linha (CE & MAT), Year, Season e Product Offer End Dt com o que ele calculou,
além dos totais por AGING, Season e Year. Cada diferença é classificada pela origem do valor no sistema
(Base GR, mês anterior, All Brazil, manual) e marcada como "sem origem" quando o valor da planilha não existe
em nenhuma das fontes permitidas (por exemplo, erro de fórmula na planilha).

Fase 2 — Diagnóstico: testa, contra a planilha, todas as combinações das regras ajustáveis do PASSO A PASSO
(pass_step_rules: passo 1 na Base GR e passo 2 no mês anterior) e sugere a que mais reduz a diferença em R$.
A sugestão só vale depois de aprovada por um administrador (Fase 3, em app.py).

Nada aqui altera dados: a conferência lê o banco, guarda o resultado em ``conference_runs`` e as linhas da
planilha em ``processed/conferences/<id>.parquet``. Roda em segundo plano no servidor local (o n8n só recebe
o resumo pronto, pela fila de avisos do Optimus).
"""

from __future__ import annotations

import json
import os
import time
from calendar import monthrange
from dataclasses import replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable

import duckdb
import pyarrow as pa
import pyarrow.parquet as pq

from .db import connect
from .ingestion.schema_resolver import SchemaContract, load_contracts, resolve_headers
from .inventory import INVENTORY_COMPANY, YEAR_BEFORE_2022, _clean_value
from .master_base import master_base_in_use
from .pass_step_rules import (FIELD_LABELS, FIELDS, STEP1_OPTIONS, STEP2_OPTIONS, load_rules, step1_sql, step2_sql)
from .sources.all_brazil import pass_step_lookup_from_all_brazil

WANTED = ("center_material_key", "company", "plant", "material", "product_offer_end_date", "season", "year", "fiscal_total_amount")
# Sugestão só quando reduz de fato a diferença: ao menos R$ 100 mil e 5% da diferença atual do campo.
MIN_GAIN_VALUE = 100_000.0
MIN_GAIN_SHARE = 0.05
Progress = Callable[[int, str], None]


def ensure_conference_schema(connection: Any) -> None:
    connection.execute(
        """create table if not exists conference_runs (
               conference_id varchar primary key,
               period varchar not null,
               file_name varchar,
               source_sha256 varchar,
               status varchar not null,
               message varchar,
               created_by varchar,
               created_at timestamp not null default current_timestamp,
               finished_at timestamp,
               summary_json varchar,
               diagnosis_json varchar
           )"""
    )


def conferences_folder(settings: Any) -> Path:
    folder = Path(settings.path("processed")) / "conferences"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


# ------------------------------------------------------------------ leitura da planilha manual


# Layout aceito na conferência: só as colunas que ela usa, reconhecidas pelo nome (mesmos nomes da Posição de
# Estoque), em qualquer ordem e com o cabeçalho em qualquer uma das 30 primeiras linhas. Obrigatórias: Material e
# Vlr Tot Fisc; Centro ou CE & MAT; e ao menos uma entre Product Offer End Dt, Season e Year. Empresa é opcional
# (quando existe, só a 7170 entra).
CONFERENCE_COLUMNS = ("center_material_key", "company", "plant", "material", "fiscal_total_amount",
                      "product_offer_end_date", "season", "year")
CONFERENCE_REQUIRED = ("material", "fiscal_total_amount")
LAYOUT_HELP = ("Layout aceito: Posição de Estoque em .xlsb ou .xlsx (de preferência a aba Base Orig.), com as colunas Material, "
               "Vlr Tot Fisc, Centro (ou CE & MAT) e ao menos uma entre Product Offer End Dt, Season e Year; Empresa é opcional. "
               "A ordem das colunas e as colunas a mais não importam.")


def conference_contract(settings: Any) -> SchemaContract:
    base = load_contracts(settings.root / "config" / "schemas.yaml")["zmm119_base_orig"]
    columns = tuple(replace(column, required=column.canonical in CONFERENCE_REQUIRED)
                    for column in base.columns if column.canonical in CONFERENCE_COLUMNS)
    return replace(base, columns=columns)


def _check_layout(positions: dict[str, int], sheet: str) -> None:
    missing = []
    if "plant" not in positions and "center_material_key" not in positions:
        missing.append("Centro (ou CE & MAT)")
    if not any(name in positions for name in ("product_offer_end_date", "season", "year")):
        missing.append("ao menos uma entre Product Offer End Dt, Season e Year")
    if missing:
        raise ValueError(f"A aba {sheet} não tem: {', '.join(missing)}. {LAYOUT_HELP}")


def read_position_file(settings: Any, path: Path, progress: Progress | None = None, expected_rows: int = 0) -> tuple[pa.Table, dict[str, Any]]:
    """Lê só as colunas da conferência (.xlsb ou .xlsx), linha a linha, informando o andamento."""
    contract = conference_contract(settings)
    report = progress or (lambda _p, _m: None)
    suffix = path.suffix.lower()
    columns: dict[str, list[Any]] = {name: [] for name in ("k", "material", "plant", "pod", "season", "year", "value", "source_row")}
    info: dict[str, Any] = {"discarded_company_rows": 0}

    def pick_sheet(names: list[str], preview_of: Callable[[str], list[list[Any]]]) -> tuple[str, Any]:
        """A aba Base Orig. (ou BASE - CKMR) primeiro; senão, a primeira aba que tenha as colunas da conferência."""
        ordered = [n for n in contract.sheet_aliases if n in names] + [n for n in names if n not in contract.sheet_aliases]
        last_error: Exception | None = None
        for name in ordered:
            try:
                resolution = resolve_headers(preview_of(name), contract)
                _check_layout(resolution.positions, name)
                return name, resolution
            except Exception as exc:  # noqa: BLE001 - tenta a próxima aba
                last_error = exc
        if isinstance(last_error, ValueError) and "não tem:" in str(last_error):
            raise last_error
        raise ValueError(f"Nenhuma aba da planilha tem as colunas da conferência (Material e Vlr Tot Fisc). {LAYOUT_HELP}")

    def consume(rows_iter: Any, header_row: int, positions: dict[str, int]) -> None:
        for source_row, row in enumerate(rows_iter, start=1):
            if source_row <= header_row:
                continue
            if source_row % 25000 == 0:
                share = min(1.0, source_row / expected_rows) if expected_rows else 0.5
                report(5 + round(23 * share), f"Lendo a planilha: {source_row:,} de ~{expected_rows:,} linhas…".replace(",", ".")
                       if expected_rows else f"Lendo a planilha: {source_row:,} linhas…".replace(",", "."))

            def cell(name: str) -> Any:
                position = positions.get(name)
                return None if position is None or position >= len(row) else _clean_value(name, row[position])

            material = cell("material")
            if material in (None, ""):
                continue
            company = cell("company")
            if company not in (None, "") and str(company).strip() != INVENTORY_COMPANY:
                info["discarded_company_rows"] += 1
                continue
            plant = cell("plant")
            key = cell("center_material_key") or f"{plant or ''}{material}"
            year = cell("year")
            columns["k"].append(str(key).strip())
            columns["material"].append(str(material).strip().upper())
            columns["plant"].append(None if plant is None else str(plant).strip())
            columns["pod"].append(cell("product_offer_end_date"))
            season = cell("season")
            columns["season"].append(None if season in (None, "") else str(season).strip().upper())
            columns["year"].append(None if year is None else int(year))
            columns["value"].append(float(cell("fiscal_total_amount") or 0))
            columns["source_row"].append(source_row)

    if suffix == ".xlsb":
        from pyxlsb import open_workbook

        with open_workbook(str(path)) as workbook:
            def preview_of(name: str) -> list[list[Any]]:
                with workbook.get_sheet(name) as worksheet:
                    return [[c.v for c in r] for _, r in zip(range(contract.header_search_rows), worksheet.rows())]

            sheet, resolution = pick_sheet(list(workbook.sheets), preview_of)
            with workbook.get_sheet(sheet) as worksheet:
                consume(([c.v for c in r] for r in worksheet.rows()), resolution.header_row, resolution.positions)
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet, resolution = pick_sheet(
                list(workbook.sheetnames),
                lambda name: [list(r) for r in workbook[name].iter_rows(min_row=1, max_row=contract.header_search_rows, values_only=True)])
            consume((list(r) for r in workbook[sheet].iter_rows(values_only=True)), resolution.header_row, resolution.positions)
        finally:
            workbook.close()
    else:
        raise ValueError(f"Envie a Posição de Estoque em .xlsb ou .xlsx. {LAYOUT_HELP}")
    if not columns["k"]:
        raise ValueError(f"A aba {sheet} foi lida, mas não tem linhas de estoque da Empresa 7170.")
    info.update(sheet=sheet, header_row=resolution.header_row, rows=len(columns["k"]))
    table = pa.table({
        "k": pa.array(columns["k"], pa.string()), "material": pa.array(columns["material"], pa.string()),
        "plant": pa.array(columns["plant"], pa.string()), "pod": pa.array(columns["pod"], pa.date32()),
        "season": pa.array(columns["season"], pa.string()), "year": pa.array(columns["year"], pa.int32()),
        "value": pa.array(columns["value"], pa.float64()), "source_row": pa.array(columns["source_row"], pa.int64()),
    })
    return table, info


# ------------------------------------------------------------------ dados do sistema (cópia em memória)


def _previous_period(period: str) -> str:
    year, month = map(int, period.split("-"))
    return (date(year, month, 1) - timedelta(days=1)).strftime("%Y-%m")


def _system_snapshot(settings: Any, period: str) -> tuple[duckdb.DuckDBPyConnection, dict[str, Any]]:
    """Copia para um banco em memória só o que a conferência usa e solta o banco do sistema na hora."""
    previous = _previous_period(period)
    queries = {
        "inventory_rows": ("""select period, source_row, center_material_key, upper(trim(material)) material, plant,
                                     product_offer_end_date, season, year, fiscal_total_amount, aging_bucket
                                from inventory_rows where period in (?, ?)""", [period, previous]),
        "pass_step_sources": ("select * from pass_step_sources where period=?", [period]),
        "mb59_material_entries": ("select * from mb59_material_entries where period=?", [period]),
        "mb59_sku_entries": ("select * from mb59_sku_entries where period=?", [period]),
        "gr_collection_rows": ("select * from gr_collection_rows where period=?", [period]),
        "master_base_rows": ("select * from master_base_rows where period=?", [previous]),
        "mapping_rules": ("select group_key, key_value, value_1 from mapping_rules where group_key='aging'", []),
    }
    memory = duckdb.connect()
    info: dict[str, Any] = {"previous_period": previous}
    with connect(settings.path("database")) as connection:
        info["rules"] = load_rules(connection)
        try:
            info["previous_from_master"] = master_base_in_use(connection, previous)
        except Exception:  # noqa: BLE001 - sem base mestre
            info["previous_from_master"] = False
        for table, (sql, params) in queries.items():
            try:
                data = connection.execute(sql, params).fetch_arrow_table()
            except Exception:  # noqa: BLE001 - tabela ainda não criada nesta máquina
                data = None
            if data is not None:
                memory.register(f"_{table}", data)
                memory.execute(f"create table {table} as select * from _{table}")
                memory.unregister(f"_{table}")
            else:
                memory.execute(f"create table {table} (period varchar)")
    return memory, info


# ------------------------------------------------------------------ comparação


def _norm(field: str, column: str) -> str:
    """Valor comparável: Year agrupado (<2022), Season em maiúsculas, data como está."""
    if field == "year":
        return f"(case when {column} is null then null when cast({column} as integer) < 2022 then '<2022' else cast(cast({column} as integer) as varchar) end)"
    if field == "season":
        return f"upper(trim(cast({column} as varchar)))"
    return f"cast({column} as date)"


def _aging_case(memory: duckdb.DuckDBPyConnection, column: str, as_of: date) -> str:
    rules = memory.execute("select try_cast(key_value as integer), value_1 from mapping_rules order by 1").fetchall()
    rules = [r for r in rules if r[0] is not None]
    if not rules:
        return "null"
    days = f"date_diff('day', {column}, date '{as_of.isoformat()}')"
    parts = [f"when {days} < 0 then '0) Futures'"]
    bounds = [r[0] for r in rules[1:]] + [10**9]
    parts += [f"when {days} >= {start} and {days} < {end} then '{label}'" for (start, label), end in zip(rules, bounds)]
    return f"(case when {column} is null then null {' '.join(parts)} end)"


def _n(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def _pct(value: float) -> str:
    return f"{value:.1f}".replace(".", ",") + "%"


def _brl(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _mi(value: float) -> str:
    return f"R$ {value / 1e6:,.1f} mi".replace(",", "_").replace(".", ",").replace("_", ".")


def run_conference(settings: Any, conference_id: str, period: str, path: Path, progress: Progress | None = None) -> dict[str, Any]:
    """Executa as Fases 1 e 2 e devolve ``{"summary": ..., "diagnosis": ...}``."""
    report = progress or (lambda _p, _m: None)
    started = time.time()
    report(3, "Lendo a planilha manual…")
    try:  # estimativa para a barra de andamento: as linhas da competência no sistema
        with connect(settings.path("database")) as connection:
            expected = int(connection.execute("select count(*) from inventory_rows where period=?", [period]).fetchone()[0])
    except Exception:  # noqa: BLE001 - sem estimativa, a barra avança por etapa
        expected = 0
    planilha, read_info = read_position_file(settings, path, report, expected)
    pq.write_table(planilha, conferences_folder(settings) / f"{conference_id}.parquet")
    report(30, f"Planilha lida: {_n(read_info['rows'])} linhas. Copiando os dados do sistema…")
    memory, info = _system_snapshot(settings, period)
    memory.register("_p", planilha)
    memory.execute("create table p as select * from _p")
    year, month = map(int, period.split("-"))
    as_of = date(year, month, monthrange(year, month)[1])

    report(40, "Consultando o All Brazil do mês para as fontes de cada linha…")
    materials = {row[0] for row in memory.execute(
        "select distinct material from p union select distinct material from inventory_rows where period=?", [period]).fetchall() if row[0]}
    try:
        ab = pass_step_lookup_from_all_brazil(
            settings, as_of, materials,
            progress_callback=lambda fraction, _message: report(40 + round(18 * max(0.0, min(1.0, float(fraction)))),
                                                                "Consultando o All Brazil do mês para as fontes de cada linha…"))
        ab_lookup, ab_snapshot = ab["lookup"], ab.get("snapshot")
    except (FileNotFoundError, ValueError, OSError):
        ab_lookup, ab_snapshot = {}, None
    keys = sorted(ab_lookup)
    memory.register("_ab", pa.table({
        "material": pa.array(keys, pa.string()),
        "pod": pa.array([ab_lookup[k].get("product_offer_end_date") for k in keys], pa.date32()),
        "season": pa.array([ab_lookup[k].get("season") for k in keys], pa.string()),
        "year": pa.array([None if ab_lookup[k].get("year") is None else int(ab_lookup[k]["year"]) for k in keys], pa.int32()),
    }))
    memory.execute("create table ab as select * from _ab")

    report(60, "Comparando linha a linha…")
    memory.execute(f"""create table j as select
            coalesce(p.k, s.center_material_key) k, coalesce(p.material, s.material) material,
            substr(coalesce(p.material, s.material), 1, 10) sku,
            p.k is not null in_p, s.center_material_key is not null in_s,
            coalesce(p.value, 0) p_value, coalesce(s.fiscal_total_amount, 0) s_value,
            p.pod p_pod, p.season p_season, p.year p_year,
            s.product_offer_end_date s_pod, s.season s_season, s.year s_year, s.aging_bucket s_aging,
            src.product_offer_end_date src_pod, src.season src_season, src.year src_year
        from p full join (select * from inventory_rows where period=?) s on s.center_material_key=p.k
        left join pass_step_sources src on src.period=? and src.source_row=s.source_row""", [period, period])
    # Valores de cada fonte por linha (todas as opções do passo 1, as chaves do passo 2 e o All Brazil).
    field_cols = {"product_offer_end_date": "pod", "season": "season", "year": "year"}
    candidate_columns: dict[str, dict[str, str]] = {field: {} for field in FIELDS}
    previous = info["previous_period"]
    for field in FIELDS:
        col = field_cols[field]
        for option in STEP1_OPTIONS:
            spec = step1_sql(option, field)
            if spec is None:
                continue
            sql, kind, _tables = spec
            name = f"s1_{option}_{col}"
            try:
                memory.execute(f"create or replace table t_{name} as select upper(trim(cast(k as varchar))) as key, v from ({sql.replace('?', repr(period))}) t(k, v)")
            except duckdb.Error:
                continue  # opção sem dados nesta máquina (ex.: MB59 carregada antes das entradas por Material)
            memory.execute(f"alter table j add column {name} {'date' if col == 'pod' else 'varchar' if col == 'season' else 'integer'}")
            memory.execute(f"update j set {name}=t.v from t_{name} t where t.key=j.{'material' if kind == 'material' else 'sku'}")
            candidate_columns[field][f"step1:{option}"] = name
        for kind_option in ("ce_mat", "material"):
            sql, kind = step2_sql(kind_option, field, bool(info["previous_from_master"]))[0]
            name = f"s2_{kind_option}_{col}"
            memory.execute(f"create or replace table t_{name} as select cast(k as varchar) as key, v from ({sql.replace('?', repr(previous))}) t(k, v)")
            memory.execute(f"alter table j add column {name} {'date' if col == 'pod' else 'varchar' if col == 'season' else 'integer'}")
            memory.execute(f"update j set {name}=t.v from t_{name} t where t.key=j.{'k' if kind == 'ce_mat' else 'material'}")
            candidate_columns[field][f"step2:{kind_option}"] = name
        memory.execute(f"alter table j add column ab_{col} {'date' if col == 'pod' else 'varchar' if col == 'season' else 'integer'}")
        memory.execute(f"update j set ab_{col}=ab.{col} from ab where ab.material=j.material")
    for field in FIELDS:
        col = field_cols[field]
        sources = [*candidate_columns[field].values(), f"ab_{col}", f"s_{col}"]
        hit = " or ".join(f"{_norm(field, f'p_{col}')} = {_norm(field, s)}" for s in sources)
        memory.execute(f"alter table j add column no_source_{col} boolean")
        memory.execute(f"update j set no_source_{col} = in_p and in_s and p_{col} is not null and not coalesce({hit}, false)")

    summary = _phase1(memory, period, path.name, read_info, ab_snapshot, as_of, field_cols)
    report(80, "Testando as variações de regra (diagnóstico)…")
    diagnosis = _phase2(memory, info, field_cols, candidate_columns)
    summary["explanation"] = _explain(summary, diagnosis)
    summary["duration_seconds"] = round(time.time() - started, 1)
    memory.close()
    report(100, "Conferência concluída.")
    return {"summary": summary, "diagnosis": diagnosis}


def _phase1(memory: duckdb.DuckDBPyConnection, period: str, file_name: str, read_info: dict[str, Any],
            ab_snapshot: str | None, as_of: date, field_cols: dict[str, str]) -> dict[str, Any]:
    q = lambda sql, params=None: memory.execute(sql, params or []).fetchall()  # noqa: E731
    counts = q("""select count(*) filter (where in_p), count(*) filter (where in_s), count(*) filter (where in_p and in_s),
                         count(*) filter (where in_p and not in_s), count(*) filter (where in_s and not in_p),
                         coalesce(sum(p_value),0), coalesce(sum(s_value),0),
                         coalesce(sum(p_value) filter (where in_p and not in_s),0), coalesce(sum(s_value) filter (where in_s and not in_p),0)
                    from j""")[0]
    summary: dict[str, Any] = {
        "period": period, "file_name": file_name, "sheet": read_info.get("sheet"), "all_brazil_snapshot": ab_snapshot,
        "rows": {"planilha": counts[0], "sistema": counts[1], "comparadas": counts[2], "so_planilha": counts[3], "so_sistema": counts[4],
                 "descartadas_outra_empresa": read_info.get("discarded_company_rows", 0)},
        "totals": {"planilha": round(counts[5], 2), "sistema": round(counts[6], 2), "diferenca": round(counts[6] - counts[5], 2),
                   "so_planilha_valor": round(counts[7], 2), "so_sistema_valor": round(counts[8], 2)},
        "fields": {},
    }
    for field in FIELDS:
        col = field_cols[field]
        equal = f"{_norm(field, f'p_{col}')} is not distinct from {_norm(field, f's_{col}')}"
        row = q(f"""select count(*) filter (where {equal}), coalesce(sum(p_value) filter (where {equal}),0),
                           count(*) filter (where not ({equal})), coalesce(sum(p_value) filter (where not ({equal})),0),
                           count(*) filter (where not ({equal}) and no_source_{col}), coalesce(sum(p_value) filter (where not ({equal}) and no_source_{col}),0)
                      from j where in_p and in_s""")[0]
        origins = q(f"""select coalesce(src_{col}, 'sem origem registrada'), count(*), coalesce(sum(p_value),0) from j
                         where in_p and in_s and not ({equal}) group by 1 order by 3 desc""")
        examples = q(f"""select k, round(p_value, 2), cast(p_{col} as varchar), cast(s_{col} as varchar), src_{col}, no_source_{col}
                          from j where in_p and in_s and not ({equal}) order by p_value desc limit 15""")
        compared = max(1, summary["rows"]["comparadas"])
        summary["fields"][field] = {
            "label": FIELD_LABELS[field],
            "iguais": {"linhas": row[0], "valor": round(row[1], 2), "pct": round(100 * row[0] / compared, 2)},
            "diferentes": {"linhas": row[2], "valor": round(row[3], 2)},
            "sem_origem": {"linhas": row[4], "valor": round(row[5], 2)},
            "por_origem": [{"origem": o[0], "linhas": o[1], "valor": round(o[2], 2)} for o in origins],
            "exemplos": [{"ce_mat": e[0], "valor": e[1], "planilha": e[2], "sistema": e[3], "origem_sistema": e[4], "sem_origem": bool(e[5])}
                         for e in examples],
        }
    aging = _aging_case(memory, "p_pod", as_of)
    reports = {
        "AGING": q(f"""select coalesce(a.k, b.k) k, coalesce(a.v,0), coalesce(b.v,0) from
                        (select {aging} k, sum(p_value) v from j where in_p group by 1) a
                        full join (select s_aging k, sum(s_value) v from j where in_s group by 1) b on a.k=b.k order by 1"""),
        "Season": q(f"""select coalesce(a.k, b.k) k, coalesce(a.v,0), coalesce(b.v,0) from
                         (select {_norm('season', 'p_season')} k, sum(p_value) v from j where in_p group by 1) a
                         full join (select {_norm('season', 's_season')} k, sum(s_value) v from j where in_s group by 1) b on a.k is not distinct from b.k order by 1"""),
        "Year": q(f"""select coalesce(a.k, b.k) k, coalesce(a.v,0), coalesce(b.v,0) from
                       (select {_norm('year', 'p_year')} k, sum(p_value) v from j where in_p group by 1) a
                       full join (select {_norm('year', 's_year')} k, sum(s_value) v from j where in_s group by 1) b on a.k is not distinct from b.k
                       order by case when coalesce(a.k, b.k)='<2022' then '0' else coalesce(a.k, b.k) end"""),
    }
    summary["reports"] = {name: [{"faixa": r[0] or "Não informado", "planilha": round(r[1], 2), "sistema": round(r[2], 2),
                                  "diferenca": round(r[2] - r[1], 2)} for r in rows] for name, rows in reports.items()}
    return summary


def _phase2(memory: duckdb.DuckDBPyConnection, info: dict[str, Any], field_cols: dict[str, str],
            candidate_columns: dict[str, dict[str, str]]) -> dict[str, Any]:
    """Testa as combinações (passo 1 × passo 2) contra a planilha; All Brazil é sempre o 3º passo."""
    rules = info["rules"]
    result: dict[str, Any] = {"rules": rules, "fields": {}}
    for field in FIELDS:
        col = field_cols[field]
        available = [option for option in STEP1_OPTIONS
                     if option == "none" or memory.execute(f"select count(*) from j where {candidate_columns[field].get(f'step1:{option}', 'null')} is not null").fetchone()[0] > 0]
        combos = []
        for option1 in available:
            for option2 in STEP2_OPTIONS:
                parts = ([] if option1 == "none" else [candidate_columns[field][f"step1:{option1}"]])
                parts += {"ce_mat": [candidate_columns[field]["step2:ce_mat"]], "material": [candidate_columns[field]["step2:material"]],
                          "ce_mat_material": [candidate_columns[field]["step2:ce_mat"], candidate_columns[field]["step2:material"]]}[option2]
                predicted = "coalesce(" + ", ".join([*parts, f"ab_{col}"]) + ")"
                differ = f"{_norm(field, f'p_{col}')} is distinct from {_norm(field, predicted)}"
                row = memory.execute(f"""select count(*) filter (where {differ}), coalesce(sum(p_value) filter (where {differ}),0),
                                                count(*) filter (where {differ} and not no_source_{col}),
                                                coalesce(sum(p_value) filter (where {differ} and not no_source_{col}),0)
                                           from j where in_p and in_s""").fetchone()
                combos.append({"step1": option1, "step2": option2, "diferentes_linhas": row[0], "diferentes_valor": round(row[1], 2),
                               "explicaveis_linhas": row[2], "explicaveis_valor": round(row[3], 2)})
        combos.sort(key=lambda c: (c["diferentes_valor"], c["diferentes_linhas"]))
        current = next((c for c in combos if c["step1"] == rules[field]["step1"] and c["step2"] == rules[field]["step2"]), None)
        best = combos[0] if combos else None
        suggestion = None
        if best and current and (best["step1"], best["step2"]) != (current["step1"], current["step2"]):
            gain = current["diferentes_valor"] - best["diferentes_valor"]
            if gain >= MIN_GAIN_VALUE and gain >= MIN_GAIN_SHARE * max(current["diferentes_valor"], 1):
                suggestion = {"field": field, "step1": best["step1"], "step2": best["step2"],
                              "step1_label": STEP1_OPTIONS[best["step1"]], "step2_label": STEP2_OPTIONS[best["step2"]],
                              "ganho_valor": round(gain, 2), "ganho_linhas": current["diferentes_linhas"] - best["diferentes_linhas"],
                              "antes": current, "depois": best}
        result["fields"][field] = {
            "label": FIELD_LABELS[field], "atual": current, "melhor": best, "sugestao": suggestion,
            "combinacoes": [{**c, "step1_label": STEP1_OPTIONS[c["step1"]], "step2_label": STEP2_OPTIONS[c["step2"]]} for c in combos[:8]],
        }
    return result


def _explain(summary: dict[str, Any], diagnosis: dict[str, Any]) -> list[str]:
    """Explicação em linguagem simples, montada só com os números da conferência (nada inventado)."""
    rows, totals = summary["rows"], summary["totals"]
    lines = [f"Foram comparadas {_n(rows['comparadas'])} linhas (CE & MAT) da competência {summary['period']}."]
    if abs(totals["diferenca"]) < 0.01:
        lines.append(f"O valor total é igual nos dois lados ({_brl(totals['sistema'])}).")
    else:
        lines.append(f"O valor total difere em {_brl(totals['diferenca'])} (sistema {_brl(totals['sistema'])} x planilha {_brl(totals['planilha'])}).")
    if rows["so_planilha"] or rows["so_sistema"]:
        lines.append(f"{rows['so_planilha']} linha(s) só na planilha e {rows['so_sistema']} só no sistema.")
    for field, data in summary["fields"].items():
        diff = data["diferentes"]
        if not diff["linhas"]:
            lines.append(f"{data['label']}: 100% igual à planilha.")
            continue
        text = f"{data['label']}: {_pct(data['iguais']['pct'])} igual; {_n(diff['linhas'])} linhas diferentes ({_mi(diff['valor'])})"
        if data["sem_origem"]["linhas"]:
            text += (f", das quais {_n(data['sem_origem']['linhas'])} ({_mi(data['sem_origem']['valor'])}) têm na planilha um valor "
                     "que não existe em nenhuma fonte (Base GR, mês anterior ou All Brazil) — provável erro na planilha")
        lines.append(text + ".")
        # Onde estão as diferenças: de qual fonte veio o valor do sistema nas linhas diferentes.
        origins = sorted(data.get("por_origem") or [], key=lambda o: -abs(o["valor"]))[:3]
        if origins:
            lines.append(f"{data['label']} — diferenças por origem do valor no sistema: "
                         + "; ".join(f"{o['origem']}: {_n(o['linhas'])} linhas ({_mi(o['valor'])})" for o in origins) + ".")
        suggestion = diagnosis["fields"][field].get("sugestao")
        if suggestion:
            lines.append(f"Sugestão para {data['label']}: {suggestion['step1_label']} → {suggestion['step2_label']} reduziria a "
                         f"diferença em {_mi(suggestion['ganho_valor'])}. Só vale depois de aprovada por um administrador.")
    # Onde o valor se concentra nos relatórios (sistema − planilha), só faixas com diferença de pelo menos R$ 50 mil.
    for name, report in (summary.get("reports") or {}).items():
        top = sorted((r for r in report if abs(r["diferenca"]) >= 5e4), key=lambda r: -abs(r["diferenca"]))[:3]
        if top:
            lines.append(f"Totais por {name} (sistema − planilha), maiores diferenças: "
                         + "; ".join(f"{r['faixa']} {_mi(r['diferenca'])}" for r in top) + ".")
    if not any(diagnosis["fields"][f].get("sugestao") for f in FIELDS):
        lines.append("Nenhuma variação das regras reduz a diferença de forma relevante: a regra atual é a mais próxima da planilha.")
    return lines


# ------------------------------------------------------------------ registro


def save_run(settings: Any, conference_id: str, *, period: str, file_name: str, sha: str, status: str, message: str,
             created_by: str, summary: dict[str, Any] | None = None, diagnosis: dict[str, Any] | None = None,
             finished: bool = False) -> None:
    with connect(settings.path("database")) as connection:
        ensure_conference_schema(connection)
        connection.execute("delete from conference_runs where conference_id=?", [conference_id])
        connection.execute(
            """insert into conference_runs(conference_id, period, file_name, source_sha256, status, message, created_by,
                   finished_at, summary_json, diagnosis_json) values (?,?,?,?,?,?,?,?,?,?)""",
            [conference_id, period, file_name, sha, status, message, created_by,
             datetime.now() if finished else None,
             json.dumps(summary, ensure_ascii=False, default=str) if summary else None,
             json.dumps(diagnosis, ensure_ascii=False, default=str) if diagnosis else None],
        )


class ConferenceCancelled(Exception):
    """Conferência cancelada pelo usuário: o andamento para no próximo ponto de controle."""


def delete_run(settings: Any, conference_id: str) -> None:
    """Apaga uma conferência cancelada (registro e linhas lidas): o último resultado da competência volta a valer."""
    with connect(settings.path("database")) as connection:
        ensure_conference_schema(connection)
        connection.execute("delete from conference_runs where conference_id=?", [conference_id])
    (conferences_folder(settings) / f"{conference_id}.parquet").unlink(missing_ok=True)


def latest_run(settings: Any, period: str | None = None, conference_id: str | None = None) -> dict[str, Any] | None:
    with connect(settings.path("database")) as connection:
        ensure_conference_schema(connection)
        where, params = ("conference_id=?", [conference_id]) if conference_id else ("period=?", [period])
        row = connection.execute(
            f"""select conference_id, period, file_name, status, message, created_by, created_at, finished_at, summary_json, diagnosis_json
                  from conference_runs where {where} order by created_at desc limit 1""", params).fetchone()
    if not row:
        return None
    keys = ("conference_id", "period", "file_name", "status", "message", "created_by", "created_at", "finished_at", "summary", "diagnosis")
    item = dict(zip(keys, row, strict=True))
    for key in ("created_at", "finished_at"):
        item[key] = item[key].isoformat(timespec="seconds") if hasattr(item[key], "isoformat") else item[key]
    for key in ("summary", "diagnosis"):
        item[key] = json.loads(item[key]) if item[key] else None
    return item


def short_facts(run: dict[str, Any]) -> dict[str, Any]:
    """Resumo compacto para o Optimus e para o controle automático (sem linhas da planilha)."""
    summary, diagnosis = run.get("summary") or {}, run.get("diagnosis") or {}
    return {
        "period": run.get("period"), "file_name": run.get("file_name"),
        "rows": summary.get("rows"), "totals": summary.get("totals"),
        "fields": {f: {"label": d["label"], "iguais_pct": d["iguais"]["pct"], "diferentes": d["diferentes"], "sem_origem": d["sem_origem"]}
                   for f, d in (summary.get("fields") or {}).items()},
        "suggestions": [d["sugestao"] for d in (diagnosis.get("fields") or {}).values() if d.get("sugestao")],
        "explanation": summary.get("explanation"),
    }
