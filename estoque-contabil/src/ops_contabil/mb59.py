from __future__ import annotations

import math
from calendar import monthrange
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable

import pyarrow as pa
from openpyxl import load_workbook

from .audit import sha256
from .db import connect
from .ingestion.schema_resolver import load_contracts, resolve_headers
from .sources.all_brazil import pass_step_lookup_from_all_brazil


def ensure_mb59_schema(connection: Any) -> None:
    connection.execute(
        """
        create table if not exists mb59_rows (
            period varchar not null,
            center_material_key varchar not null,
            material varchar not null,
            plant varchar not null,
            product_offer_end_date date,
            season varchar,
            year integer,
            season_year varchar,
            posting_date date,
            source_row bigint,
            primary key(period, center_material_key)
        );
        create table if not exists mb59_imports (
            period varchar primary key,
            source_path varchar not null,
            source_sha256 varchar not null,
            source_sheet varchar not null,
            header_row integer not null,
            source_rows bigint not null,
            lookup_rows bigint not null,
            schema_fingerprint varchar not null,
            imported_at timestamp not null default current_timestamp
        );
        """
    )


def _text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not math.isnan(float(value)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    return None


def _year(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def import_mb59_xlsx(settings: Any, period: str, path: Path) -> dict[str, Any]:
    datetime.strptime(period, "%Y-%m")
    path = path.resolve()
    contract = load_contracts(settings.root / "config" / "schemas.yaml")["mb59_base"]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        best: tuple[Any, Any] | None = None
        for worksheet in workbook.worksheets:
            preview = [
                list(row)
                for row in worksheet.iter_rows(
                    min_row=1, max_row=contract.header_search_rows, values_only=True
                )
            ]
            try:
                candidate = resolve_headers(preview, contract)
            except Exception:
                continue
            if best is None or len(candidate.positions) > len(best[1].positions):
                best = (worksheet, candidate)
        if best is None:
            raise ValueError("Nenhuma aba da exportação MB59 atende ao contrato de colunas.")
        worksheet, resolution = best
        selected: dict[str, dict[str, Any]] = {}
        source_rows = 0
        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=resolution.header_row + 1, values_only=True),
            start=resolution.header_row + 1,
        ):
            def raw(field: str) -> Any:
                position = resolution.positions.get(field)
                return None if position is None or position >= len(row) else row[position]

            material = _text(raw("material"))
            plant = _text(raw("plant"))
            if not material or not plant:
                continue
            source_rows += 1
            posting_date = _date(raw("posting_date"))
            product_offer_end_date = _date(raw("product_offer_end_date"))
            season = _text(raw("season"))
            year = _year(raw("year"))
            season_year = _text(raw("season_year")) or (
                f"{season}{year}" if season is not None and year is not None else None
            )
            key = f"{plant}{material}"
            candidate = {
                "period": period,
                "center_material_key": key,
                "material": material,
                "plant": plant,
                "product_offer_end_date": product_offer_end_date,
                "season": season,
                "year": year,
                "season_year": season_year,
                "posting_date": posting_date,
                "source_row": source_row,
            }
            current = selected.get(key)
            if current is None or (posting_date or date.min) >= (current["posting_date"] or date.min):
                if current is not None:
                    for field in ("product_offer_end_date", "season", "year", "season_year"):
                        if candidate[field] is None:
                            candidate[field] = current[field]
                selected[key] = candidate
        rows = list(selected.values())
        batch = pa.Table.from_pylist(rows) if rows else None
        source_hash = sha256(path)
        with connect(settings.path("database")) as connection:
            ensure_mb59_schema(connection)
            connection.execute("begin")
            try:
                connection.execute("delete from mb59_rows where period=?", [period])
                if batch is not None:
                    connection.register("mb59_batch", batch)
                    connection.execute("insert into mb59_rows select * from mb59_batch")
                connection.execute("delete from mb59_imports where period=?", [period])
                connection.execute(
                    """insert into mb59_imports(period,source_path,source_sha256,source_sheet,
                       header_row,source_rows,lookup_rows,schema_fingerprint)
                       values (?,?,?,?,?,?,?,?)""",
                    [period, str(path), source_hash, worksheet.title, resolution.header_row,
                     source_rows, len(rows), resolution.fingerprint],
                )
                connection.execute("commit")
            except Exception:
                connection.execute("rollback")
                raise
        return {
            "status": "imported",
            "period": period,
            "source_rows": source_rows,
            "lookup_rows": len(rows),
            "path": str(path),
            "header_row": resolution.header_row,
        }
    finally:
        workbook.close()


def _previous_period(period: str) -> str:
    year_number, month_number = map(int, period.split("-"))
    first_day = date(year_number, month_number, 1)
    return (first_day - timedelta(days=1)).strftime("%Y-%m")


def enrich_inventory_pass_step(
    settings: Any,
    period: str,
    as_of: date | None = None,
    progress_callback: Callable[[float, str], None] | None = None,
) -> dict[str, Any]:
    """Preenche campos PASSO A PASSO com prioridade operacional homologada.

    Prioridade por campo:
    1. Base GR / MB59 da competência, usando CE & MAT (Centro + Material).
    2. Posição de Estoque do mês anterior, usando CE & MAT.
    3. All Brazil sem Centro: primeiro Material completo; se ele não existir no
       cadastro, fallback pelo Estilo-Cor de 10 caracteres.

    Se nenhuma das três fontes trouxer valor, o campo permanece NULL e deve ser
    apresentado ao usuário como "Não localizado". Nenhum valor é inferido ou inventado.
    """
    datetime.strptime(period, "%Y-%m")
    if as_of is None:
        year_number, month_number = map(int, period.split("-"))
        as_of = date(year_number, month_number, monthrange(year_number, month_number)[1])
    previous_period = _previous_period(period)

    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        inventory_exists = bool(
            connection.execute(
                """select count(*) > 0 from information_schema.tables
                   where lower(table_name)='inventory_rows'"""
            ).fetchone()[0]
        )
        if not inventory_exists:
            return {
                "status": "inventory_missing",
                "period": period,
                "previous_period": previous_period,
                "rows": 0,
                "message": "A Base de Estoque ainda não foi carregada.",
            }

        current_rows = connection.execute(
            """select source_row, center_material_key,
                      upper(trim(material)) as material_key,
                      product_offer_end_date, season, year
               from inventory_rows
               where period=?
               order by source_row""",
            [period],
        ).fetchall()
        if not current_rows:
            return {
                "status": "inventory_missing",
                "period": period,
                "previous_period": previous_period,
                "rows": 0,
                "message": f"A competência {period} não possui linhas na Base de Estoque.",
            }

        mb_rows = connection.execute(
            """select center_material_key, product_offer_end_date, season, year
               from mb59_rows where period=?""",
            [period],
        ).fetchall()
        mb_lookup = {
            str(row[0]): {
                "product_offer_end_date": row[1],
                "season": row[2],
                "year": int(row[3]) if row[3] is not None else None,
            }
            for row in mb_rows
            if row[0] not in (None, "")
        }

        previous_rows = connection.execute(
            """select center_material_key, product_offer_end_date, season, year
               from inventory_rows where period=?""",
            [previous_period],
        ).fetchall()
        previous_lookup = {
            str(row[0]): {
                "product_offer_end_date": row[1],
                "season": row[2],
                "year": int(row[3]) if row[3] is not None else None,
            }
            for row in previous_rows
            if row[0] not in (None, "")
        }

    material_keys = {
        str(row[2]).strip()
        for row in current_rows
        if row[2] not in (None, "")
    }
    try:
        all_brazil = pass_step_lookup_from_all_brazil(
            settings,
            as_of,
            material_keys,
            progress_callback=(
                (lambda fraction, message: progress_callback(fraction * 0.42, message))
                if progress_callback is not None
                else None
            ),
        )
    except (FileNotFoundError, ValueError, OSError):
        all_brazil = {"lookup": {}, "snapshot": None, "snapshot_date": None, "conflict_keys": 0}
    all_brazil_lookup = all_brazil["lookup"]

    field_sources = {
        "product_offer_end_date": {"mb59": 0, "previous_inventory": 0, "all_brazil": 0, "not_found": 0},
        "season": {"mb59": 0, "previous_inventory": 0, "all_brazil": 0, "not_found": 0},
        "year": {"mb59": 0, "previous_inventory": 0, "all_brazil": 0, "not_found": 0},
    }
    updates: list[dict[str, Any]] = []

    def choose(field: str, center_material_key: str, material_key: str) -> Any:
        candidates = (
            ("mb59", mb_lookup.get(center_material_key, {}).get(field)),
            ("previous_inventory", previous_lookup.get(center_material_key, {}).get(field)),
            ("all_brazil", all_brazil_lookup.get(material_key, {}).get(field)),
        )
        for source, value in candidates:
            if value not in (None, ""):
                field_sources[field][source] += 1
                return value
        field_sources[field]["not_found"] += 1
        return None

    total_current_rows = max(1, len(current_rows))
    if progress_callback is not None:
        progress_callback(0.44, f"Aplicando prioridades do PASSO A PASSO em {len(current_rows):,} linhas.")
    for processed, (source_row, center_material_key, material_key, existing_offer, existing_season, existing_year) in enumerate(current_rows, start=1):
        if progress_callback is not None and (processed % 5000 == 0 or processed >= total_current_rows):
            progress_callback(
                0.44 + (processed / total_current_rows) * 0.46,
                f"PASSO A PASSO: {processed:,} de {total_current_rows:,} linhas processadas.",
            )
        center_key = "" if center_material_key is None else str(center_material_key)
        material_only_key = "" if material_key is None else str(material_key).strip().upper()
        offer_end = choose("product_offer_end_date", center_key, material_only_key)
        season = choose("season", center_key, material_only_key)
        year_value = choose("year", center_key, material_only_key)
        year_number = int(year_value) if year_value is not None else None
        updates.append({
            "source_row": int(source_row),
            "product_offer_end_date": offer_end,
            "season": None if season is None else str(season).strip(),
            "year": year_number,
        })

    if progress_callback is not None:
        progress_callback(0.92, "Gravando resultado do PASSO A PASSO na Base de Estoque.")
    batch = pa.Table.from_pylist(updates)
    with connect(settings.path("database")) as connection:
        connection.register("pass_step_batch", batch)
        connection.execute("begin")
        try:
            connection.execute(
                """update inventory_rows as inventory set
                     product_offer_end_date=step.product_offer_end_date,
                     season=step.season,
                     year=step.year
                   from pass_step_batch as step
                   where inventory.period=? and inventory.source_row=step.source_row""",
                [period],
            )
            connection.execute("commit")
        except Exception:
            connection.execute("rollback")
            raise

    rows = len(current_rows)
    if progress_callback is not None:
        progress_callback(1.0, f"PASSO A PASSO concluído em {rows:,} linhas.")
    return {
        "status": "enriched",
        "period": period,
        "previous_period": previous_period,
        "as_of": as_of.isoformat(),
        "rows": rows,
        "mb59_ce_mat_matches": sum(1 for row in current_rows if str(row[1]) in mb_lookup),
        "previous_ce_mat_matches": sum(1 for row in current_rows if str(row[1]) in previous_lookup),
        "all_brazil_material_matches": sum(1 for row in current_rows if str(row[2]).strip() in all_brazil_lookup),
        "all_brazil_snapshot": all_brazil.get("snapshot"),
        "all_brazil_snapshots_available": all_brazil.get("snapshots_available", 0),
        "all_brazil_snapshots_scanned": all_brazil.get("snapshots_scanned", 0),
        "all_brazil_conflict_keys": all_brazil.get("conflict_keys", 0),
        "field_sources": field_sources,
        "not_found": {
            field: counts["not_found"]
            for field, counts in field_sources.items()
        },
        "message": (
            "PASSO A PASSO aplicado com prioridade MB59 por CE & MAT, "
            f"posição {previous_period} por CE & MAT e All Brazil por Material. "
            "Campos sem correspondência foram marcados como Não localizado."
        ),
    }


def enrich_inventory_from_mb59(
    settings: Any,
    period: str,
    as_of: date | None = None,
) -> dict[str, Any]:
    """Compatibilidade: usa a rotina PASSO A PASSO sem inferir valores ausentes."""
    return enrich_inventory_pass_step(settings, period, as_of)

def has_mb59_period(settings: Any, period: str) -> bool:
    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        return connection.execute(
            "select count(*)>0 from mb59_imports where period=?", [period]
        ).fetchone()[0]


def mb59_status(settings: Any, period: str) -> dict[str, Any]:
    with connect(settings.path("database")) as connection:
        ensure_mb59_schema(connection)
        row = connection.execute(
            """select imports.source_path, imports.lookup_rows, imports.imported_at,
                      count(rows.center_material_key) as indexed_rows,
                      count(rows.product_offer_end_date) as offer_rows,
                      count(rows.season) as season_rows,
                      count(rows.year) as year_rows
                 from mb59_imports imports
                 left join mb59_rows rows on rows.period=imports.period
                where imports.period=?
                group by imports.source_path, imports.lookup_rows, imports.imported_at""",
            [period],
        ).fetchone()
    if row is None:
        return {
            "period": period,
            "updated": False,
            "message": "MB59 pendente para esta competência",
        }
    return {
        "period": period,
        "updated": True,
        "source_path": row[0],
        "file_name": Path(row[0]).name,
        "lookup_rows": int(row[1]),
        "imported_at": row[2].isoformat() if hasattr(row[2], "isoformat") else str(row[2]),
        "indexed_rows": int(row[3]),
        "offer_rows": int(row[4]),
        "season_rows": int(row[5]),
        "year_rows": int(row[6]),
        "pass_step_value_rows": max(int(row[4]), int(row[5]), int(row[6])),
        "message": "MB59 atualizada para esta competência",
    }
